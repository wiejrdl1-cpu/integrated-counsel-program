import io
import re
import json
import sys
import os
import subprocess
import copy
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, filedialog, messagebox
from pathlib import Path
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Protection

from pastel_theme import PALETTE, apply_button_palette, apply_pastel_theme, style_plain_widgets


try:
    import msoffcrypto
except Exception:
    msoffcrypto = None


_NativeButton = tk.Button


class RaisedButton(_NativeButton):
    """샘플 UI와 같은 얇은 테두리의 평면형 공통 버튼입니다."""
    def __init__(self, master=None, **kwargs):
        kwargs.pop('style', None)
        kwargs['relief'] = 'solid'
        kwargs['bd'] = 1
        kwargs['highlightthickness'] = 1
        kwargs.setdefault('font', ('Malgun Gothic', 10, 'bold'))
        kwargs.setdefault('padx', 12)
        kwargs.setdefault('pady', 7)
        kwargs.setdefault('cursor', 'hand2')
        apply_button_palette(kwargs)
        super().__init__(master, **kwargs)


tk.Button = RaisedButton
ttk.Button = RaisedButton

SERVICE_OPTIONS = [
    '기초연금', '차상위계층확인', '차상위본인부담경감', '장애인연금', '차상위장애인',
    '저소득 한부모', '경기도 한부모', '초중고교육비', '차상위자활', '청년내일저축계좌', '타법의료급여(유공자)'
]
BASIC_PENSION_HOUSEHOLDS = ['노인단독가구', '노인부부1인', '노인부부1인->2인(배우자 추가신청)', '노인부부동시신청', '직접입력']
SINGLE_PARENT_HOUSEHOLDS = ['부자가구', '모자가구', '청소년한부모가구', '조손가구', '직접입력']
YOUTH_RENT_HOUSEHOLDS = ['청년 단독가구', '혼인한 청년가구', '기타가구', '직접입력']
GENERAL_HOUSEHOLDS = ['단독가구', '부부가구', '부부+자녀가구', '한부모가구', '조손가구', '기타가구', '직접입력']
HOME_OPTIONS = [
    '자가', '사용대차전체', '사용대차부분', '사용대차(급여미지급)', '시설입소',
    '공공건설임대(영구임대)', '공공건설임대(국민임대)', '공공건설임대(매입임대)',
    '전세임대', '민간(전세)', '민간(보증부월세)', '민간(월세)', '기타(직접입력)'
]
OLD_HOME_OPTION_MAP = {
    '사용대차': '사용대차전체',
    '사용대차(전체)': '사용대차전체',
    '사용대차(부분)': '사용대차부분',
    '영구임대주택': '공공건설임대(영구임대)',
    '국민임대주택': '공공건설임대(국민임대)',
    '전세': '민간(전세)',
    '월세': '민간(보증부월세)',
    '고시원': '기타(직접입력)',
    '직접입력': '기타(직접입력)',
}
THRESHOLDS = {
    '기초연금': {1: 2470000, 2: 3952000},
    '차상위계층': {1: 1282119, 2: 2099646, 3: 2679518, 4: 3247369, 5: 3778360, 6: 4277976, 7: 4765161},
    '저소득 한부모': {2: 2729540, 3: 3483373, 4: 4221580, 5: 4911867, 6: 5561369, 7: 6173447},
    '경기도 한부모': {1: 2564238, 2: 4199292, 3: 5359036, 4: 6494738, 5: 7556720, 6: 8555952, 7: 9530322},
    '장애인연금': {1: 1400000, 2: 2240000},
    '초중고교육비': {1: 2051390, 2: 3359434, 3: 4287229, 4: 5195790, 5: 6045376, 6: 6844761, 7: 7624258},
    '타법의료급여_일반가구': {1: 2051390, 2: 3359434, 3: 4287229, 4: 5195790, 5: 6045375, 6: 6844762},
    '타법의료급여_취약가구': {1: 2564238, 2: 4199292, 3: 5359036, 4: 6494738, 5: 7556719, 6: 8555952}
}
SUPPORTER_ABILITY_BY_SIZE = {1: 3077086, 2: 5039150, 3: 6430843, 4: 7793686, 5: 9088063, 6: 10267142, 7: 11418180}

SERVICE_TO_STANDARD = {
    '기초연금': '기초연금',
    '차상위계층확인': '차상위계층',
    '차상위본인부담경감': '차상위계층',
    '차상위장애인': '차상위계층',
    '차상위자활': '차상위계층',
    '청년내일저축계좌': '차상위계층',
    '장애인연금': '장애인연금',
    '저소득 한부모': '저소득 한부모',
    '경기도 한부모': '경기도 한부모',
    '초중고교육비': '초중고교육비',
    '타법의료급여(유공자)': '타법의료급여_일반가구'
}

class WrongPasswordError(Exception):
    pass

class BenefitPopup(tk.Toplevel):
    def __init__(self, master, current):
        super().__init__(master)
        self.title('수급자 책정 여부')
        self.resizable(False, False)
        self.result = None
        self.vars = {k: tk.BooleanVar(value=k in current) for k in ['생계', '의료', '주거', '교육']}
        ttk.Label(self, text='책정된 보장을 선택하세요').pack(anchor='w', padx=14, pady=(14, 8))
        for k, v in self.vars.items():
            ttk.Checkbutton(self, text=k, variable=v).pack(anchor='w', padx=18, pady=2)
        f = ttk.Frame(self)
        f.pack(fill='x', padx=12, pady=12)
        ttk.Button(f, text='확인', command=self.ok).pack(side='left', padx=4)
        ttk.Button(f, text='취소', command=self.cancel).pack(side='left', padx=4)
        self.transient(master)
        self.grab_set()
        self.lift()
        self.focus_force()
        self.protocol('WM_DELETE_WINDOW', self.cancel)
        self.wait_window(self)
    def ok(self):
        self.result = [k for k, v in self.vars.items() if v.get()]
        self.destroy()
    def cancel(self):
        self.result = None
        self.destroy()

class MultiIncomePopup(tk.Toplevel):
    def __init__(self, master, services, current):
        super().__init__(master)
        self.title('보장별 소득인정액 입력')
        self.resizable(False, False)
        # 확인 또는 취소 전까지 다른 창 뒤로 숨지 않도록 항상 맨 앞에 표시합니다.
        self.attributes('-topmost', True)
        self.result = None
        self.entries = {}
        ttk.Label(self, text='보장별 소득인정액을 입력하세요').pack(anchor='w', padx=14, pady=(14, 8))
        body = ttk.Frame(self)
        body.pack(fill='both', padx=12, pady=4)
        for i, s in enumerate(services):
            ttk.Label(body, text=s, width=18).grid(row=i, column=0, sticky='w', padx=4, pady=3)
            e = ttk.Entry(body, width=18)
            e.grid(row=i, column=1, sticky='w', padx=4, pady=3)
            if s in current and current[s]:
                e.insert(0, current[s])
            self.entries[s] = e
        btns = ttk.Frame(self)
        btns.pack(fill='x', padx=12, pady=12)
        ttk.Button(btns, text='확인', command=self.ok).pack(side='left', padx=4)
        ttk.Button(btns, text='취소', command=self.cancel).pack(side='left', padx=4)
        self.transient(master)
        self.grab_set()
        self.protocol('WM_DELETE_WINDOW', self.cancel)
        self.wait_window(self)
    def ok(self):
        out = {}
        for k, e in self.entries.items():
            v = re.sub(r'[^0-9]', '', e.get().strip())
            out[k] = f"{int(v):,}" if v else ''
        self.result = out
        self.destroy()
    def cancel(self):
        self.result = None
        self.destroy()

class FileRow:
    def __init__(self, master, idx, remove_callback, picked_callback=None):
        self.frame = ttk.Frame(master)
        self.path = ''
        self.display_name = f'가구원 {idx}'
        self.picked_callback = picked_callback
        self.lbl_no = ttk.Label(self.frame, text=self.display_name)
        self.lbl_no.grid(row=0, column=0, padx=4, pady=4, sticky='w')
        ttk.Button(self.frame, text='엑셀 선택', command=self.pick).grid(row=0, column=1, padx=4, pady=4)
        self.lbl_path = ttk.Label(self.frame, text='선택된 파일 없음', width=60)
        self.lbl_path.grid(row=0, column=2, padx=4, pady=4, sticky='w')
        ttk.Button(self.frame, text='삭제', command=lambda: remove_callback(self)).grid(row=0, column=3, padx=4, pady=4)
    def pick(self):
        path = filedialog.askopenfilename(filetypes=[('Excel files', '*.xlsx *.xls')])
        if path:
            self.path = path
            self.lbl_path.config(text=Path(path).name)
            if self.picked_callback:
                self.picked_callback(self)
    def set_display_name(self, name):
        self.display_name = name
        self.lbl_no.config(text=name)

class App:
    def get_config_path(self):
        """엑셀 비밀번호 저장 파일 경로를 반환합니다.
        exe 실행 시 PyInstaller 임시폴더가 아니라 exe가 있는 폴더에 저장합니다.
        exe 폴더에 저장이 어려우면 사용자 AppData 폴더를 사용합니다.
        """
        try:
            if getattr(sys, 'frozen', False):
                base = Path(sys.executable).resolve().parent
            else:
                base = Path(__file__).resolve().parent

            test_file = base / '.consult_auto_write_test'
            try:
                test_file.write_text('ok', encoding='utf-8')
                test_file.unlink(missing_ok=True)
                return base / 'consult_auto_settings.json'
            except Exception:
                appdata = os.environ.get('APPDATA') or str(Path.home())
                fallback = Path(appdata) / 'consult_auto'
                fallback.mkdir(parents=True, exist_ok=True)
                return fallback / 'consult_auto_settings.json'
        except Exception:
            fallback = Path.home() / 'consult_auto_settings.json'
            return fallback

    def load_saved_password(self):
        try:
            path = self.get_config_path()
            if path.exists():
                data = json.loads(path.read_text(encoding='utf-8'))
                return data.get('excel_password', '')
        except Exception:
            pass
        return ''

    def load_saved_gu(self):
        try:
            path = self.get_config_path()
            if path.exists():
                data = json.loads(path.read_text(encoding='utf-8'))
                saved = data.get('selected_gu', '분당구')
                if saved in ('분당구', '중원구', '수정구'):
                    return saved
        except Exception:
            pass
        return '분당구'

    def save_selected_gu(self, event=None):
        selected = self.gu_var.get().strip()
        if selected not in ('분당구', '중원구', '수정구'):
            return
        try:
            path = self.get_config_path()
            data = {}
            if path.exists():
                try:
                    data = json.loads(path.read_text(encoding='utf-8'))
                except Exception:
                    data = {}
            data['selected_gu'] = selected
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
        except Exception:
            pass

    def save_excel_password(self):
        pwd = self.e_file_password.get().strip()
        try:
            path = self.get_config_path()
            data = {}
            if path.exists():
                try:
                    data = json.loads(path.read_text(encoding='utf-8'))
                except Exception:
                    data = {}
            data['excel_password'] = pwd
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            messagebox.showinfo('안내', '엑셀 비밀번호가 저장되었습니다.')
        except Exception as e:
            messagebox.showerror('오류', f'비밀번호 저장 중 오류가 발생했습니다.\n{e}')

    def clear_saved_password(self):
        try:
            path = self.get_config_path()
            data = {}
            if path.exists():
                try:
                    data = json.loads(path.read_text(encoding='utf-8'))
                except Exception:
                    data = {}
            data.pop('excel_password', None)
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            self.e_file_password.delete(0, 'end')
            messagebox.showinfo('안내', '저장된 엑셀 비밀번호가 초기화되었습니다.')
        except Exception as e:
            messagebox.showerror('오류', f'비밀번호 초기화 중 오류가 발생했습니다.\n{e}')


    def load_saved_asset_date(self):
        try:
            path = self.get_config_path()
            if path.exists():
                data = json.loads(path.read_text(encoding='utf-8'))
                return data.get('asset_date', '')
        except Exception:
            pass
        return ''

    def save_asset_date(self):
        val = self.e_asset_date.get().strip()
        try:
            path = self.get_config_path()
            data = {}
            if path.exists():
                try:
                    data = json.loads(path.read_text(encoding='utf-8'))
                except Exception:
                    data = {}
            data['asset_date'] = val
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            messagebox.showinfo('안내', '금융재산 조회기준일이 저장되었습니다.')
        except Exception as e:
            messagebox.showerror('오류', f'조회기준일 저장 중 오류가 발생했습니다.\n{e}')

    def clear_saved_asset_date(self):
        try:
            path = self.get_config_path()
            data = {}
            if path.exists():
                try:
                    data = json.loads(path.read_text(encoding='utf-8'))
                except Exception:
                    data = {}
            data.pop('asset_date', None)
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            self.e_asset_date.delete(0, 'end')
            messagebox.showinfo('안내', '저장된 금융재산 조회기준일이 초기화되었습니다.')
        except Exception as e:
            messagebox.showerror('오류', f'조회기준일 초기화 중 오류가 발생했습니다.\n{e}')


    def resource_path(self, filename):
        try:
            base = Path(sys._MEIPASS)
        except Exception:
            base = Path(__file__).resolve().parent
        return base / filename

    def open_selection_standard_pdf(self):
        path = self.resource_path('selection_standard_2026.pdf')
        if not path.exists():
            messagebox.showerror('오류', f'붙임파일을 찾을 수 없습니다.\n{path}')
            return
        try:
            if sys.platform.startswith('win'):
                os.startfile(str(path))
            elif sys.platform == 'darwin':
                subprocess.Popen(['open', str(path)])
            else:
                subprocess.Popen(['xdg-open', str(path)])
        except Exception as e:
            messagebox.showerror('오류', f'붙임파일 열기 중 오류가 발생했습니다.\n{e}')


    def __init__(self, root):
        self.root = root
        self.return_to_launcher = False
        self.root.title('통합조사 상담 프로그램 v1 - 2팀(차상위, 기초연금 등)')
        self.root.geometry('1720x940')
        self.root.minsize(1360, 720)
        self.app_font_size = 10
        apply_pastel_theme(self.root, self.app_font_size)
        self.setup_fonts(self.app_font_size)
        self.current_df = None
        self.file_rows = []
        self.selected_beneficiary = []
        self.last_snapshot = None
        self.multi_income_popup_open = False
        self.multi_income_values = {}
        self._last_service_snapshot = tuple()
        self._last_valid_services = tuple()
        self.current_services = []
        self._service_selection_alerting = False
        self.excel_entries = []
        self.generated_excel_password = ''
        self._live_preview_job = None
        self._autosave_job = None
        self._last_preview_text = ''
        self._restoring_state = False
        # 엑셀 반영 저장목록과 생성 엑셀 비밀번호는 입력내용 초기화로 지우지 않습니다.
        # 저장목록은 우측 '선택 삭제'/'전체 삭제'로만 삭제합니다.

        # 1팀과 같은 좌우 분할 구조: 왼쪽은 입력, 오른쪽은 상담내역입니다.
        main_pane = ttk.PanedWindow(root, orient='horizontal')
        main_pane.pack(fill='both', expand=True)
        left_container = ttk.Frame(main_pane, style='Window.TFrame')
        right_container = ttk.Frame(main_pane, style='Window.TFrame')
        main_pane.add(left_container, weight=1)
        main_pane.add(right_container, weight=1)

        def _stabilize_sash():
            try:
                # 기초연금 선택 시 오른쪽의 상담내역과 저장목록이 항상 함께 보이도록
                # 신청정보 입력영역의 기본 폭을 제한합니다.
                target = min(840, max(760, int(self.root.winfo_width() * 0.49)))
                main_pane.sashpos(0, target)
            except tk.TclError:
                pass
        for delay in (80, 300, 800):
            self.root.after(delay, _stabilize_sash)

        left_canvas = tk.Canvas(left_container, highlightthickness=0, background=PALETTE['window'])
        left_scroll = ttk.Scrollbar(left_container, orient='vertical', command=left_canvas.yview)
        left_canvas.configure(yscrollcommand=left_scroll.set)
        left_canvas.pack(side='left', fill='both', expand=True)
        left_scroll.pack(side='right', fill='y')
        self.input_root = ttk.Frame(left_canvas, style='Window.TFrame', padding=(4, 4))
        left_window = left_canvas.create_window((0, 0), window=self.input_root, anchor='nw')

        def _sync_left_scroll(event=None):
            left_canvas.configure(scrollregion=left_canvas.bbox('all'))
            left_canvas.itemconfigure(left_window, width=left_canvas.winfo_width())
        self.input_root.bind('<Configure>', _sync_left_scroll)
        left_canvas.bind('<Configure>', _sync_left_scroll)

        def _is_descendant(widget, ancestor):
            current = widget
            while current is not None:
                if current == ancestor:
                    return True
                try:
                    parent_name = current.winfo_parent()
                    current = current._nametowidget(parent_name) if parent_name else None
                except Exception:
                    return False
            return False

        def _on_mousewheel(event):
            if _is_descendant(event.widget, left_container):
                left_canvas.yview_scroll((-1 if event.delta > 0 else 1) * 3, 'units')
                return 'break'
            return None
        left_canvas.bind_all('<MouseWheel>', _on_mousewheel, add='+')

        nav_frame = ttk.Frame(self.input_root, padding=(12, 10, 12, 4), style='Window.TFrame')
        nav_frame.pack(fill='x')
        ttk.Button(nav_frame, text='← 뒤로가기', command=self.go_back_to_launcher).pack(side='left')
        self.btn_selection_standard = tk.Button(
            nav_frame, text='복지대상자 선정기준표 보기', command=self.open_selection_standard_pdf,
            fg='#b00000', font=('Malgun Gothic', self.app_font_size, 'bold')
        )
        self.btn_selection_standard.pack(side='left', padx=(8, 0))

        top = ttk.LabelFrame(self.input_root, text='신청 정보', padding=12, style='Card.TLabelframe')
        top.pack(fill='x', padx=12, pady=(4, 8))
        ttk.Label(top, text='구 선택').grid(row=0, column=0, sticky='w', padx=4, pady=4)
        self.gu_var = tk.StringVar(value=self.load_saved_gu())
        self.cb_gu = ttk.Combobox(top, state='readonly', width=14, values=['분당구', '중원구', '수정구'], textvariable=self.gu_var)
        self.cb_gu.grid(row=0, column=1, sticky='w', padx=4, pady=4)
        self.cb_gu.bind('<<ComboboxSelected>>', self.save_selected_gu)

        ttk.Label(top, text='신청일자').grid(row=1, column=0, sticky='w', padx=4, pady=4)
        self.e_date = ttk.Entry(top, width=16, style='Normal.TEntry')
        self.e_date.grid(row=1, column=1, sticky='w', padx=4, pady=4)
        self.e_date.bind('<FocusOut>', self.validate_date)
        tk.Label(top, text='자유 입력', fg='gray', font=('Malgun Gothic', 9, 'bold')).grid(row=2, column=1, sticky='w', padx=4)

        ttk.Label(top, text='보장구분').grid(row=0, column=2, sticky='w', padx=4, pady=4)
        self.multi_var = tk.BooleanVar(value=False)
        self.service_var = tk.StringVar()
        self.cb_service = ttk.Combobox(top, state='readonly', width=20, values=SERVICE_OPTIONS, textvariable=self.service_var)
        self.cb_service.grid(row=0, column=3, sticky='w', padx=4, pady=4)
        self.cb_service.bind('<<ComboboxSelected>>', self.on_service_change)
        self.service_var.trace_add('write', lambda *args: self.root.after(10, self.on_service_change))
        self.lb_service = tk.Listbox(top, selectmode='multiple', exportselection=False, height=5, width=22)
        for item in SERVICE_OPTIONS:
            self.lb_service.insert('end', item)

        # 기본 다중선택 방식은 유지하되,
        # 클릭한 항목의 선택 상태를 클릭 후 다시 확인해서 선택 누락을 막습니다.
        self.lb_service.bind('<Button-1>', self.remember_service_click, add='+')
        self.lb_service.bind('<<ListboxSelect>>', self.schedule_service_refresh)
        self.lb_service.bind('<ButtonRelease-1>', self.finish_service_click)
        self.lb_service.bind('<KeyRelease>', self.schedule_service_refresh)
        opt = ttk.Frame(top)
        opt.grid(row=1, column=3, sticky='w')
        ttk.Checkbutton(opt, text='여러보장 선택', variable=self.multi_var, command=self.toggle_multi).pack(side='left', padx=(0, 10))
        self.beneficiary_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(opt, text='수급자 책정 여부', variable=self.beneficiary_var, command=self.on_beneficiary).pack(side='left')

        ttk.Label(top, text='가구유형').grid(row=0, column=4, sticky='w', padx=4, pady=4)
        self.cb_house = ttk.Combobox(top, state='disabled', width=18)
        self.cb_house.grid(row=0, column=5, sticky='w', padx=4, pady=4)
        self.cb_house.bind('<<ComboboxSelected>>', self.on_house_change)
        self.e_house_direct = ttk.Entry(top, width=18, state='disabled')
        self.e_house_direct.grid(row=1, column=5, sticky='w', padx=4)
        # 부양의무자 입력 영역: 주거유형 위쪽에 별도 줄로 표시합니다.
        self.supporter_outer = ttk.LabelFrame(self.input_root, text='부양의무자', padding=8, style='Peach.TLabelframe')
        self.supporter_wrap = ttk.Frame(self.supporter_outer)
        self.supporter_wrap.pack(anchor='w')

        self.supporter_header = ttk.Frame(self.supporter_wrap)
        ttk.Label(self.supporter_header, text='', width=3, anchor='center').grid(row=0, column=0, padx=(0,4), pady=1)
        ttk.Label(self.supporter_header, text='수급자와의 관계', width=14, anchor='center').grid(row=0, column=1, padx=2, pady=1)
        ttk.Label(self.supporter_header, text='성명', width=10, anchor='center').grid(row=0, column=2, padx=2, pady=1)
        ttk.Label(self.supporter_header, text='생년월일', width=10, anchor='center').grid(row=0, column=3, padx=2, pady=1)
        ttk.Label(self.supporter_header, text='소득인정액', width=13, anchor='center').grid(row=0, column=4, padx=2, pady=1)
        ttk.Label(self.supporter_header, text='가구원 수', width=8, anchor='center').grid(row=0, column=5, padx=2, pady=1)

        self.supporter_rows = []
        self.btn_add_supporter = None

        # 타법의료급여(유공자) 선택 시 일반가구/취약가구를 선택하는 영역입니다.
        # 평소에는 숨기고, 타법의료급여(유공자)를 선택했을 때만 표시합니다.
        self.tabeop_type_var = tk.StringVar(value='일반가구')
        self.tabeop_outer = ttk.LabelFrame(self.input_root, text='타법의료급여(유공자) 가구유형', padding=8, style='Mint.TLabelframe')
        self.tabeop_general_var = tk.BooleanVar(value=True)
        self.tabeop_vulnerable_var = tk.BooleanVar(value=False)
        self.tabeop_general_chk = ttk.Checkbutton(self.tabeop_outer, text='일반가구(중위 80%)', variable=self.tabeop_general_var, command=lambda: self.set_tabeop_type('일반가구'))
        self.tabeop_general_chk.pack(side='left', padx=8)
        self.tabeop_vulnerable_chk = ttk.Checkbutton(self.tabeop_outer, text='취약가구(중위 100%)', variable=self.tabeop_vulnerable_var, command=lambda: self.set_tabeop_type('취약가구'))
        self.tabeop_vulnerable_chk.pack(side='left', padx=8)
        self.tabeop_hint = ttk.Label(self.tabeop_outer, text='※ 타법의료급여(유공자) 기준 선택값이 조사자의견에 반영됩니다.', foreground='gray')
        self.tabeop_hint.pack(side='left', padx=12)
        self.tabeop_type_var.trace_add('write', lambda *args: self.root.after(10, self.update_tabeop_visibility))
        # 처음 실행 시에는 화면에 표시하지 않습니다.

        # 청년내일저축계좌 선택 시 표시되는 추가 입력 영역입니다.
        self.youth_work_outer = ttk.LabelFrame(self.input_root, text='청년내일저축계좌 추가정보', padding=8, style='Blue.TLabelframe')
        self.youth_not_working_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(self.youth_work_outer, text='현재 미근로여부', variable=self.youth_not_working_var).pack(side='left', padx=8)
        ttk.Label(self.youth_work_outer, text='※ 체크 시 조사자의견에 미근로 부적합 문구가 반영됩니다.', foreground='gray').pack(side='left', padx=8)

        # 장애인연금/차상위장애인 선택 시 표시되는 장애심사 정보 입력 영역입니다.
        self.disability_outer = ttk.LabelFrame(self.input_root, text='장애심사 정보', padding=8, style='Lavender.TLabelframe')
        ttk.Label(self.disability_outer, text='공단심사번호').pack(side='left', padx=(4, 2))
        self.e_disability_review_no = ttk.Entry(self.disability_outer, width=24)
        self.e_disability_review_no.pack(side='left', padx=(0, 12))
        ttk.Label(self.disability_outer, text='장애정도 결정일').pack(side='left', padx=(4, 2))
        self.e_disability_decision_date = ttk.Entry(self.disability_outer, width=18)
        self.e_disability_decision_date.pack(side='left', padx=(0, 12))
        ttk.Label(self.disability_outer, text='장애정도 및 유형').pack(side='left', padx=(4, 2))
        self.e_disability_type = ttk.Entry(self.disability_outer, width=22)
        self.e_disability_type.pack(side='left', padx=(0, 8))

        ttk.Label(top, text='주거유형').grid(row=2, column=0, sticky='w', padx=4, pady=4)
        self.cb_home = ttk.Combobox(top, state='readonly', width=28, values=HOME_OPTIONS)
        self.cb_home.grid(row=2, column=1, sticky='w', padx=4, pady=4)
        self.cb_home.bind('<<ComboboxSelected>>', self.on_home_change)

        ttk.Label(top, text='소득인정액').grid(row=2, column=2, sticky='w', padx=4, pady=4)
        self.e_income = ttk.Entry(top, width=16)
        self.e_income.grid(row=2, column=3, sticky='w', padx=4, pady=4)
        self.e_income.bind('<FocusOut>', self.format_entry_money)
        self.e_income.bind('<Button-1>', self.open_multi_income_popup)
        self.ref_var = tk.BooleanVar(value=False)

        self.multi_income_frame = ttk.LabelFrame(self.input_root, text='보장별 소득인정액', padding=8, style='Blue.TLabelframe')
        self.multi_income_entries = {}

        self.hf = ttk.Frame(self.input_root, padding=(12, 0, 12, 8))
        self.hf.pack(fill='x')
        ttk.Label(self.hf, text='주거유형 직접입력').grid(row=0, column=0, padx=4, pady=4)
        self.e_home_direct = ttk.Entry(self.hf, width=24, state='disabled')
        self.e_home_direct.grid(row=0, column=1, columnspan=3, padx=4, pady=4, sticky='w')
        ttk.Label(self.hf, text='보증금(선택사항)').grid(row=1, column=0, padx=4, pady=4)
        self.e_deposit = ttk.Entry(self.hf, width=16, state='disabled')
        self.e_deposit.grid(row=1, column=1, padx=4, pady=4)
        self.e_deposit.bind('<FocusOut>', self.format_entry_money)
        ttk.Label(self.hf, text='임대료(선택사항)').grid(row=1, column=2, padx=4, pady=4)
        self.e_rent = ttk.Entry(self.hf, width=16, state='disabled')
        self.e_rent.grid(row=1, column=3, padx=4, pady=4)
        self.e_rent.bind('<FocusOut>', self.format_entry_money)

        fw = ttk.LabelFrame(self.input_root, text='가구원별 엑셀 파일', padding=12, style='Lavender.TLabelframe')
        fw.pack(fill='x', padx=12, pady=(0, 8))
        pwf = ttk.Frame(fw)
        pwf.pack(fill='x', pady=(0, 6))
        ttk.Label(pwf, text='엑셀 비밀번호').pack(side='left', padx=4)
        self.e_file_password = ttk.Entry(pwf, width=20, show='*')
        self.e_file_password.pack(side='left', padx=4)
        saved_pwd = self.load_saved_password()
        if saved_pwd:
            self.e_file_password.insert(0, saved_pwd)
        ttk.Button(pwf, text='비밀번호 저장', command=self.save_excel_password).pack(side='left', padx=4)
        ttk.Button(pwf, text='비밀번호 초기화', command=self.clear_saved_password).pack(side='left', padx=4)
        self.file_area = ttk.Frame(fw)
        self.file_area.pack(fill='x')
        fb = ttk.Frame(fw)
        fb.pack(fill='x', pady=(8, 0))
        ttk.Button(fb, text='가구원 추가', command=self.add_file_row).pack(side='left', padx=4)
        tk.Button(fb, text='첨부 파일 반영하기', command=self.load_all_excels, fg='black', font=('Malgun Gothic', self.app_font_size, 'bold')).pack(side='left', padx=4)

        asset_df = ttk.Frame(self.input_root, padding=(12, 0, 12, 4))
        asset_df.pack(fill='x')
        ttk.Label(asset_df, text='금융재산 조회기준일(선택사항)').pack(side='left', padx=4)
        self.e_asset_date = ttk.Entry(asset_df, width=18)
        self.e_asset_date.pack(side='left', padx=4)
        saved_asset_date = self.load_saved_asset_date()
        if saved_asset_date:
            self.e_asset_date.insert(0, saved_asset_date)
        ttk.Button(asset_df, text='조회기준일 저장', command=self.save_asset_date).pack(side='left', padx=4)
        ttk.Button(asset_df, text='조회기준일 초기화', command=self.clear_saved_asset_date).pack(side='left', padx=4)

        bf = ttk.Frame(self.input_root, padding=12)
        bf.pack(fill='x')
        self.action_frame = bf
        ttk.Button(bf, text='상담내역 생성', command=self.generate_preview).pack(side='left', padx=4)
        ttk.Button(bf, text='상담내역 저장', command=self.copy_clipboard).pack(side='left', padx=4)
        tk.Button(
            bf, text='입력내용 초기화', command=self.confirm_reset,
            fg='#b00000', font=('Malgun Gothic', self.app_font_size, 'bold'),
            bd=0, relief='flat', highlightthickness=1, padx=12, pady=7
        ).pack(side='left', padx=4)
        ttk.Button(bf, text='초기화 되돌리기', command=self.undo_reset).pack(side='left', padx=4)
        ttk.Button(bf, text='최근 입력 10건', command=self.show_recent_inputs).pack(side='left', padx=4)
        ttk.Label(bf, text='글자 크기').pack(side='left', padx=(18, 4))
        self.font_size_var = tk.StringVar(value='10')
        self.cb_font_size = ttk.Combobox(bf, state='readonly', width=5, values=['9', '10', '11', '12', '13', '14', '16', '18'], textvariable=self.font_size_var)
        self.cb_font_size.pack(side='left', padx=4)
        self.cb_font_size.bind('<<ComboboxSelected>>', self.change_text_font_size)

        ref_frame = ttk.Frame(self.input_root, padding=(12, 0, 12, 8))
        ref_frame.pack(fill='x')
        ttk.Label(ref_frame, text='참고자료 미조사').pack(side='left', padx=(4, 2))
        ttk.Checkbutton(ref_frame, variable=self.ref_var).pack(side='left', padx=(0, 8))
        ttk.Label(ref_frame, text='* 현재 소득인정액이 기준을 초과하여 참고자료 조사가 필요없을 경우',
                  foreground='gray').pack(side='left', padx=4)

        shortcut_row = ttk.Frame(self.input_root, padding=(12, 0, 12, 10))
        shortcut_row.pack(fill='x')
        shortcut_card = tk.Frame(shortcut_row, bd=1, relief='solid', padx=8, pady=4, background=PALETTE['blue'], highlightbackground=PALETTE['border'])
        shortcut_card.pack(side='left', padx=4)
        tk.Label(shortcut_card, text='⌨ 빠른 단축키', background=PALETTE['blue'], foreground=PALETTE['primary'],
                 font=('Malgun Gothic', 9, 'bold')).grid(row=0, column=0, columnspan=5, sticky='w', pady=(0, 2))
        shortcuts = [('Ctrl+Enter', '생성'), ('Ctrl+S', '저장'), ('Ctrl+R', '초기화'), ('Ctrl+Z', '되돌리기'), ('F1', '안내')]
        for col, (key, action) in enumerate(shortcuts):
            cell = tk.Frame(shortcut_card, background=PALETTE['blue'])
            cell.grid(row=1, column=col, padx=(0 if col == 0 else 7, 0), sticky='w')
            tk.Label(cell, text=key, bd=1, relief='ridge', padx=4, pady=1,
                     background=PALETTE['entry'], foreground=PALETTE['primary'], font=('Malgun Gothic', 8, 'bold')).pack(side='left')
            tk.Label(cell, text=action, background=PALETTE['blue'], foreground=PALETTE['text'],
                     font=('Malgun Gothic', 8)).pack(side='left', padx=(3, 0))

        self.excel_bar = ttk.Frame(self.input_root, padding=(12, 0, 12, 8))
        ttk.Button(self.excel_bar, text='엑셀에 반영', command=self.open_excel_entry_popup).pack(side='left', padx=4)
        ttk.Button(self.excel_bar, text='엑셀 파일 만들기', command=self.create_excel_file).pack(side='left', padx=(18, 4))
        ttk.Label(self.excel_bar, text='생성 엑셀 열기 비밀번호').pack(side='left', padx=(12, 4))
        self.e_generated_excel_password = ttk.Entry(self.excel_bar, width=20, show='*')
        self.e_generated_excel_password.pack(side='left', padx=4)
        ttk.Button(self.excel_bar, text='비밀번호 저장', command=self.save_generated_excel_password).pack(side='left', padx=4)
        ttk.Button(self.excel_bar, text='비밀번호 초기화', command=self.clear_generated_excel_password).pack(side='left', padx=4)

        body = ttk.Frame(right_container, padding=(8, 12, 12, 12), style='Window.TFrame')
        body.pack(fill='both', expand=True)
        text_wrap = ttk.Frame(body, style='Card.TFrame')
        text_wrap.pack(side='left', fill='both', expand=True)
        ttk.Label(
            text_wrap,
            text='상담내역 미리보기',
            font=('Malgun Gothic', self.app_font_size + 1, 'bold'),
            foreground=PALETTE['primary'],
            background=PALETTE['surface'],
        ).pack(side='top', fill='x', padx=12, pady=(10, 6))
        text_inner = ttk.Frame(text_wrap, style='Card.TFrame')
        text_inner.pack(side='top', fill='both', expand=True, padx=10, pady=(0, 10))
        self.txt = tk.Text(text_inner, wrap='word', font=('Malgun Gothic', int(self.font_size_var.get())),
                           background=PALETTE['entry'], foreground=PALETTE['text'], insertbackground=PALETTE['text'],
                           selectbackground=PALETTE['blue_active'], relief='solid', bd=1)
        text_scroll = ttk.Scrollbar(text_inner, orient='vertical', command=self.txt.yview)
        self.txt.configure(yscrollcommand=text_scroll.set)
        self.txt.pack(side='left', fill='both', expand=True)
        text_scroll.pack(side='right', fill='y')
        self.char_count_var = tk.StringVar(value='0자')
        ttk.Label(body, textvariable=self.char_count_var, anchor='e').pack(side='bottom', fill='x', pady=(3, 0))
        self.txt.tag_configure('changed_line', background='#fff3a3')
        self.right_panel = ttk.LabelFrame(body, text='미리보기 및 저장목록', padding=8, style='Mint.TLabelframe', width=238)
        self.right_panel.grid_propagate(False)
        # 상단 정보와 하단 삭제 버튼을 고정하고, 목록만 남는 공간에 맞춰 늘어나도록 배치합니다.
        self.right_panel.columnconfigure(0, weight=1)
        self.right_panel.rowconfigure(5, weight=1)
        ttk.Label(self.right_panel, text='신청 가구원 미리보기', font=('Malgun Gothic', self.app_font_size, 'bold')).grid(row=0, column=0, sticky='w')
        self.applicant_preview_var = tk.StringVar(value='첨부 파일 반영 후 표시됩니다.')
        ttk.Label(self.right_panel, textvariable=self.applicant_preview_var, justify='left', wraplength=210).grid(row=1, column=0, sticky='ew', pady=(2, 10))
        ttk.Separator(self.right_panel).grid(row=2, column=0, sticky='ew', pady=(0, 8))
        ttk.Label(self.right_panel, text='엑셀 반영 저장목록', font=('Malgun Gothic', self.app_font_size, 'bold')).grid(row=3, column=0, sticky='w')
        self.excel_entry_listbox = tk.Listbox(self.right_panel, width=24, height=8, exportselection=False)
        self.excel_entry_listbox.grid(row=5, column=0, sticky='nsew', pady=(4, 6))
        delete_buttons = ttk.Frame(self.right_panel)
        delete_buttons.grid(row=6, column=0, sticky='ew')
        delete_buttons.columnconfigure(0, weight=1)
        delete_buttons.columnconfigure(1, weight=1)
        ttk.Button(delete_buttons, text='선택 삭제', command=self.delete_selected_excel_entry).grid(row=0, column=0, sticky='ew', padx=(0, 3))
        ttk.Button(delete_buttons, text='전체 삭제', command=self.clear_excel_entries).grid(row=0, column=1, sticky='ew', padx=(3, 0))

        self.add_file_row()
        self.root.after(120, self.update_excel_visibility)
        self.root.after(300, self.watch_service_selection)
        self.bind_productivity_features()
        self.restore_autosave_state()
        self.schedule_autosave()
        self.schedule_live_preview()
        self.root.protocol('WM_DELETE_WINDOW', self.on_close)

        def _fit_initial_window():
            try:
                self.root.update_idletasks()
                width = min(1720, max(1360, self.root.winfo_screenwidth() - 60))
                needed = self.input_root.winfo_reqheight() + 45
                height = min(max(needed, 700), self.root.winfo_screenheight() - 90)
                self.root.geometry(f'{width}x{height}')
                _stabilize_sash()
            except tk.TclError:
                pass
        self.root.after(350, _fit_initial_window)
        style_plain_widgets(self.root)


    def get_runtime_data_path(self, filename):
        base = self.get_config_path().parent
        base.mkdir(parents=True, exist_ok=True)
        return base / filename

    def bind_productivity_features(self):
        self.root.bind('<Control-Return>', lambda e: self.generate_preview())
        self.root.bind('<Control-s>', lambda e: self.copy_clipboard())
        self.root.bind('<Control-r>', lambda e: self.confirm_reset())
        self.root.bind('<Control-z>', lambda e: self.undo_reset())
        self.root.bind('<F1>', lambda e: self.show_shortcut_guide())
        for w in [self.e_date, self.cb_service, self.cb_house, self.e_house_direct, self.cb_home,
                  self.e_income, self.e_deposit, self.e_rent,
                  self.e_asset_date, self.e_disability_review_no, self.e_disability_decision_date,
                  self.e_disability_type]:
            try:
                w.bind('<KeyRelease>', lambda e: self.schedule_live_preview(), add='+')
                w.bind('<<ComboboxSelected>>', lambda e: self.schedule_live_preview(), add='+')
            except Exception:
                pass
        for v in [self.ref_var, self.beneficiary_var, self.youth_not_working_var,
                  self.tabeop_general_var, self.tabeop_vulnerable_var]:
            try:
                v.trace_add('write', lambda *args: self.schedule_live_preview())
            except Exception:
                pass

    def show_shortcut_guide(self):
        messagebox.showinfo('단축키 안내', 'Ctrl+Enter  상담내역 생성\nCtrl+S  상담내역 저장(클립보드)\nCtrl+R  입력내용 초기화\nCtrl+Z  초기화 되돌리기\nF1  단축키 안내')

    def confirm_reset(self):
        if messagebox.askyesno('입력내용 초기화', '현재 입력내용을 초기화할까요?\n초기화 후에는 되돌리기 버튼으로 한 번 복원할 수 있습니다.'):
            self.reset_all()

    def schedule_live_preview(self, delay=350):
        if self._restoring_state:
            return
        try:
            if self._live_preview_job:
                self.root.after_cancel(self._live_preview_job)
        except Exception:
            pass
        self._live_preview_job = self.root.after(delay, self.live_preview)

    def live_preview(self):
        self._live_preview_job = None
        self.update_required_highlights()
        self.update_applicant_preview()
        if self.current_df is None or getattr(self.current_df, 'empty', True):
            self.update_char_count()
            return
        if not self.get_selected_services():
            self.update_char_count()
            return
        try:
            new_text = self.build_text()
        except Exception:
            return
        old_text = self.txt.get('1.0', 'end-1c')
        if new_text != old_text:
            self.txt.delete('1.0', 'end')
            self.txt.insert('1.0', new_text)
            self.highlight_changed_lines(old_text, new_text)
        self._last_preview_text = new_text
        self.update_char_count()

    def highlight_changed_lines(self, old_text, new_text):
        old = old_text.splitlines()
        new = new_text.splitlines()
        changed = [i for i, line in enumerate(new, start=1) if i-1 >= len(old) or old[i-1] != line]
        self.txt.tag_remove('changed_line', '1.0', 'end')
        for i in changed[:80]:
            self.txt.tag_add('changed_line', f'{i}.0', f'{i}.end')
        if changed:
            self.root.after(1800, lambda: self.txt.tag_remove('changed_line', '1.0', 'end'))

    def update_char_count(self):
        try:
            self.char_count_var.set(f"{len(self.txt.get('1.0', 'end-1c')):,}자")
        except Exception:
            pass

    def update_required_highlights(self):
        try:
            style = ttk.Style()
            style.configure('Required.TEntry', fieldbackground=PALETTE['rose'])
            style.configure('Required.TCombobox', fieldbackground=PALETTE['rose'])
        except Exception:
            pass
        checks = [
            (self.e_date, bool(self.e_date.get().strip()), 'Normal.TEntry', 'Required.TEntry'),
            (self.cb_service, bool(self.get_selected_services()), 'TCombobox', 'Required.TCombobox'),
            (self.cb_house, bool(self.household_display().strip()), 'TCombobox', 'Required.TCombobox'),
            (self.cb_home, bool(self.cb_home.get().strip()), 'TCombobox', 'Required.TCombobox'),
        ]
        for w, ok, normal, bad in checks:
            try:
                w.configure(style=normal if ok else bad)
            except Exception:
                pass

    def update_applicant_preview(self):
        try:
            if self.current_df is None or self.current_df.empty:
                self.applicant_preview_var.set('첨부 파일 반영 후 표시됩니다.')
                return
            rows = self.current_df[['성명', '주민등록번호']].drop_duplicates()
            vals = []
            for _, r in rows.head(8).iterrows():
                vals.append(f"{self.clean(r['성명'])}  {self.mask_rrn(r['주민등록번호'])}")
            if len(rows) > 8:
                vals.append(f'외 {len(rows)-8}명')
            self.applicant_preview_var.set('\n'.join(vals) if vals else '표시할 가구원이 없습니다.')
        except Exception:
            pass

    def autosave_payload(self):
        return {
            'date': self.e_date.get(), 'gu': self.cb_gu.get(), 'multi': self.multi_var.get(),
            'single_service': self.cb_service.get(), 'multi_sel': list(self.lb_service.curselection()),
            'house': self.cb_house.get(), 'house_direct': self.e_house_direct.get(), 'home': self.cb_home.get(),
            'home_direct': self.e_home_direct.get(),
            'income': self.e_income.get(), 'ref': self.ref_var.get(), 'beneficiary': self.beneficiary_var.get(),
            'deposit': self.e_deposit.get(), 'rent': self.e_rent.get(),
            'asset_date': self.e_asset_date.get(), 'file_paths': [r.path for r in self.file_rows if r.path],
            'text': self.txt.get('1.0', 'end-1c'), 'saved_at': datetime.now().isoformat(timespec='seconds')
        }

    def schedule_autosave(self):
        try:
            if self._autosave_job:
                self.root.after_cancel(self._autosave_job)
        except Exception:
            pass
        self._autosave_job = self.root.after(1500, self.autosave_state)

    def autosave_state(self):
        self._autosave_job = None
        try:
            path = self.get_runtime_data_path('team2_autosave.json')
            tmp = path.with_suffix('.tmp')
            tmp.write_text(json.dumps(self.autosave_payload(), ensure_ascii=False, indent=2), encoding='utf-8')
            tmp.replace(path)
        except Exception:
            pass
        self.schedule_autosave()

    def restore_autosave_state(self):
        path = self.get_runtime_data_path('team2_autosave.json')
        if not path.exists():
            return
        try:
            data = json.loads(path.read_text(encoding='utf-8'))
        except Exception:
            return
        if not any([data.get('date'), data.get('single_service'), data.get('multi_sel'), data.get('text'), data.get('file_paths')]):
            return
        if not messagebox.askyesno('자동저장 복원', f"이전 입력내용을 복원할까요?\n저장시각: {data.get('saved_at', '')}"):
            return
        self.restore_recent_item(data)

    def recent_path(self):
        return self.get_runtime_data_path('team2_recent_inputs.json')

    def save_recent_input(self):
        try:
            item = self.autosave_payload()
            item['label'] = self.recent_label()
            path = self.recent_path()
            items = []
            if path.exists():
                try:
                    items = json.loads(path.read_text(encoding='utf-8'))
                except Exception:
                    items = []
            items = [x for x in items if x.get('label') != item['label']]
            items.insert(0, item)
            items = items[:10]
            tmp = path.with_suffix('.tmp')
            tmp.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding='utf-8')
            tmp.replace(path)
        except Exception as e:
            messagebox.showwarning('최근 입력 저장 오류', f'최근 입력내역을 저장하지 못했습니다.\n{e}')

    def recent_label(self):
        names = []
        try:
            if self.current_df is not None and not self.current_df.empty:
                names = [self.clean(x) for x in self.current_df['성명'].dropna().unique().tolist() if self.clean(x)]
        except Exception:
            pass
        return f"{datetime.now().strftime('%m/%d %H:%M')} | {', '.join(names[:2]) or '성명 없음'} | {self.format_service_line() or '보장 미선택'}"

    def show_recent_inputs(self):
        path = self.recent_path()
        try:
            items = json.loads(path.read_text(encoding='utf-8')) if path.exists() else []
        except Exception:
            items = []
        if not items:
            messagebox.showinfo('최근 입력 10건', '저장된 최근 입력내역이 없습니다.\n상담내역 생성 또는 저장 후 다시 확인하세요.')
            return
        win = tk.Toplevel(self.root)
        win.title('최근 입력 10건')
        win.geometry('620x340')
        win.transient(self.root)
        lb = tk.Listbox(win, exportselection=False)
        lb.pack(fill='both', expand=True, padx=10, pady=10)
        for x in items:
            lb.insert('end', x.get('label', '최근 입력'))
        def restore():
            sel = lb.curselection()
            if not sel:
                return
            self.restore_recent_item(items[sel[0]])
            win.destroy()
        ttk.Button(win, text='선택 복원', command=restore).pack(pady=(0, 10))
        lb.bind('<Double-Button-1>', lambda e: restore())

    def restore_recent_item(self, data):
        self._restoring_state = True
        try:
            self.e_date.delete(0, 'end'); self.e_date.insert(0, data.get('date', ''))
            self.cb_gu.set(data.get('gu', '분당구'))
            self.multi_var.set(bool(data.get('multi'))); self.toggle_multi()
            if data.get('multi'):
                self.lb_service.selection_clear(0, 'end')
                for i in data.get('multi_sel', []):
                    try: self.lb_service.selection_set(int(i))
                    except Exception: pass
            else:
                self.cb_service.set(data.get('single_service', ''))
            self.on_service_change()
            self.cb_house.set(data.get('house', '')); self.on_house_change()
            self.e_house_direct.config(state='normal'); self.e_house_direct.delete(0, 'end'); self.e_house_direct.insert(0, data.get('house_direct', ''))
            if self.cb_house.get().strip() != '직접입력':
                self.e_house_direct.config(state='disabled')
            old_home = data.get('home', '')
            restored_home = OLD_HOME_OPTION_MAP.get(old_home, old_home)
            if restored_home not in HOME_OPTIONS:
                restored_home = '기타(직접입력)'
            self.cb_home.set(restored_home); self.on_home_change()
            if restored_home == '기타(직접입력)':
                direct_home = data.get('home_direct', '') or (old_home if old_home not in ('', '직접입력', '기타(직접입력)') else '')
                self.e_home_direct.insert(0, direct_home)
            for w, k in [(self.e_income, 'income'), (self.e_asset_date, 'asset_date')]:
                try:
                    w.config(state='normal'); w.delete(0, 'end'); w.insert(0, data.get(k, ''))
                except Exception:
                    pass
            for w, k in [(self.e_deposit, 'deposit'), (self.e_rent, 'rent')]:
                try:
                    if str(w.cget('state')) != 'disabled':
                        w.delete(0, 'end'); w.insert(0, data.get(k, ''))
                except Exception:
                    pass
            self.ref_var.set(bool(data.get('ref')))
            self.beneficiary_var.set(bool(data.get('beneficiary')))
            self.on_beneficiary()
            self.txt.delete('1.0', 'end'); self.txt.insert('1.0', data.get('text', ''))
        finally:
            self._restoring_state = False
        self.update_char_count()
        self.schedule_live_preview(50)

    def setup_fonts(self, size=10):
        self.app_font_size = int(size)
        for name in ['TkDefaultFont', 'TkTextFont', 'TkFixedFont', 'TkMenuFont', 'TkHeadingFont', 'TkCaptionFont', 'TkSmallCaptionFont', 'TkIconFont', 'TkTooltipFont']:
            try:
                tkfont.nametofont(name).configure(family='Malgun Gothic', size=self.app_font_size)
            except Exception:
                pass
        try:
            style = ttk.Style()
            style.configure('.', font=('Malgun Gothic', self.app_font_size))
            style.configure('TButton', font=('Malgun Gothic', self.app_font_size, 'bold'), padding=(10, 4))
            style.configure('TLabel', font=('Malgun Gothic', self.app_font_size))
            style.configure('TCheckbutton', font=('Malgun Gothic', self.app_font_size))
            style.configure('TCombobox', font=('Malgun Gothic', self.app_font_size))
            style.configure('TLabelframe.Label', font=('Malgun Gothic', self.app_font_size))
            # 신청일자 입력칸도 다른 입력칸과 동일한 테두리 두께로 보이도록 명시합니다.
            style.configure('Normal.TEntry', borderwidth=1, relief='solid')
        except Exception:
            pass

    def apply_font_to_children(self, widget, size):
        for child in widget.winfo_children():
            try:
                cls = child.winfo_class()
                if cls in ['Text', 'Listbox', 'Entry', 'Label', 'Button', 'Checkbutton', 'Radiobutton']:
                    child.configure(font=('Malgun Gothic', size))
            except Exception:
                pass
            self.apply_font_to_children(child, size)

    def change_text_font_size(self, event=None):
        try:
            size = int(self.font_size_var.get())
        except Exception:
            size = 10
        self.setup_fonts(size)
        self.apply_font_to_children(self.root, size)
        try:
            self.txt.configure(font=('Malgun Gothic', size))
        except Exception:
            pass

    def normalize_service(self, text):
        return text.replace('(개발예정)', '').strip()

    def validate_date(self, event=None):
        # 신청일자는 자유 입력으로 사용합니다. 형식검사/오류팝업을 표시하지 않습니다.
        return

    def format_date(self):
        # 상담내역에는 사용자가 입력한 신청일자를 그대로 반영합니다.
        return self.e_date.get().strip()

    def format_date_for_excel(self, value=None):
        """
        엑셀 파일 만들기에는 신청일자를 YYYY.MM.DD 형식으로 저장합니다.
        입력칸은 자유 입력을 허용하므로 20260729, 2026-7-29, 2026.7.29,
        7월29일 같은 값을 가능한 범위에서 표준 형식으로 변환합니다.
        """
        raw = (value if value is not None else self.e_date.get()).strip()
        if not raw:
            return ''
        text = raw.strip()

        # 이미 datetime/date 형태로 들어온 경우
        try:
            if hasattr(value, 'year') and hasattr(value, 'month') and hasattr(value, 'day'):
                return f'{value.year:04d}.{value.month:02d}.{value.day:02d}'
        except Exception:
            pass

        # 20260729 / 260729 같은 숫자만 있는 입력
        digits = re.sub(r'\D', '', text)
        candidates = []
        if len(digits) == 8:
            candidates.append((int(digits[:4]), int(digits[4:6]), int(digits[6:8])))
        elif len(digits) == 6:
            yy = int(digits[:2])
            yyyy = 2000 + yy if yy < 70 else 1900 + yy
            candidates.append((yyyy, int(digits[2:4]), int(digits[4:6])))

        # 2026.7.29 / 2026-07-29 / 2026/7/29 / 2026년 7월 29일
        m = re.search(r'(20\d{2}|19\d{2})\D+(\d{1,2})\D+(\d{1,2})', text)
        if m:
            candidates.append((int(m.group(1)), int(m.group(2)), int(m.group(3))))

        # 7월29일처럼 연도가 없는 입력은 현재 연도 기준
        m = re.search(r'(\d{1,2})\s*월\s*(\d{1,2})\s*일?', text)
        if m and not re.search(r'(20\d{2}|19\d{2})', text):
            candidates.append((datetime.now().year, int(m.group(1)), int(m.group(2))))

        for y, mth, d in candidates:
            try:
                dt = datetime(y, mth, d)
                return dt.strftime('%Y.%m.%d')
            except Exception:
                continue

        # 변환할 수 없는 값은 임의 변경하지 않고 그대로 둡니다.
        return raw

    def remember_service_click(self, event=None):
        try:
            idx = self.lb_service.nearest(event.y)
            if idx is None or idx < 0 or idx >= self.lb_service.size():
                self._service_click_info = None
                return 'break'
            was_selected = idx in self.lb_service.curselection()
            self._service_click_info = (idx, was_selected)
        except Exception:
            self._service_click_info = None
        # Listbox 기본 선택 처리와 수동 선택 처리가 겹치면 팝업이 2번 뜨고
        # 선택 상태가 꼬일 수 있어, 아래 finish_service_click에서 직접 처리합니다.
        return 'break'

    def finish_service_click(self, event=None):
        # 여러보장 선택은 직접 토글 처리합니다.
        # 충돌이 생기면 "방금 누른 항목"만 되돌려, 기초연금 단독 선택일 때만 엑셀 반영 기능이 켜지게 합니다.
        def apply_and_refresh():
            try:
                clicked_idx = None
                if self.multi_var.get() and self._service_click_info is not None:
                    idx, was_selected = self._service_click_info
                    clicked_idx = idx
                    if 0 <= idx < self.lb_service.size():
                        if was_selected:
                            self.lb_service.selection_clear(idx)
                        else:
                            self.lb_service.selection_set(idx)
                self._service_click_info = None
                self.resolve_service_conflicts(clicked_idx=clicked_idx)
                self.on_service_change()
                self.update_supporter_visibility()
                self.update_tabeop_visibility()
                self.update_special_visibility()
                self.update_excel_visibility()
            except Exception:
                self._service_click_info = None
                self.on_service_change()
                self.update_excel_visibility()
        self.root.after(20, apply_and_refresh)
        return 'break'

    def schedule_service_refresh(self, event=None):
        # 마우스 클릭 중에는 finish_service_click에서 한 번만 처리합니다.
        # 이중 호출을 막아 같은 팝업이 2개 뜨는 현상을 방지합니다.
        if getattr(self, '_service_click_info', None) is not None:
            return
        self.root.after(200, self.on_service_change)
        self.root.after(230, self.update_supporter_visibility)
        self.root.after(240, self.update_tabeop_visibility)
        self.root.after(250, self.update_special_visibility)
        self.root.after(260, self.update_excel_visibility)

    def update_service_listbox_display(self):
        # 현재 Listbox의 선택 상태를 그대로 유지합니다.
        # 외부에서 복원할 때만 selection_set으로 음영을 다시 표시합니다.
        return

    def toggle_multi(self):
        if self.multi_var.get():
            # 단일 선택값이 있으면 여러보장 목록에도 반영
            single = self.cb_service.get().strip()
            self.cb_service.grid_remove()
            self.lb_service.grid(row=0, column=3, sticky='w', padx=4, pady=4)
            if single:
                target = self.normalize_service(single)
                for i in range(self.lb_service.size()):
                    if self.normalize_service(self.lb_service.get(i)) == target:
                        self.lb_service.selection_set(i)
                        break
        else:
            selected = self.get_selected_services()
            self.lb_service.grid_remove()
            self.cb_service.grid()
            if selected:
                for x in SERVICE_OPTIONS:
                    if self.normalize_service(x) == selected[0]:
                        self.cb_service.set(x)
                        break
            self.lb_service.selection_clear(0, 'end')

        self.sync_current_services_from_widgets()
        self.on_service_change()
        self.root.after(80, self.update_supporter_visibility)
        self.root.after(90, self.update_tabeop_visibility)
        self.root.after(100, self.update_excel_visibility)

    def _read_services_from_widgets(self):
        """현재 화면 위젯에서 선택된 보장을 직접 읽습니다.

        단일 선택(Combobox)과 여러보장 선택(Listbox)의 값을 한곳에서만 읽어
        엑셀 반영 표시 로직이 서로 다른 기준을 보지 않도록 합니다.
        """
        try:
            if self.multi_var.get():
                return [self.normalize_service(self.lb_service.get(i)) for i in self.lb_service.curselection()]
            v = ''
            try:
                v = self.service_var.get().strip()
            except Exception:
                v = ''
            if not v:
                v = self.cb_service.get().strip()
            return [self.normalize_service(v)] if v else []
        except Exception:
            return []

    def sync_current_services_from_widgets(self):
        self.current_services = self._read_services_from_widgets()
        return list(self.current_services)

    def get_selected_services(self):
        # 프로그램 전체가 같은 선택 상태를 보도록 current_services를 기준으로 반환합니다.
        # current_services가 아직 없거나 초기 상태면 위젯에서 직접 읽어 보정합니다.
        if not hasattr(self, 'current_services'):
            self.current_services = []
        if not self.current_services:
            self.sync_current_services_from_widgets()
        return list(self.current_services)

    def watch_service_selection(self):
        try:
            current = tuple(self._read_services_from_widgets())
            if current != self._last_service_snapshot:
                self._last_service_snapshot = current
                self.current_services = list(current)
                self.on_service_change()
            else:
                self.current_services = list(current)
                self.update_supporter_visibility()
                self.update_tabeop_visibility()
                self.update_special_visibility()
            self.update_excel_visibility()
            self.refresh_excel_entry_list()
        finally:
            self.root.after(120, self.update_excel_visibility)
        self.root.after(300, self.watch_service_selection)

    def on_beneficiary(self):
        if self.beneficiary_var.get():
            p = BenefitPopup(self.root, self.selected_beneficiary)
            if p.result is None or not p.result:
                self.beneficiary_var.set(False)
                self.selected_beneficiary = []
            else:
                self.selected_beneficiary = p.result
        else:
            self.selected_beneficiary = []

    def on_service_change(self, event=None):
        # 단일/다중 선택값을 먼저 하나의 상태값으로 확정합니다.
        services = self.sync_current_services_from_widgets()
        if self.resolve_service_conflicts():
            services = self.sync_current_services_from_widgets()

        if not services:
            self.cb_house.set('')
            self.cb_house.config(state='disabled', values=[])
            self.on_house_change()
            self.refresh_multi_income_inputs()
            self.update_supporter_visibility()
            self.update_tabeop_visibility()
            self.update_special_visibility()
            self.update_excel_visibility()
            return

        values = []

        # 기초연금은 단독 선택으로 제한되며, 선택 시 기초연금 가구유형을 표시합니다.
        if '기초연금' in services:
            values.extend(BASIC_PENSION_HOUSEHOLDS)

        non_basic_services = [s for s in services if s != '기초연금']
        if non_basic_services:
            if set(non_basic_services).issubset({'저소득 한부모', '경기도 한부모'}):
                values.extend(SINGLE_PARENT_HOUSEHOLDS)
            else:
                values.extend(GENERAL_HOUSEHOLDS)

        if not values:
            values = GENERAL_HOUSEHOLDS

        values = list(dict.fromkeys(values))

        self.cb_house.config(state='readonly', values=values)
        if self.cb_house.get() not in values:
            self.cb_house.set(values[0])

        self.on_house_change()
        self.refresh_multi_income_inputs()

        # 즉시 1회 + 선택 반영 후 1회 더 확인
        self.update_supporter_visibility()
        self.root.after(80, self.update_supporter_visibility)
        self.root.after(90, self.update_tabeop_visibility)
        self.root.after(100, self.update_special_visibility)
        self.root.after(110, self.update_excel_visibility)

    def update_supporter_visibility(self):
        if self.supporter_enabled():
            self.show_supporter_area()
        else:
            self.clear_supporters()


    def _listbox_selected_services(self):
        return [self.normalize_service(self.lb_service.get(i)) for i in self.lb_service.curselection()]

    def _show_service_conflict_popup(self, message):
        if getattr(self, '_service_selection_alerting', False):
            return
        try:
            self._service_selection_alerting = True
            messagebox.showinfo('안내', message)
        finally:
            self._service_selection_alerting = False

    def resolve_service_conflicts(self, clicked_idx=None):
        """보장 동시선택 제한을 정리합니다.

        핵심 원칙:
        - 기초연금 단독 선택일 때만 엑셀 반영 기능을 활성화합니다.
        - 다른 보장 선택 상태에서 기초연금을 누르면 기초연금 선택만 취소합니다.
        - 기초연금 선택 상태에서 다른 보장을 누르면 방금 누른 다른 보장만 취소합니다.
        - 충돌 팝업은 1회만 표시합니다.
        """
        try:
            if not self.multi_var.get():
                self.current_services = self._read_services_from_widgets()
                self._last_valid_services = tuple(self.current_services)
                return False

            selected_idx = list(self.lb_service.curselection())
            selected = [self.normalize_service(self.lb_service.get(i)) for i in selected_idx]
            clicked_service = None
            if clicked_idx is not None and 0 <= clicked_idx < self.lb_service.size():
                clicked_service = self.normalize_service(self.lb_service.get(clicked_idx))

            # 기초연금 충돌: 방금 누른 항목만 되돌립니다.
            if '기초연금' in selected and len(selected) > 1:
                if clicked_idx is not None and 0 <= clicked_idx < self.lb_service.size():
                    self.lb_service.selection_clear(clicked_idx)
                else:
                    previous = set(getattr(self, '_last_valid_services', tuple()))
                    if '기초연금' in previous:
                        # 기존에 기초연금이 단독 선택되어 있었다면 새로 추가된 다른 보장들을 제거
                        for i in selected_idx:
                            if self.normalize_service(self.lb_service.get(i)) != '기초연금':
                                self.lb_service.selection_clear(i)
                    else:
                        # 기존에 다른 보장들이 선택되어 있었다면 새로 추가된 기초연금을 제거
                        for i in selected_idx:
                            if self.normalize_service(self.lb_service.get(i)) == '기초연금':
                                self.lb_service.selection_clear(i)
                                break
                self._show_service_conflict_popup('기초연금은 다른 보장과 동시 선택이 불가합니다.')
                self.current_services = self._listbox_selected_services()
                self._last_valid_services = tuple(self.current_services)
                self.update_excel_visibility()
                return True

            # 청년내일저축계좌도 기존처럼 단독 선택만 허용하되, 방금 누른 항목만 되돌립니다.
            if '청년내일저축계좌' in selected and len(selected) > 1:
                if clicked_idx is not None and 0 <= clicked_idx < self.lb_service.size():
                    self.lb_service.selection_clear(clicked_idx)
                else:
                    previous = set(getattr(self, '_last_valid_services', tuple()))
                    if '청년내일저축계좌' in previous:
                        for i in selected_idx:
                            if self.normalize_service(self.lb_service.get(i)) != '청년내일저축계좌':
                                self.lb_service.selection_clear(i)
                    else:
                        for i in selected_idx:
                            if self.normalize_service(self.lb_service.get(i)) == '청년내일저축계좌':
                                self.lb_service.selection_clear(i)
                                break
                self._show_service_conflict_popup('청년내일저축계좌는 다른 보장과 함께 선택할 수 없습니다.')
                self.current_services = self._listbox_selected_services()
                self._last_valid_services = tuple(self.current_services)
                self.update_excel_visibility()
                return True

            self.current_services = list(selected)
            self._last_valid_services = tuple(self.current_services)
            return False
        except Exception:
            return False

    def enforce_single_basic_pension_selection(self):
        return self.resolve_service_conflicts()

    def enforce_single_youth_selection(self):
        return self.resolve_service_conflicts()

    def youth_work_enabled(self):
        return '청년내일저축계좌' in self.get_selected_services()

    def disability_info_enabled(self):
        return any(s in self.get_selected_services() for s in ['장애인연금', '차상위장애인'])

    def basic_pension_excel_enabled(self):
        """기초연금이 단독 선택된 경우 엑셀 생성 기능을 표시합니다."""
        try:
            services = self.get_selected_services()
            return len(services) == 1 and services[0] == '기초연금'
        except Exception:
            return False

    def update_excel_visibility(self):
        try:
            show = self.basic_pension_excel_enabled()
            if show:
                if not self.excel_bar.winfo_manager():
                    self.excel_bar.pack(fill='x', padx=0, pady=0, after=self.action_frame)
                if not self.right_panel.winfo_manager():
                    self.right_panel.pack(side='right', fill='y', padx=(8, 0))
            else:
                if self.excel_bar.winfo_manager():
                    self.excel_bar.pack_forget()
                if self.right_panel.winfo_manager():
                    self.right_panel.pack_forget()
        except Exception as e:
            try:
                with open('error_log.txt', 'a', encoding='utf-8') as f:
                    f.write(f'update_excel_visibility error: {e}\n')
            except Exception:
                pass

    def has_input_content(self):
        try:
            return bool(
                self.e_date.get().strip()
                or self.get_selected_services()
                or self.txt.get('1.0', 'end-1c').strip()
                or any(row.path for row in self.file_rows)
            )
        except Exception:
            return True

    def save_autosave_now(self):
        try:
            path = self.get_runtime_data_path('team2_autosave.json')
            tmp = path.with_suffix('.tmp')
            tmp.write_text(json.dumps(self.autosave_payload(), ensure_ascii=False, indent=2), encoding='utf-8')
            tmp.replace(path)
        except Exception:
            pass

    def go_back_to_launcher(self):
        if self.has_input_content() and not messagebox.askyesno(
            '초기화면으로 돌아가기',
            '입력 중인 내용이 있습니다. 자동저장 후 초기화면으로 돌아갈까요?'
        ):
            return
        self.save_autosave_now()
        self.return_to_launcher = True
        self.root.destroy()

    def on_close(self):
        self.save_autosave_now()
        self.root.destroy()

    def update_special_visibility(self):
        """청년내일저축계좌/장애심사 정보 입력 영역을 선택 보장에 맞춰 표시합니다."""
        try:
            if self.youth_work_enabled():
                if not self.youth_work_outer.winfo_ismapped():
                    self.youth_work_outer.pack(fill='x', padx=12, pady=(0, 8), before=self.hf)
            else:
                if self.youth_work_outer.winfo_ismapped():
                    self.youth_work_outer.pack_forget()
                self.youth_not_working_var.set(False)
        except Exception:
            pass

        try:
            if self.disability_info_enabled():
                if not self.disability_outer.winfo_ismapped():
                    self.disability_outer.pack(fill='x', padx=12, pady=(0, 8), before=self.hf)
            else:
                if self.disability_outer.winfo_ismapped():
                    self.disability_outer.pack_forget()
                self.e_disability_review_no.delete(0, 'end')
                self.e_disability_decision_date.delete(0, 'end')
                self.e_disability_type.delete(0, 'end')
        except Exception:
            pass

    def format_disability_section(self):
        if not self.disability_info_enabled():
            return ''
        review_no = self.e_disability_review_no.get().strip()
        decision_date = self.e_disability_decision_date.get().strip()
        disability_type = self.e_disability_type.get().strip()
        lines = []
        if review_no:
            lines.append(f'- 공단심사번호: {review_no}')
        if decision_date:
            lines.append(f'- 장애정도 결정일: {decision_date}')
        if disability_type:
            lines.append(f'- 장애정도 및 유형: {disability_type}')
        return '\n'.join(lines) if lines else ''


    def tabeop_household_type(self):
        """타법의료급여(유공자) 가구유형 선택값을 반환합니다."""
        try:
            v = self.tabeop_type_var.get().strip()
        except Exception:
            v = ''
        return v if v in ('일반가구', '취약가구') else '일반가구'

    def set_tabeop_type(self, value):
        """일반가구/취약가구 체크박스를 라디오 버튼처럼 1개만 선택되게 합니다."""
        if value not in ('일반가구', '취약가구'):
            value = '일반가구'
        self.tabeop_type_var.set(value)
        try:
            self.tabeop_general_var.set(value == '일반가구')
            self.tabeop_vulnerable_var.set(value == '취약가구')
        except Exception:
            pass

    def update_tabeop_visibility(self):
        """타법의료급여(유공자) 선택 여부에 따라 가구유형 선택 영역을 표시/숨김 처리합니다."""
        try:
            services = self.get_selected_services()
            raw_single = self.cb_service.get().strip() if hasattr(self, 'cb_service') else ''
            enabled = ('타법의료급여(유공자)' in services) or (self.normalize_service(raw_single) == '타법의료급여(유공자)')
        except Exception:
            enabled = False

        try:
            if enabled:
                # 주거 보증금 입력 영역 바로 위에 표시합니다.
                if not self.tabeop_outer.winfo_ismapped():
                    if hasattr(self, 'hf') and self.hf.winfo_exists():
                        self.tabeop_outer.pack(fill='x', padx=12, pady=(0, 8), before=self.hf)
                    else:
                        self.tabeop_outer.pack(fill='x', padx=12, pady=(0, 8))
                self.tabeop_outer.lift()
                self.tabeop_general_chk.configure(state='normal')
                self.tabeop_vulnerable_chk.configure(state='normal')
                self.tabeop_hint.configure(text='※ 타법의료급여(유공자) 기준 선택값이 조사자의견에 반영됩니다.')
                if self.tabeop_household_type() not in ('일반가구', '취약가구'):
                    self.set_tabeop_type('일반가구')
            else:
                if self.tabeop_outer.winfo_ismapped():
                    self.tabeop_outer.pack_forget()
        except Exception as e:
            try:
                with open('error_log.txt', 'a', encoding='utf-8') as f:
                    f.write(f'update_tabeop_visibility error: {e}\n')
            except Exception:
                pass

    def show_supporter_area(self):
        if not self.supporter_outer.winfo_ismapped():
            self.supporter_outer.pack(fill='x', padx=12, pady=(0, 8), before=self.hf)
        self.supporter_header.grid(row=0, column=0, sticky='w', pady=(0, 2))
        if not self.supporter_rows:
            self.add_supporter_row()
        self.place_add_supporter_button()

    def clear_supporters(self):
        for item in self.supporter_rows:
            item['frame'].destroy()
        self.supporter_rows = []
        self.supporter_header.grid_forget()
        if self.btn_add_supporter is not None:
            self.btn_add_supporter.destroy()
            self.btn_add_supporter = None
        self.supporter_outer.pack_forget()

    def refresh_supporter_labels(self):
        for i, item in enumerate(self.supporter_rows, start=1):
            item['label'].config(text=f'{i}')

    def remove_supporter_row(self, rowf):
        if len(self.supporter_rows) == 1:
            messagebox.showwarning('안내', '최소 1개의 부양의무자 입력칸은 유지됩니다.')
            return
        target = None
        for item in self.supporter_rows:
            if item['frame'] is rowf:
                target = item
                break
        if target:
            target['frame'].destroy()
            self.supporter_rows.remove(target)
            self.refresh_supporter_labels()
            for idx, item in enumerate(self.supporter_rows, start=1):
                item['frame'].grid_configure(row=idx)
            self.place_add_supporter_button()

    def place_add_supporter_button(self):
        if self.btn_add_supporter is not None:
            self.btn_add_supporter.destroy()
            self.btn_add_supporter = None
        if not self.supporter_enabled():
            return
        row_index = len(self.supporter_rows) + 1
        self.btn_add_supporter = ttk.Button(self.supporter_wrap, text='부양의무자 추가', command=self.add_supporter_row)
        self.btn_add_supporter.grid(row=row_index, column=0, sticky='e', padx=2, pady=2)

    def add_supporter_row(self, name='', birth='', income='', members='', relation=''):
        if self.btn_add_supporter is not None:
            self.btn_add_supporter.destroy()
            self.btn_add_supporter = None

        idx = len(self.supporter_rows) + 1
        rowf = ttk.Frame(self.supporter_wrap)
        rowf.grid(row=idx, column=0, sticky='w', pady=2)

        lbl = ttk.Label(rowf, text=f'{idx}', width=3, anchor='center')
        lbl.grid(row=0, column=0, padx=(0,4), pady=1)

        er = ttk.Entry(rowf, width=14)
        er.grid(row=0, column=1, padx=2, pady=1)

        en = ttk.Entry(rowf, width=10)
        en.grid(row=0, column=2, padx=2, pady=1)

        eb = ttk.Entry(rowf, width=10)
        eb.grid(row=0, column=3, padx=2, pady=1)
        eb.bind('<FocusOut>', lambda e: self.format_birth_entry(e.widget))

        ei = ttk.Entry(rowf, width=13)
        ei.grid(row=0, column=4, padx=2, pady=1)
        ei.bind('<FocusOut>', self.format_entry_money)

        cm = ttk.Combobox(rowf, state='readonly', width=6, values=['1인','2인','3인','4인','5인','6인','7인'])
        cm.grid(row=0, column=5, padx=2, pady=1)

        ttk.Button(rowf, text='삭제', command=lambda rf=rowf: self.remove_supporter_row(rf)).grid(row=0, column=6, padx=2, pady=1)

        if relation:
            er.insert(0, relation)
        if members:
            cm.set(members)
        if name:
            en.insert(0, name)
        if birth:
            eb.insert(0, birth)
        if income:
            ei.insert(0, income)

        self.supporter_rows.append({'frame': rowf, 'label': lbl, 'relation': er, 'name': en, 'birth': eb, 'income': ei, 'members': cm})
        self.refresh_supporter_labels()
        self.place_add_supporter_button()

    def format_birth_entry(self, entry):
        v = re.sub(r'[^0-9]', '', entry.get().strip())[:6]
        entry.delete(0, 'end')
        if v:
            entry.insert(0, v)

    def supporter_enabled(self):
        return '차상위본인부담경감' in self.get_selected_services()

    def get_supporters(self):
        vals = []
        for r in self.supporter_rows:
            relation = r.get('relation').get().strip() if r.get('relation') else ''
            name = r['name'].get().strip()
            birth = re.sub(r'[^0-9]', '', r['birth'].get().strip())[:6]
            income = r['income'].get().strip()
            members = r['members'].get().strip()
            if relation or name or birth or income or members:
                vals.append({'relation': relation, 'name': name, 'birth': birth, 'income': income, 'members': members})
        return vals

    def validate_supporters(self):
        if not self.supporter_enabled():
            return True
        for i, s in enumerate(self.get_supporters(), start=1):
            if (s['name'] or s['birth'] or s['income']) and not s['members']:
                messagebox.showwarning('안내', f'부양의무자 {i}의 세대원 수를 선택하세요.')
                return False
        return True

    def has_supporter_with_ability(self):
        if not self.supporter_enabled():
            return False
        for s in self.get_supporters():
            ability = self.supporter_ability_text(s.get('income', ''), s.get('members', ''))
            if ability == '부양능력 있음':
                return True
        return False

    def supporter_ability_text(self, income_text, members_text):
        v = re.sub(r'[^0-9]', '', income_text or '')
        if not v:
            return ''
        m = re.sub(r'[^0-9]', '', members_text or '')
        if not m:
            return ''
        income = int(v)
        size = int(m)
        threshold = SUPPORTER_ABILITY_BY_SIZE.get(size, SUPPORTER_ABILITY_BY_SIZE[1])
        return '부양능력 있음' if income >= threshold else '부양능력 없음'

    def on_house_change(self, event=None):
        if self.cb_house.get().strip() == '직접입력':
            self.e_house_direct.config(state='normal')
        else:
            self.e_house_direct.config(state='normal')
            self.e_house_direct.delete(0, 'end')
            self.e_house_direct.config(state='disabled')

    def household_display(self):
        if '타법의료급여(유공자)' in self.get_selected_services():
            return self.tabeop_household_type()
        return self.e_house_direct.get().strip() if self.cb_house.get().strip() == '직접입력' else self.cb_house.get().strip()

    def set_state(self, entry, state, clear=False):
        entry.config(state='normal')
        if clear:
            entry.delete(0, 'end')
        entry.config(state=state)

    def on_home_change(self, event=None):
        for e in [self.e_deposit, self.e_rent]:
            self.set_state(e, 'disabled', clear=True)
        self.set_state(self.e_home_direct, 'disabled', clear=True)
        k = self.cb_home.get().strip()
        if k == '민간(전세)':
            self.set_state(self.e_deposit, 'normal')
        elif k in (
            '공공건설임대(영구임대)', '공공건설임대(국민임대)',
            '공공건설임대(매입임대)', '민간(보증부월세)'
        ):
            self.set_state(self.e_deposit, 'normal')
            self.set_state(self.e_rent, 'normal')
        elif k in ('전세임대', '민간(월세)'):
            self.set_state(self.e_rent, 'normal')
        elif k == '기타(직접입력)':
            self.set_state(self.e_home_direct, 'normal')
            self.set_state(self.e_deposit, 'normal')
            self.set_state(self.e_rent, 'normal')

    def format_entry_money(self, event=None):
        e = event.widget if event else None
        if e is None or str(e.cget('state')) == 'disabled':
            return

        # 여러보장 선택 시 소득인정액 대표 입력칸에는 금액이 아니라 '입력완료' 표시를 유지합니다.
        if e == self.e_income and self.multi_var.get():
            if self.multi_income_values:
                e.delete(0, 'end')
                e.insert(0, '입력완료')
            return

        if e.get().strip() == '입력완료':
            return

        v = re.sub(r'[^0-9]', '', e.get().strip())
        e.delete(0, 'end')
        if v:
            e.insert(0, f'{int(v):,}')

    def refresh_multi_income_inputs(self):
        """여러 보장 선택 시 별도 보장별 소득인정액 영역은 표시하지 않습니다.

        각 보장별 금액 입력은 기존처럼 대표 소득인정액 칸을 클릭해 팝업에서 처리합니다.
        """
        services = self.get_selected_services()
        self.multi_income_values = {s: self.multi_income_values.get(s, '') for s in services}
        self.multi_income_frame.pack_forget()

    def open_multi_income_popup(self, event=None):
        services = self.get_selected_services()
        if not (self.multi_var.get() and len(services) > 1):
            return
        if self.multi_income_popup_open:
            return 'break'
        self.multi_income_popup_open = True
        try:
            popup = MultiIncomePopup(self.root, services, self.multi_income_values)
            if popup.result is not None:
                self.multi_income_values = popup.result
                self.refresh_multi_income_inputs()
                filled = [v for v in self.multi_income_values.values() if v]
                self.e_income.delete(0, 'end')
                self.e_income.insert(0, '입력완료' if filled else '')
        finally:
            self.multi_income_popup_open = False
        return 'break'

    def get_income_text(self, service):
        if self.multi_var.get() and len(self.get_selected_services()) > 1:
            return self.multi_income_values.get(service, '')
        return self.e_income.get().strip()

    def get_income_num(self, service):
        v = re.sub(r'[^0-9]', '', self.get_income_text(service))
        return int(v) if v else None

    def add_file_row(self):
        row = FileRow(self.file_area, len(self.file_rows)+1, self.remove_file_row, self.update_file_row_identity)
        row.frame.pack(fill='x', pady=2)
        self.file_rows.append(row)
        self.refresh_row_names()

    def update_file_row_identity(self, row):
        """엑셀을 선택하는 즉시 성명과 생년월일을 가구원 표시란에 보여줍니다."""
        if not row.path:
            return
        try:
            df = self.parse_single_excel(row.path, self.e_file_password.get().strip())
            names = [x for x in df['성명'].dropna().unique().tolist() if self.clean(x)] if '성명' in df.columns else []
            labels = []
            for name in names:
                rrn = ''
                if '주민등록번호' in df.columns:
                    matched = df[df['성명'].astype(str).map(self.clean) == self.clean(name)]
                    if not matched.empty:
                        rrn = self.mask_rrn(matched.iloc[0].get('주민등록번호', ''))
                labels.append(f'{name}({rrn})' if rrn else str(name))
            if len(labels) == 1:
                row.set_display_name(labels[0])
            elif len(labels) > 1:
                row.set_display_name(', '.join(labels[:2]) + (' 외' if len(labels) > 2 else ''))
        except WrongPasswordError:
            row.set_display_name('비밀번호 확인 필요')
        except Exception:
            # 선택 단계에서는 파일명은 유지하고, 자세한 오류는 '첨부 파일 반영하기'에서 안내합니다.
            pass

    def remove_file_row(self, row):
        if len(self.file_rows) == 1:
            messagebox.showwarning('안내', '최소 1개의 입력칸은 유지됩니다.')
            return
        row.frame.destroy()
        self.file_rows.remove(row)
        self.refresh_row_names()

    def refresh_row_names(self):
        for i, r in enumerate(self.file_rows, start=1):
            if r.display_name.startswith('가구원 '):
                r.set_display_name(f'가구원 {i}')

    def clean(self, text):
        t = str(text).replace('_x000D_', ' ').replace('\n', ' ').replace('\r', ' ')
        return re.sub(r'\s+', ' ', t).strip()

    def normalize_rrn(self, rrn):
        """주민등록번호/생년월일을 문자열로 정리하고 앞자리 0 누락을 보정합니다.
        엑셀에서 010101 같은 값이 숫자로 읽히면 10101처럼 바뀌는 문제를 막기 위한 처리입니다.
        """
        if rrn is None:
            return ''
        text = str(rrn).strip()
        if text.lower() in ('nan', 'none'):
            return ''
        # 엑셀 숫자값이 10101.0처럼 들어오는 경우 .0 제거
        text = re.sub(r'\.0$', '', text)
        digits = re.sub(r'[^0-9]', '', text)
        if not digits:
            return self.clean(text)
        # 생년월일 6자리만 있는 경우: 010101 보정
        if len(digits) <= 6:
            return digits.zfill(6)
        # 주민등록번호 13자리인 경우: 010101-3****** 보정
        if len(digits) < 13:
            return digits.zfill(13)
        return digits[:13]

    def mask_rrn(self, rrn):
        s = self.normalize_rrn(rrn)
        digits = re.sub(r'[^0-9]', '', str(s))
        return digits[:6] + '-' + digits[6] if len(digits) >= 7 else digits

    def read_excel(self, path, password):
        if not password:
            raw = pd.read_excel(path, sheet_name=0, header=None)
            return raw, path
        if msoffcrypto is None:
            raise ValueError('비밀번호 기능 사용 불가')
        try:
            with open(path, 'rb') as f:
                office = msoffcrypto.OfficeFile(f)
                office.load_key(password=password)
                bio = io.BytesIO()
                office.decrypt(bio)
                bio.seek(0)
                return pd.read_excel(bio, sheet_name=0, header=None), bio
        except Exception:
            raise WrongPasswordError()

    def parse_single_excel(self, path, password):
        raw, src = self.read_excel(path, password)
        if hasattr(src, 'seek'):
            src.seek(0)
        header_row = None
        # 엑셀 서식에 따라 헤더가 아래쪽에 있을 수 있어 30행까지 확인
        for i in range(min(30, len(raw))):
            vals = [self.clean(x) for x in raw.iloc[i].tolist()]
            if '성명' in vals and '주민등록번호' in vals and '소득재산상세분류코드명' in vals:
                header_row = i
                break
        if header_row is None:
            raise ValueError('헤더 행을 찾을 수 없습니다.')
        if hasattr(src, 'seek'):
            src.seek(0)
        df = pd.read_excel(src, sheet_name=0, header=header_row)
        df.columns = [self.clean(c) for c in df.columns]

        need = ['성명','주민등록번호','소득재산상세분류코드명','최신금액','반영금액','반영여부']
        for c in need:
            if c not in df.columns:
                raise ValueError(f'필수 열 없음: {c}')

        # 반영내용은 상담내역에 출력하지 않지만, 엑셀에 있으면 정리만 해둠
        for c in ['성명','주민등록번호','소득재산상세분류코드명','반영여부','반영내용','소득재산상세분류코드']:
            if c in df.columns:
                df[c] = df[c].fillna('').astype(str).map(self.clean)
        if '주민등록번호' in df.columns:
            df['주민등록번호'] = df['주민등록번호'].map(self.normalize_rrn)

        base = df[df['성명'].ne('')].copy()
        if base.empty:
            return base

        # 원칙: 반영여부에 '반영'이 포함된 행만 사용
        reflected = base[base['반영여부'].astype(str).str.contains('반영', na=False)].copy()

        # 다만 엑셀 양식/값 문제로 반영 행이 0건이면, 신청인·소득·재산이 전부 비는 문제를 막기 위해
        # 금액이 있는 행을 예비로 사용
        if reflected.empty:
            def has_amount(row):
                return self.amount_int(row) not in (None, 0)
            fallback = base[base.apply(has_amount, axis=1)].copy()
            return fallback if not fallback.empty else base

        return reflected

    def load_all_excels(self):
        rows = [r for r in self.file_rows if r.path]
        if not rows:
            messagebox.showwarning('안내', '먼저 엑셀 파일을 선택하세요.')
            return
        frames = []
        try:
            for r in rows:
                df = self.parse_single_excel(r.path, self.e_file_password.get().strip())
                frames.append(df)
                names = [x for x in df['성명'].dropna().unique().tolist() if self.clean(x)] if '성명' in df.columns else []
                labels = []
                for name in names:
                    rrn = ''
                    if '주민등록번호' in df.columns:
                        matched = df[df['성명'].astype(str).map(self.clean) == self.clean(name)]
                        if not matched.empty:
                            rrn = self.mask_rrn(matched.iloc[0].get('주민등록번호', ''))
                    labels.append(f'{name}({rrn})' if rrn else str(name))
                if len(labels) == 1:
                    r.set_display_name(labels[0])
                elif len(labels) > 1:
                    r.set_display_name(', '.join(labels[:2]) + (' 외' if len(labels) > 2 else ''))
            self.current_df = pd.concat(frames, ignore_index=True) if frames else None
            self.txt.delete('1.0', 'end')
            if self.current_df is None or self.current_df.empty:
                messagebox.showwarning('안내', '엑셀은 읽었지만 불러올 행이 없습니다. 반영여부, 성명, 금액 칸을 확인하세요.')
            else:
                messagebox.showinfo('완료', f'{len(rows)}개 파일, {len(self.current_df)}개 행을 반영했습니다.')
        except WrongPasswordError:
            messagebox.showerror('오류', '비밀번호가 틀렸습니다')
        except Exception as e:
            messagebox.showerror('오류', f'엑셀 읽기 실패\n{e}')

    def amount_int(self, row):
        for col in ['반영금액', '최신금액']:
            v = row.get(col, None)
            if pd.isna(v) or str(v).strip() == '':
                continue
            try:
                return int(float(str(v).replace(',', '').strip()))
            except Exception:
                digits = re.sub(r'[^0-9-]', '', str(v))
                if digits not in ['', '-']:
                    try:
                        return int(digits)
                    except Exception:
                        pass
        return None

    def pick_amount(self, row):
        iv = self.amount_int(row)
        if iv is None or iv == 0:
            return ''
        return f'{iv:,}원'

    def classify(self, code_value, name=''):
        s = str(code_value).strip()
        if s.startswith('1'):
            return '소득사항'
        if s.startswith('2'):
            return '재산사항'
        if s.startswith('31') or s.startswith('32') or s.startswith('33'):
            return '공제'
        if s.startswith('4'):
            return '부채'
        if s.startswith('3'):
            return '공제'
        n = str(name)
        if '공제' in n:
            return '공제'
        if any(k in n for k in ['부채', '대출', '차용']):
            return '부채'
        if any(k in n for k in ['재산', '증여', '주택', '건물', '토지', '차량', '금융', '예금', '보험', '증권']):
            return '재산사항'
        if any(k in n for k in ['급여', '연금', '수당', '소득', '무료임차료', '보상금']):
            return '소득사항'
        return '기타'

    def is_daily_income(self, name):
        n = self.clean(name)
        return '일용' in n or '일용근로' in n

    def hide_row_for_service_rule(self, row):
        nm = self.clean(row.get('소득재산상세분류코드명', ''))
        services = self.get_selected_services()

        # 사적이전소득은 저소득 한부모 또는 경기도 한부모 보장을 선택한 경우에만 표시합니다.
        if '사적이전소득' in nm and not any(s in services for s in ['저소득 한부모', '경기도 한부모']):
            return True

        # 기초연금 단독 신청일 때만 일용근로/일용소득/실업급여/기초연금 소득을 제외합니다.
        # 기초연금과 다른 보장을 동시에 선택한 경우에는 모든 소득내역을 그대로 표시합니다.
        if services == ['기초연금'] and self.is_basic_pension_excluded_income(nm):
            return True
        # 기초연금 신청 시 장애인연금 소득은 소득사항에서 제외하고 소득공제로 옮깁니다.
        if self.is_basic_pension_disability_income(row):
            return True
        return False

    def is_basic_pension_disability_income(self, row):
        if self.get_selected_services() != ['기초연금']:
            return False
        nm = self.clean(row.get('소득재산상세분류코드명', ''))
        code_val = self.clean(row.get('소득재산상세분류코드', ''))
        return '장애인연금' in nm and self.classify(code_val, nm) == '소득사항'


    def is_financial_asset(self, name):
        n = self.clean(name)
        return any(k in n for k in ['금융', '예금', '적금', '보험', '증권', '주식', '펀드', '채권', '퇴직연금', '연금저축', '기타일시금'])

    def total_amount_for(self, category, financial=None):
        if self.current_df is None or self.current_df.empty:
            return 0
        total = 0
        for _, row in self.current_df.iterrows():
            nm = self.clean(row.get('소득재산상세분류코드명', ''))
            code_val = self.clean(row.get('소득재산상세분류코드', ''))
            if self.classify(code_val, nm) != category:
                continue
            if self.hide_row_for_service_rule(row):
                continue
            if financial is not None and self.is_financial_asset(nm) != financial:
                continue
            iv = self.amount_int(row)
            if iv is None or iv == 0:
                continue
            total += iv
        return total

    def section_summary_lines(self, category):
        if category == '소득사항':
            return [f'소득 총액: {self.total_amount_for("소득사항"):,}원']
        if category == '재산사항':
            general = self.total_amount_for('재산사항', financial=False)
            financial = self.total_amount_for('재산사항', financial=True)
            return [f'일반재산 총액: {general:,}원', f'금융재산 총액: {financial:,}원']
        return []

    def collect_person_items(self, category, financial=None):
        person_map = {}
        seen = set()
        if self.current_df is None or self.current_df.empty:
            return person_map
        for _, row in self.current_df.iterrows():
            nm = self.clean(row.get('소득재산상세분류코드명', ''))
            code_val = self.clean(row.get('소득재산상세분류코드', ''))
            if self.classify(code_val, nm) != category:
                continue
            if self.hide_row_for_service_rule(row):
                continue
            if financial is not None and self.is_financial_asset(nm) != financial:
                continue
            amt = self.pick_amount(row)
            if not amt:
                continue
            person = self.clean(row.get('성명', '')) or '성명미상'
            item = f'{nm} {amt}'.strip()
            key = (person, item)
            if item and key not in seen:
                seen.add(key)
                person_map.setdefault(person, []).append(item)
        return person_map

    def format_person_map(self, person_map):
        if not person_map:
            return '해당사항 없음'
        blocks = []
        for person, items in person_map.items():
            block = [f'- {person}:']
            for item in items:
                block.append(f'  · {item}')
            blocks.append('\n'.join(block))
        return '\n\n'.join(blocks)

    def collect_asset_text(self):
        general_total = self.total_amount_for('재산사항', financial=False)
        financial_total = self.total_amount_for('재산사항', financial=True)
        general_map = self.collect_person_items('재산사항', financial=False)
        financial_map = self.collect_person_items('재산사항', financial=True)

        blocks = [
            f'일반재산 총액: {general_total:,}원',
            self.format_person_map(general_map),
            f'금융재산 총액: {financial_total:,}원',
            self.format_person_map(financial_map),
        ]
        return '\n\n'.join(blocks)

    def deduction_kind(self, code_value):
        s = str(code_value).strip()
        if s.startswith('31'):
            return '소득공제'
        if s.startswith('32'):
            return '일반재산 공제'
        if s.startswith('33'):
            return '금융재산 공제'
        return '기타 공제'

    def collect_deduction_person_items(self, kind):
        person_map = {}
        seen = set()
        if self.current_df is None or self.current_df.empty:
            return person_map
        for _, row in self.current_df.iterrows():
            nm = self.clean(row.get('소득재산상세분류코드명', ''))
            code_val = self.clean(row.get('소득재산상세분류코드', ''))
            if self.classify(code_val, nm) != '공제':
                continue
            if self.deduction_kind(code_val) != kind:
                continue
            amt = self.pick_amount(row)
            if not amt:
                continue
            person = self.clean(row.get('성명', '')) or '성명미상'
            item = f'{nm} {amt}'.strip()
            key = (person, item)
            if item and key not in seen:
                seen.add(key)
                person_map.setdefault(person, []).append(item)
        return person_map

    def deduction_total_for_kind(self, kind):
        if self.current_df is None or self.current_df.empty:
            return 0
        total = 0
        for _, row in self.current_df.iterrows():
            nm = self.clean(row.get('소득재산상세분류코드명', ''))
            code_val = self.clean(row.get('소득재산상세분류코드', ''))
            if self.classify(code_val, nm) != '공제':
                continue
            if self.deduction_kind(code_val) != kind:
                continue
            iv = self.amount_int(row)
            if iv is None or iv == 0:
                continue
            total += iv
        return total

    def collect_deduction_text(self):
        blocks = []
        for kind in ['소득공제', '일반재산 공제', '금융재산 공제', '기타 공제']:
            person_map = self.collect_deduction_person_items(kind)
            total = self.deduction_total_for_kind(kind)
            if total == 0 and not person_map:
                continue
            blocks.append(f'{kind} 총액: {total:,}원')
            blocks.append(self.format_person_map(person_map))
        return '\n\n'.join(blocks) if blocks else '해당사항 없음'


    def asset_date_text(self):
        try:
            val = self.e_asset_date.get().strip()
        except Exception:
            val = ''
        if not val:
            return ''
        digits = re.sub(r'[^0-9]', '', val)
        if len(digits) == 8:
            val = f'{digits[:4]}.{digits[4:6]}.{digits[6:8]}.'
        return f' (조회기준일 {val})'


    def collect_category_text(self, category):
        if self.current_df is None or self.current_df.empty:
            return '해당사항 없음'

        # 소득사항
        if category == '소득사항':
            total = self.total_amount_for('소득사항')
            items = {}

            for _, row in self.current_df.iterrows():
                nm = self.clean(row.get('소득재산상세분류코드명', ''))
                code_val = self.clean(row.get('소득재산상세분류코드', ''))

                if self.classify(code_val, nm) != '소득사항':
                    continue

                if self.hide_row_for_service_rule(row):
                    continue

                iv = self.amount_int(row)
                if iv is None or iv == 0:
                    continue

                items[nm] = items.get(nm, 0) + iv

            lines = [f'소득총액 {total:,}원']

            for item_name, amount in items.items():
                lines.append(f'- {item_name} {amount:,}원')

            return '\n'.join(lines)

        # 재산사항
        if category == '재산사항':
            general_total = self.total_amount_for('재산사항', financial=False)
            financial_total = self.total_amount_for('재산사항', financial=True)

            general_items = {}
            financial_items = {}

            for _, row in self.current_df.iterrows():
                nm = self.clean(row.get('소득재산상세분류코드명', ''))
                code_val = self.clean(row.get('소득재산상세분류코드', ''))

                if self.classify(code_val, nm) != '재산사항':
                    continue

                iv = self.amount_int(row)
                if iv is None or iv == 0:
                    continue

                if self.is_financial_asset(nm):
                    financial_items[nm] = financial_items.get(nm, 0) + iv
                else:
                    general_items[nm] = general_items.get(nm, 0) + iv

            lines = [f'일반재산 총액 {general_total:,}원']

            for item_name, amount in general_items.items():
                lines.append(f'- {item_name} {amount:,}원')

            lines.append('')
            lines.append(f'금융재산 총액 {financial_total:,}원{self.asset_date_text()}')

            for item_name, amount in financial_items.items():
                lines.append(f'- {item_name} {amount:,}원')

            return '\n'.join(lines)

        # 공제
        if category == '공제':
            income_deduction = {}
            general_deduction = {}
            financial_deduction = {}

            income_total = 0
            general_total = 0
            financial_total = 0

            for _, row in self.current_df.iterrows():
                nm = self.clean(row.get('소득재산상세분류코드명', ''))
                code_val = self.clean(row.get('소득재산상세분류코드', ''))

                iv = self.amount_int(row)
                if iv is None or iv == 0:
                    continue

                if code_val.startswith('31') or self.is_basic_pension_disability_income(row):
                    income_total += iv
                    income_deduction[nm] = income_deduction.get(nm, 0) + iv

                elif code_val.startswith('32'):
                    general_total += iv
                    general_deduction[nm] = general_deduction.get(nm, 0) + iv

                elif code_val.startswith('33'):
                    financial_total += iv
                    financial_deduction[nm] = financial_deduction.get(nm, 0) + iv

            lines = []

            if income_total:
                lines.append(f'소득공제 총액 {income_total:,}원')
                for item_name, amount in income_deduction.items():
                    if amount:
                        lines.append(f'- {item_name} {amount:,}원')

            if general_total:
                if lines:
                    lines.append('')
                lines.append(f'일반재산 공제 총액 {general_total:,}원')
                for item_name, amount in general_deduction.items():
                    if amount:
                        lines.append(f'- {item_name} {amount:,}원')

            if financial_total:
                if lines:
                    lines.append('')
                lines.append(f'금융재산 공제 총액 {financial_total:,}원')
                for item_name, amount in financial_deduction.items():
                    if amount:
                        lines.append(f'- {item_name} {amount:,}원')

            return '\n'.join(lines) if lines else '해당사항 없음'

        # 부채
        if category == '부채':
            total = 0
            items = {}

            for _, row in self.current_df.iterrows():
                nm = self.clean(row.get('소득재산상세분류코드명', ''))
                code_val = self.clean(row.get('소득재산상세분류코드', ''))

                if not str(code_val).startswith('4'):
                    continue

                iv = self.amount_int(row)
                if iv is None or iv == 0:
                    continue

                total += iv
                items[nm] = items.get(nm, 0) + iv

            lines = [f'부채 총액 {total:,}원']

            for item_name, amount in items.items():
                lines.append(f'- {item_name} {amount:,}원')

            return '\n'.join(lines)

        return '해당사항 없음'

    def format_service_line(self):
        return ', '.join(self.get_selected_services())

    def format_home(self):
        selected = self.cb_home.get().strip()
        k = self.e_home_direct.get().strip() if selected == '기타(직접입력)' else selected
        k = k or '직접입력'
        extra = []
        if selected == '민간(전세)' and self.e_deposit.get().strip():
            extra.append(f'보증금 {self.e_deposit.get().strip()}원')
        elif selected in (
            '공공건설임대(영구임대)', '공공건설임대(국민임대)',
            '공공건설임대(매입임대)', '민간(보증부월세)'
        ):
            if self.e_deposit.get().strip(): extra.append(f'보증금 {self.e_deposit.get().strip()}원')
            if self.e_rent.get().strip(): extra.append(f'임대료 {self.e_rent.get().strip()}원')
        elif selected in ('전세임대', '민간(월세)'):
            if self.e_rent.get().strip(): extra.append(f'임대료 {self.e_rent.get().strip()}원')
        elif selected == '기타(직접입력)':
            if self.e_deposit.get().strip(): extra.append(f'보증금 {self.e_deposit.get().strip()}원')
            if self.e_rent.get().strip(): extra.append(f'임대료 {self.e_rent.get().strip()}원')
        return f"{k} ({', '.join(extra)})" if k and extra else k

    def household_count(self):
        if self.current_df is None or self.current_df.empty:
            return 0
        return len(self.current_df[['성명','주민등록번호']].drop_duplicates())

    def suitability_by_threshold(self, service):
        amt = self.get_income_num(service)
        if amt is None:
            return None, ''

        if service == '타법의료급여(유공자)':
            family_type = self.tabeop_household_type()
            standard = '타법의료급여_취약가구' if family_type == '취약가구' else '타법의료급여_일반가구'
            threshold = THRESHOLDS.get(standard, {}).get(self.household_count())
            if threshold is None:
                return None, f'{service}: 기준 없음'
            ok = amt <= threshold
            base = f'{family_type} {self.household_count()}인 가구 기준 소득인정액 {threshold:,}원'
            return ok, f'{service}: {base} 이하로 적합' if ok else f'{service}: {base} 초과로 제외'

        standard = SERVICE_TO_STANDARD.get(service)
        if not standard:
            return None, f'{service}: 개발예정'
        threshold = THRESHOLDS.get(standard, {}).get(self.household_count())
        if threshold is None:
            return None, f'{service}: 기준 없음'
        ok = amt <= threshold
        base = f'{self.household_count()}인 가구 기준 소득인정액 {threshold:,}원'
        return ok, f'{service}: {base} 이하로 적합' if ok else f'{service}: {base} 초과로 제외'

    def investigator_opinion_for_service(self, service, selected):
        if service == '청년내일저축계좌' and self.youth_not_working_var.get():
            return '청년내일저축계좌: 현재 미근로로 청년내일저축계좌 부적합'
        if service == '차상위본인부담경감' and self.has_supporter_with_ability():
            return '차상위본인부담경감: 부양의무자 부양능력 있음으로 보장 제외'
        if service == '경기도 한부모' and '저소득 한부모' in selected:
            ok1, _ = self.suitability_by_threshold('저소득 한부모')
            ok2, _ = self.suitability_by_threshold('경기도 한부모')
            if ok1 is True and ok2 is True:
                return '경기도 한부모: 상위보장(저소득 한부모) 책정으로 제외'
        if service == '차상위계층확인':
            ok_check, _ = self.suitability_by_threshold('차상위계층확인')

            # 차상위본인부담경감은 부양의무자 부양능력 있음으로 제외되는 경우
            # 차상위계층확인 제외 사유로 적용하지 않습니다.
            if '차상위본인부담경감' in selected:
                ok_med, _ = self.suitability_by_threshold('차상위본인부담경감')
                if ok_check is True and ok_med is True and not self.has_supporter_with_ability():
                    return '차상위계층확인: 차상위본인부담경감 책정으로 보장 제외'

            # 차상위장애인, 차상위자활이 최종 적합이면 차상위계층확인은 중복 책정 불가로 제외
            for upper_service in ['차상위장애인', '차상위자활']:
                if upper_service in selected:
                    ok_upper, _ = self.suitability_by_threshold(upper_service)
                    if ok_check is True and ok_upper is True:
                        return f'차상위계층확인: {upper_service} 책정으로 보장 제외'
        if service == '차상위계층확인' and any(x in self.selected_beneficiary for x in ['생계', '주거', '교육', '의료']):
            vals = [('의료급여' if x == '의료' else x) for x in self.selected_beneficiary if x in ['생계', '주거', '교육', '의료']]
            return f'차상위계층확인: 상위보장 책정({"·".join(vals)})으로 보장 제외'
        if service == '차상위본인부담경감' and '의료' in self.selected_beneficiary:
            return '차상위본인부담경감: 상위보장 책정(의료급여)으로 보장 제외'
        _, txt = self.suitability_by_threshold(service)
        return txt

    def service_result_tag(self, service, selected):
        if service == '차상위본인부담경감' and self.has_supporter_with_ability():
            return '제외'
        op = self.investigator_opinion_for_service(service, selected)
        return '적합' if ('적합' in op and '부적합' not in op and '제외' not in op and '초과' not in op) else '제외'

    def format_multiline_section(self, no, title, content):
        return f'{no}. {title}: {content}' if '\n' not in content else f'{no}. {title}\n\n{content}'

    def extra_note(self):
        return '현재 소득인정액만으로 기준 초과하여 참고자료 조사가 필요하지 않아 참고자료 미조사' if self.ref_var.get() else ''

    def format_income_section(self, services):
        if self.multi_var.get() and len(services) > 1:
            income_lines = [
                f'- {s}: {self.get_income_text(s)}원' if self.get_income_text(s) else f'- {s}:'
                for s in services
            ]
            return '\n'.join(income_lines)
        income_text = self.get_income_text(services[0] if services else '')
        return f'{income_text}원' if income_text else ''

    def format_supporter_section(self):
        supporters = self.get_supporters()
        if not supporters:
            return ''
        lines = []
        for i, s in enumerate(supporters, start=1):
            ability = self.supporter_ability_text(s.get('income', ''), s.get('members', ''))
            parts = [
                f"관계 {s.get('relation', '')}" if s.get('relation') else '',
                f"성명 {s.get('name', '')}" if s.get('name') else '',
                f"생년월일 {s.get('birth', '')}" if s.get('birth') else '',
                f"소득인정액 {s.get('income', '')}원" if s.get('income') else '',
                f"세대원 수 {s.get('members', '')}" if s.get('members') else '',
                ability,
            ]
            lines.append(f"- 부양의무자 {i}: " + ', '.join([x for x in parts if x]))
        return '\n'.join(lines)

    def is_basic_pension_excluded_income(self, item_name):
        text = str(item_name)
        return any(x in text for x in ['일용근로', '일용소득', '실업급여', '기초연금'])

    def build_text(self):
        services = self.get_selected_services()
        uniq = self.current_df[['성명','주민등록번호']].drop_duplicates()
        applicants = '\n'.join([f"- {self.clean(r['성명'])} ({self.mask_rrn(r['주민등록번호'])})" for _, r in uniq.iterrows()])

        lines = []
        for s in services:
            lines.append(f'[{s} 신청 조사 결과 -> {self.service_result_tag(s, services)}]')

        lines.extend([
            f'1. 신청 가구원\n\n{applicants}',
            f'2. 신청일자: {self.format_date()}',
            f'3. 보장구분: {self.format_service_line()}',
            f'4. 가구유형: {self.household_display()}',
            f'5. 주거유형: {self.format_home()}',
            self.format_multiline_section(6, '소득사항', self.collect_category_text('소득사항')),
            self.format_multiline_section(7, '재산사항', self.collect_category_text('재산사항')),
            self.format_multiline_section(8, '공제', self.collect_category_text('공제')),
            self.format_multiline_section(9, '부채', self.collect_category_text('부채')),
        ])

        next_no = 10

        # 부양의무자는 차상위본인부담경감 선택 시에만 표시
        if self.supporter_enabled():
            supporter_section = self.format_supporter_section()
            if supporter_section:
                lines.append(f'{next_no}. 부양의무자\n\n' + supporter_section)
            else:
                lines.append(f'{next_no}. 부양의무자\n\n해당사항 없음')
            next_no += 1

        income_section = self.format_income_section(services)
        lines.append(self.format_multiline_section(next_no, '소득인정액', income_section))
        next_no += 1

        disability_section = self.format_disability_section()
        if disability_section:
            lines.append(f'{next_no}. 장애심사 정보\n\n' + disability_section)
            next_no += 1

        opinions = [self.investigator_opinion_for_service(s, services) for s in services if self.investigator_opinion_for_service(s, services)]
        if len(opinions) <= 1:
            lines.append(f'{next_no}. 조사자의견: {opinions[0] if opinions else ""}')
        else:
            lines.append(f'{next_no}. 조사자의견')
            for op in opinions:
                lines.append(f'- {op}')
        next_no += 1

        if self.extra_note():
            lines.append(f'{next_no}. 참고사항: {self.extra_note()}')

        text = '\n'.join(lines).strip()
        text = text.replace('\n1. ', '\n\n1. ', 1)
        for n in range(2, 15):
            text = text.replace(f'\n{n}. ', f'\n\n{n}. ')
        return text

    def generate_preview(self):
        self.validate_date()
        if not self.e_date.get().strip():
            return

        if self.current_df is None:
            messagebox.showwarning('안내', '먼저 첨부 파일 반영하기를 실행하세요.')
            return
        if self.current_df.empty:
            messagebox.showwarning('안내', '첨부 파일은 반영됐지만 불러온 행이 없습니다. 엑셀의 반영여부, 성명, 금액 칸을 확인하세요.')
            return

        if not self.validate_supporters():
            return

        services_before_generate = self.get_selected_services()

        try:
            result_text = self.build_text()
        except Exception as e:
            messagebox.showerror('오류', f'상담내역 생성 중 오류가 발생했습니다.\n{e}')
            return

        old_text = self.txt.get('1.0', 'end-1c')
        self.txt.delete('1.0', 'end')
        self.txt.insert('1.0', result_text)
        self.highlight_changed_lines(old_text, result_text)
        self.update_char_count()
        self.update_applicant_preview()
        self.save_recent_input()


    def current_excel_decision(self):
        services = self.get_selected_services()
        if not services:
            return '제외'
        # 기초연금 등 현재 선택된 보장 중 하나라도 제외 판정이면 제외, 모두 적합이면 책정
        tags = [self.service_result_tag(s, services) for s in services]
        return '책정' if tags and all(t == '적합' for t in tags) else '제외'

    def current_excel_income(self):
        services = self.get_selected_services()
        target = services[0] if services else ''
        return self.get_income_num(target)

    def find_rrn_by_name(self, name):
        name = self.clean(name)
        if not name or self.current_df is None or self.current_df.empty or '성명' not in self.current_df.columns:
            return ''
        rows = self.current_df[self.current_df['성명'].astype(str).map(self.clean) == name]
        if rows.empty and name.replace(' ', ''):
            rows = self.current_df[self.current_df['성명'].astype(str).map(lambda x: self.clean(x).replace(' ', '')) == name.replace(' ', '')]
        if rows.empty or '주민등록번호' not in rows.columns:
            return ''
        rrn = self.normalize_rrn(rows.iloc[0]['주민등록번호'])
        digits = re.sub(r'[^0-9]', '', rrn)
        return f'{digits[:6]}-{digits[6:13]}' if len(digits) >= 13 else rrn

    def get_selected_excel_file_rows(self):
        return [r for r in getattr(self, 'file_rows', []) if getattr(r, 'path', '')]

    def read_single_file_applicant_name(self):
        """가구원별 엑셀 파일이 1개일 때 A열의 '성명' 바로 아래 값을 신청인 이름으로 사용합니다."""
        rows = self.get_selected_excel_file_rows()
        if len(rows) != 1:
            return ''
        try:
            raw, _ = self.read_excel(rows[0].path, self.e_file_password.get().strip())
            # A열에서 '성명'을 찾고, 바로 아래의 첫 유효값을 반환합니다.
            first_col = raw.iloc[:, 0].tolist() if raw is not None and raw.shape[1] >= 1 else []
            for idx, value in enumerate(first_col):
                if self.clean(value) == '성명':
                    for next_idx in range(idx + 1, len(first_col)):
                        name = self.clean(first_col[next_idx])
                        if name:
                            return name
                    break
        except Exception:
            pass
        # A열 기준으로 찾지 못한 경우, 이미 반영된 데이터의 성명 1건을 보조적으로 사용합니다.
        try:
            if self.current_df is not None and not self.current_df.empty and '성명' in self.current_df.columns:
                names = [self.clean(x) for x in self.current_df['성명'].dropna().unique().tolist() if self.clean(x)]
                if len(names) == 1:
                    return names[0]
        except Exception:
            pass
        return ''

    def open_excel_entry_popup(self):
        if not self.basic_pension_excel_enabled():
            messagebox.showwarning('안내', '엑셀 반영 기능은 기초연금 보장 선택 시에만 사용할 수 있습니다.')
            return
        if self.current_df is None or self.current_df.empty:
            messagebox.showwarning('안내', '먼저 첨부 파일 반영하기를 실행하세요.')
            return
        self.validate_date()
        if not self.e_date.get().strip():
            return
        pop = tk.Toplevel(self.root)
        pop.title('엑셀 반영 정보 입력')
        pop.transient(self.root)
        pop.grab_set()
        frm = ttk.Frame(pop, padding=14)
        frm.pack(fill='both', expand=True)
        fields = {}
        def add_row(r, label, widget):
            ttk.Label(frm, text=label).grid(row=r, column=0, sticky='w', padx=4, pady=5)
            widget.grid(row=r, column=1, sticky='ew', padx=4, pady=5)
            return widget
        frm.columnconfigure(1, weight=1)
        single_file_mode = len(self.get_selected_excel_file_rows()) == 1
        e_dong = add_row(0, '행정동', ttk.Entry(frm, width=28))
        e_name = add_row(1, '신청인 이름', ttk.Entry(frm, width=28))
        e_spouse = add_row(2, '배우자 이름', ttk.Entry(frm, width=28))
        cb_type = add_row(3, '수급유형', ttk.Combobox(frm, state='readonly', width=25, values=['노인단독', '노인부부1인', '노인부부2인']))
        cb_bigo1 = add_row(4, '비고1', ttk.Combobox(frm, state='readonly', width=25, values=['', '부부동시신청', '배우자 기수급자']))

        if single_file_mode:
            applicant_name = self.read_single_file_applicant_name()
            if applicant_name:
                e_name.insert(0, applicant_name)
            e_name.config(state='readonly')
            e_spouse.config(state='disabled')
            cb_type.set('노인단독')
            cb_type.config(state='disabled')
            cb_bigo1.set('')
            cb_bigo1.config(state='disabled')
            ttk.Label(frm, text='가구원별 엑셀 파일이 1개라서 신청인/노인단독/비고1 공란으로 자동 처리됩니다.', foreground='gray').grid(row=5, column=0, columnspan=2, sticky='w', padx=4, pady=(2, 0))
            decision_row = 6
            button_row = 7
        else:
            house = self.cb_house.get().strip()
            if '부부' in house:
                cb_type.set('노인부부2인' if '2인' in house or '동시' in house else '노인부부1인')
            else:
                cb_type.set('노인단독')
            cb_bigo1.set('')
            decision_row = 5
            button_row = 6

        ttk.Label(frm, text=f'조사자 결정은 소득인정액 기준으로 자동 반영됩니다: {self.current_excel_decision()}', foreground='gray').grid(row=decision_row, column=0, columnspan=2, sticky='w', padx=4, pady=(4, 8))
        def save():
            name = e_name.get().strip()
            if not name:
                messagebox.showwarning('안내', '신청인 이름을 입력하세요.', parent=pop)
                return
            entry = {
                'gu': self.cb_gu.get().strip() or '분당구',
                'dong': e_dong.get().strip(),
                'applicant_name': name,
                'spouse_name': e_spouse.get().strip(),
                'benefit_type': cb_type.get().strip(),
                'decision': self.current_excel_decision(),
                'bigo1': cb_bigo1.get().strip(),
                'applicant_rrn': self.find_rrn_by_name(name),
                'spouse_rrn': self.find_rrn_by_name(e_spouse.get().strip()) if e_spouse.get().strip() else '',
                'apply_date': self.format_date(),
                'income': self.current_excel_income(),
            }
            self.excel_entries.append(entry)
            self.refresh_excel_entry_list()
            pop.destroy()
        btn = ttk.Frame(frm)
        btn.grid(row=button_row, column=0, columnspan=2, sticky='e', pady=(8, 0))
        ttk.Button(btn, text='저장', command=save).pack(side='left', padx=4)
        ttk.Button(btn, text='닫기', command=pop.destroy).pack(side='left', padx=4)
        pop.wait_window()

    def refresh_excel_entry_list(self):
        if not hasattr(self, 'excel_entry_listbox'):
            return
        self.excel_entry_listbox.delete(0, 'end')
        for i, e in enumerate(self.excel_entries, start=1):
            self.excel_entry_listbox.insert('end', f"{i}. {e.get('applicant_name','')} ({e.get('decision','')})")

    def delete_selected_excel_entry(self):
        sel = list(self.excel_entry_listbox.curselection()) if hasattr(self, 'excel_entry_listbox') else []
        if not sel:
            messagebox.showwarning('안내', '삭제할 신청인을 선택하세요.')
            return
        for idx in reversed(sel):
            if 0 <= idx < len(self.excel_entries):
                del self.excel_entries[idx]
        self.refresh_excel_entry_list()

    def clear_excel_entries(self):
        if not self.excel_entries:
            return
        if messagebox.askyesno('확인', '저장된 엑셀 반영 목록을 모두 삭제할까요?'):
            self.excel_entries.clear()
            self.refresh_excel_entry_list()

    def save_generated_excel_password(self):
        self.generated_excel_password = self.e_generated_excel_password.get().strip()
        messagebox.showinfo('완료', '생성된 엑셀 비밀번호를 저장했습니다.')

    def clear_generated_excel_password(self):
        self.generated_excel_password = ''
        self.e_generated_excel_password.delete(0, 'end')
        messagebox.showinfo('완료', '생성된 엑셀 비밀번호를 초기화했습니다.')

    def apply_excel_open_password(self, path, password):
        """Microsoft Excel COM 기능으로 파일 열기 암호를 적용합니다.
        openpyxl의 시트 보호와 달리, 이 방식은 파일을 열 때 암호 입력창이 뜨게 합니다.
        Windows + Microsoft Excel + pywin32 환경에서만 동작합니다.
        """
        if not password:
            return
        if os.name != 'nt':
            raise RuntimeError('파일 열기 암호는 Windows에서 Microsoft Excel이 설치된 경우에만 적용할 수 있습니다.')
        try:
            import pythoncom
            import win32com.client
        except Exception as e:
            raise RuntimeError('pywin32가 설치되어 있지 않아 파일 열기 암호를 적용할 수 없습니다. BUILD_EXE를 다시 실행해 주세요.') from e

        abs_path = os.path.abspath(path)
        excel = None
        wb_com = None
        pythoncom.CoInitialize()
        try:
            excel = win32com.client.DispatchEx('Excel.Application')
            excel.Visible = False
            excel.DisplayAlerts = False
            wb_com = excel.Workbooks.Open(abs_path, UpdateLinks=0, ReadOnly=False)
            # 51 = xlOpenXMLWorkbook(.xlsx). Password 인자가 파일 열기 암호입니다.
            wb_com.SaveAs(Filename=abs_path, FileFormat=51, Password=password)
            wb_com.Close(SaveChanges=False)
            wb_com = None
        finally:
            if wb_com is not None:
                try:
                    wb_com.Close(SaveChanges=False)
                except Exception:
                    pass
            if excel is not None:
                try:
                    excel.DisplayAlerts = True
                    excel.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()

    def copy_excel_row_style(self, ws, source_row, target_row, max_col=14):
        """템플릿의 샘플 행 서식을 생성 행 전체에 복사합니다."""
        try:
            ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
        except Exception:
            pass
        for c in range(1, max_col + 1):
            src = ws.cell(source_row, c)
            dst = ws.cell(target_row, c)
            if src.has_style:
                dst._style = copy.copy(src._style)
            if src.number_format:
                dst.number_format = src.number_format
            if src.font:
                dst.font = copy.copy(src.font)
            if src.fill:
                dst.fill = copy.copy(src.fill)
            if src.border:
                dst.border = copy.copy(src.border)
            if src.alignment:
                dst.alignment = copy.copy(src.alignment)
            if src.protection:
                dst.protection = copy.copy(src.protection)

    def apply_excel_template_styles(self, ws, start_row, data_count, max_col=14):
        """
        기본 양식에 들어있는 샘플 데이터 5줄의 서식을 모든 생성 행에 적용합니다.
        1~5번째 생성 행은 각 샘플 행의 서식을 사용하고, 6번째 이후는 5번째
        샘플 행 서식을 반복 적용합니다.
        """
        if data_count <= 0:
            return
        sample_count = 5
        last_row = start_row + data_count - 1
        # 행이 부족하면 먼저 삽입하지 않고 해당 행 객체에 서식만 입힙니다.
        for r in range(start_row, last_row + 1):
            sample_row = start_row + min(r - start_row, sample_count - 1)
            if sample_row > ws.max_row:
                sample_row = start_row
            self.copy_excel_row_style(ws, sample_row, r, max_col=max_col)

    def create_excel_file(self):
        if not self.basic_pension_excel_enabled():
            messagebox.showwarning('안내', '엑셀 파일 만들기는 기초연금 보장 선택 시에만 사용할 수 있습니다.')
            return
        if not self.excel_entries:
            messagebox.showwarning('안내', '먼저 엑셀에 반영할 신청인을 저장하세요.')
            return
        template = self.resource_path('basic_pension_review_template.xlsx')
        if not template.exists():
            messagebox.showerror('오류', f'엑셀 양식 파일을 찾을 수 없습니다.\n{template}')
            return
        out = filedialog.asksaveasfilename(title='엑셀 파일 저장', defaultextension='.xlsx', filetypes=[('Excel 파일', '*.xlsx')], initialfile=f'기초연금_심의내역_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx')
        if not out:
            return
        try:
            wb = load_workbook(template)
            ws = wb['책정제외중지'] if '책정제외중지' in wb.sheetnames else wb.active
            start_row = 3
            needed_last_row = start_row + len(self.excel_entries) - 1
            clear_last_row = max(ws.max_row, needed_last_row)
            # 예시 행을 지우고 필요한 만큼 행 준비
            for r in range(start_row, clear_last_row + 1):
                for c in range(1, 15):
                    ws.cell(r, c).value = None
            # 템플릿의 샘플 5줄 서식을 6번째 이후 생성 행에도 반복 적용
            self.apply_excel_template_styles(ws, start_row, len(self.excel_entries), max_col=14)
            for i, entry in enumerate(self.excel_entries, start=1):
                r = start_row + i - 1
                ws.cell(r, 1).value = f'=ROW(A{r})-ROW($A$2)'
                ws.cell(r, 2).value = entry.get('gu') or self.cb_gu.get() or '분당구'
                ws.cell(r, 3).value = entry.get('dong')
                ws.cell(r, 4).value = self.format_date_for_excel(entry.get('apply_date'))
                ws.cell(r, 5).value = entry.get('decision') or '제외'
                ws.cell(r, 6).value = entry.get('applicant_name')
                ws.cell(r, 7).value = entry.get('applicant_rrn') or self.find_rrn_by_name(entry.get('applicant_name'))
                ws.cell(r, 8).value = entry.get('benefit_type')
                ws.cell(r, 9).value = entry.get('income')
                ws.cell(r, 10).value = None
                ws.cell(r, 11).value = entry.get('spouse_name') or None
                ws.cell(r, 12).value = (entry.get('spouse_rrn') or self.find_rrn_by_name(entry.get('spouse_name'))) if entry.get('spouse_name') else None
                ws.cell(r, 13).value = entry.get('bigo1') or None
                ws.cell(r, 14).value = None
            pwd = self.e_generated_excel_password.get().strip() or self.generated_excel_password
            if pwd:
                # 보조 보호: 시트/통합문서 구조 보호도 함께 적용합니다.
                for wsx in wb.worksheets:
                    wsx.protection.sheet = True
                    wsx.protection.password = pwd
                    for row in wsx.iter_rows():
                        for cell in row:
                            cell.protection = Protection(locked=True)
                try:
                    if getattr(wb, 'security', None) is not None:
                        wb.security.workbookPassword = pwd
                        wb.security.lockStructure = True
                except Exception:
                    pass
            wb.save(out)

            if pwd:
                try:
                    self.apply_excel_open_password(out, pwd)
                except Exception as e:
                    messagebox.showerror('오류', f'엑셀 파일은 생성되었지만 파일 열기 암호 적용에 실패했습니다.\n\n{out}\n\n원인: {e}')
                    return

            if pwd:
                messagebox.showinfo('완료', f'파일 열기 암호가 적용된 엑셀 파일을 만들었습니다.\n{out}')
            else:
                messagebox.showinfo('완료', f'엑셀 파일을 만들었습니다.\n{out}')
        except Exception as e:
            messagebox.showerror('오류', f'엑셀 파일 생성 중 오류가 발생했습니다.\n{e}')


    def copy_clipboard(self):
        text = self.txt.get('1.0', 'end').strip()
        if not text:
            messagebox.showwarning('안내', '먼저 상담내역 생성을 눌러주세요.')
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self.root.update()
        self.save_recent_input()
        self.update_char_count()
        messagebox.showinfo('완료', '상담내역이 클립보드에 저장되었습니다.')

    def snapshot(self):
        return {
            'date': self.e_date.get(), 'multi': self.multi_var.get(), 'single_service': self.cb_service.get(),
            'multi_sel': list(self.lb_service.curselection()), 'house': self.cb_house.get(), 'house_direct': self.e_house_direct.get(),
            'home': self.cb_home.get(), 'home_direct': self.e_home_direct.get(), 'income': self.e_income.get(), 'ref': self.ref_var.get(),
            'beneficiary': self.beneficiary_var.get(), 'selected_beneficiary': list(self.selected_beneficiary),
            'deposit': self.e_deposit.get(), 'rent': self.e_rent.get(),
            'multi_income': dict(self.multi_income_values),
            'current_df': self.current_df.copy() if self.current_df is not None else None,
            'text': self.txt.get('1.0','end'),
            'rows': [{'path':r.path,'label':r.display_name,'fname':r.lbl_path.cget('text')} for r in self.file_rows],
            'file_password': self.e_file_password.get(), 'asset_date': self.e_asset_date.get(),
            'supporters': [(r.get('relation').get() if r.get('relation') else '', r['name'].get(), r['birth'].get(), r['income'].get(), r['members'].get()) for r in self.supporter_rows],
            'tabeop_type': self.tabeop_household_type(),
            'youth_not_working': self.youth_not_working_var.get(),
            'disability_review_no': self.e_disability_review_no.get(),
            'disability_decision_date': self.e_disability_decision_date.get(),
            'disability_type': self.e_disability_type.get(),
            'gu': self.cb_gu.get(),
            'excel_entries': list(self.excel_entries),
            'generated_excel_password': self.e_generated_excel_password.get()
        }

    def reset_inputs(self, keep_result=False):
        """입력값과 내부 상태값을 초기화합니다. keep_result=True이면 상담내역 출력은 유지합니다."""
        result_text = self.txt.get('1.0', 'end') if keep_result else ''

        try:
            self.last_snapshot = self.snapshot()
        except Exception:
            pass

        try:
            self.e_date.delete(0, 'end')
        except Exception:
            pass

        try:
            self.multi_var.set(False)
            self.lb_service.selection_clear(0, 'end')
            self.lb_service.grid_remove()
            self.cb_service.grid()
            self.cb_service.set('')
        except Exception:
            pass

        try:
            self.cb_house.set('')
            self.cb_house.config(state='disabled', values=[])
            self.e_house_direct.config(state='normal')
            self.e_house_direct.delete(0, 'end')
            self.e_house_direct.config(state='disabled')
        except Exception:
            pass

        try:
            self.cb_home.set('')
            self.on_home_change()
        except Exception:
            pass

        # 주거 관련 입력값 전체 초기화
        for entry in [self.e_home_direct, self.e_deposit, self.e_rent]:
            try:
                entry.config(state='normal')
                entry.delete(0, 'end')
                entry.config(state='disabled')
            except Exception:
                pass

        try:
            self.e_income.config(state='normal')
            self.e_income.delete(0, 'end')
        except Exception:
            pass

        try:
            self.ref_var.set(False)
            self.beneficiary_var.set(False)
        except Exception:
            pass

        self.selected_beneficiary = []
        self.multi_income_values = {}
        self.current_df = None
        self._last_service_snapshot = tuple()
        self._last_valid_services = tuple()
        self.current_services = []
        self._service_selection_alerting = False
        # 엑셀 반영 저장목록과 생성 엑셀 비밀번호는 입력내용 초기화로 지우지 않습니다.
        # 저장목록은 우측 '선택 삭제'/'전체 삭제'로만 삭제합니다.

        self.multi_income_popup_open = False
        try:
            self.set_tabeop_type('일반가구')
            self.update_tabeop_visibility()
        except Exception:
            pass

        try:
            self.youth_not_working_var.set(False)
            self.e_disability_review_no.delete(0, 'end')
            self.e_disability_decision_date.delete(0, 'end')
            self.update_special_visibility()
            self.update_excel_visibility()
        except Exception:
            pass

        # 저장된 엑셀 비밀번호는 입력내용 초기화로 지우지 않습니다.
        # 비밀번호는 '비밀번호 초기화' 버튼으로만 삭제합니다.

        try:
            self.clear_supporters()
        except Exception:
            pass

        try:
            self.multi_income_frame.pack_forget()
        except Exception:
            pass

        try:
            for r in self.file_rows:
                try:
                    r.frame.destroy()
                except Exception:
                    pass
            self.file_rows = []
            self.add_file_row()
        except Exception:
            pass

        try:
            self.update_service_listbox_display()
        except Exception:
            pass

        try:
            self.update_supporter_visibility()
        except Exception:
            pass

        try:
            self.txt.delete('1.0', 'end')
            if keep_result:
                self.txt.insert('1.0', result_text.rstrip() + '\n')
            self.update_char_count()
            self.update_applicant_preview()
            self.update_required_highlights()
        except Exception:
            pass


    def clear_inputs_keep_result(self):
        self.reset_inputs(keep_result=True)

    def reset_all(self):
        self.reset_inputs(keep_result=False)


    def undo_reset(self):
        s = self.last_snapshot
        if not s:
            messagebox.showwarning('안내', '되돌릴 내용이 없습니다.')
            return
        self.e_date.delete(0,'end'); self.e_date.insert(0, s['date'])
        self.multi_var.set(s['multi'])
        if s['multi']:
            self.cb_service.grid_remove(); self.lb_service.grid(row=0, column=3, sticky='w', padx=4, pady=4)
            self.lb_service.selection_clear(0,'end')
            for i in s.get('multi_sel', []):
                self.lb_service.selection_set(i)
        else:
            self.lb_service.grid_remove(); self.cb_service.grid(); self.cb_service.set(s['single_service'])
            self.lb_service.selection_clear(0,'end')
        self.on_service_change()
        try:
            self.set_tabeop_type(s.get('tabeop_type', '일반가구'))
            self.update_tabeop_visibility()
        except Exception:
            pass
        try:
            self.youth_not_working_var.set(s.get('youth_not_working', False))
            self.e_disability_review_no.delete(0, 'end'); self.e_disability_review_no.insert(0, s.get('disability_review_no', ''))
            self.e_disability_decision_date.delete(0, 'end'); self.e_disability_decision_date.insert(0, s.get('disability_decision_date', ''))
            self.e_disability_type.delete(0, 'end'); self.e_disability_type.insert(0, s.get('disability_type', ''))
            self.cb_gu.set(s.get('gu', '분당구'))
            self.excel_entries = list(s.get('excel_entries', []))
            self.refresh_excel_entry_list()
            self.e_generated_excel_password.delete(0, 'end'); self.e_generated_excel_password.insert(0, s.get('generated_excel_password', ''))
            self.update_special_visibility()
            self.update_excel_visibility()
        except Exception:
            pass
        self.cb_house.set(s['house'])
        self.on_house_change()
        self.e_house_direct.config(state='normal')
        self.e_house_direct.delete(0,'end')
        self.e_house_direct.insert(0, s['house_direct'])
        if self.cb_house.get().strip() != '직접입력':
            self.e_house_direct.config(state='disabled')
        old_home = s.get('home', '')
        restored_home = OLD_HOME_OPTION_MAP.get(old_home, old_home)
        if restored_home not in HOME_OPTIONS:
            restored_home = '기타(직접입력)'
        self.cb_home.set(restored_home); self.on_home_change()
        if restored_home == '기타(직접입력)':
            direct_home = s.get('home_direct', '') or (old_home if old_home not in ('', '직접입력', '기타(직접입력)') else '')
            self.e_home_direct.insert(0, direct_home)
        for entry, val in [(self.e_deposit,s.get('deposit','')),(self.e_rent,s.get('rent',''))]:
            if str(entry.cget('state')) != 'disabled':
                entry.delete(0,'end'); entry.insert(0,val)
        self.e_income.delete(0,'end'); self.e_income.insert(0, s['income'])
        self.e_asset_date.delete(0,'end'); self.e_asset_date.insert(0, s.get('asset_date',''))
        self.ref_var.set(s['ref'])
        self.e_file_password.delete(0,'end'); self.e_file_password.insert(0, s.get('file_password',''))
        self.beneficiary_var.set(s['beneficiary'])
        self.selected_beneficiary = list(s['selected_beneficiary'])
        self.multi_income_values = dict(s['multi_income'])
        self.refresh_multi_income_inputs()
        if self.multi_var.get() and len(self.get_selected_services()) > 1:
            self.e_income.delete(0,'end')
            self.e_income.insert(0, '입력완료' if any(self.multi_income_values.values()) else '')
        self.clear_supporters()
        if self.supporter_enabled():
            self.show_supporter_area()
            for item in s.get('supporters', []):
                if len(item) == 5:
                    rel, nm, bd, inc, mem = item
                    self.add_supporter_row(nm, bd, inc, mem, rel)
                else:
                    nm, bd, inc, mem = item
                    self.add_supporter_row(nm, bd, inc, mem)
            if not s.get('supporters'):
                self.add_supporter_row()
            self.place_add_supporter_button()
        self.current_df = s['current_df'].copy() if s['current_df'] is not None else None
        self.txt.delete('1.0','end'); self.txt.insert('1.0', s['text'])
        for r in self.file_rows:
            r.frame.destroy()
        self.file_rows = []
        for item in s['rows']:
            self.add_file_row()
            r = self.file_rows[-1]
            r.path = item['path']; r.set_display_name(item['label']); r.lbl_path.config(text=item['fname'])
        self.update_service_listbox_display()
        self.update_supporter_visibility()
        self._last_service_snapshot = tuple(self.get_selected_services())
        self.last_snapshot = None
        self.multi_income_popup_open = False


def run_app():
    root = tk.Tk()
    app = App(root)
    root.mainloop()

if __name__ == '__main__':
    run_app()
