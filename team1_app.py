import io
import json
import os
import re
import subprocess
import sys
import tkinter as tk
import tkinter.font as tkfont
from pathlib import Path
from tkinter import ttk, filedialog, messagebox

from openpyxl import load_workbook

from pastel_theme import PALETTE, apply_button_palette, apply_pastel_theme, style_plain_widgets

try:
    import msoffcrypto
except Exception:
    msoffcrypto = None


_NativeButton = tk.Button


class RaisedButton(_NativeButton):
    """샘플 UI와 같은 얇은 테두리의 평면형 공통 버튼입니다."""
    def __init__(self, master=None, **kwargs):
        kwargs.pop('style', None)
        kwargs['relief'] = 'solid'
        kwargs['bd'] = 1
        kwargs['highlightthickness'] = 1
        kwargs.setdefault('font', ('Malgun Gothic', 10, 'bold'))
        kwargs.setdefault('padx', 12)
        kwargs.setdefault('pady', 7)
        kwargs.setdefault('cursor', 'hand2')
        apply_button_palette(kwargs)
        super().__init__(master, **kwargs)


tk.Button = RaisedButton
ttk.Button = RaisedButton

APP_TITLE = '통합조사 상담 프로그램 v1 - 1팀(국기초)'
SERVICE_OPTIONS = ['생계급여', '의료급여', '주거급여', '교육급여', '청년주거급여 분리지급']
EXISTING_BENEFIT_OPTIONS = ['생계급여', '의료급여', '주거급여', '교육급여']
HOUSEHOLD_TYPES = ['단독가구', '부부가구', '부부+자녀가구', '한부모가구', '조손가구', '기타가구', '직접입력']
HOME_OPTIONS = [
    '자가', '사용대차전체', '사용대차부분', '사용대차(급여미지급)', '시설입소',
    '공공건설임대(영구임대)', '공공건설임대(국민임대)', '공공건설임대(매입임대)',
    '전세임대', '민간(전세)', '민간(보증부월세)', '민간(월세)', '기타(직접입력)'
]
OLD_HOME_OPTION_MAP = {
    '사용대차(전체)': '사용대차전체',
    '사용대차(부분)': '사용대차부분',
    '영구임대주택': '공공건설임대(영구임대)',
    '국민임대주택': '공공건설임대(국민임대)',
    '전세': '민간(전세)',
    '월세': '민간(보증부월세)',
    '고시원': '기타(직접입력)',
    '직접입력': '기타(직접입력)',
}
RELATION_OPTIONS = ['가구주', '배우자', '자녀', '부', '모', '직접입력']

# 사용자가 제공한 2026년 복지대상자 선정기준표 값 그대로 반영합니다. 계산/반올림하지 않습니다.
THRESHOLDS_2026 = {
    '생계급여': {1: 820556, 2: 1343773, 3: 1714892, 4: 2078316, 5: 2418150, 6: 2737905, 7: 3044848},
    '의료급여': {1: 1025695, 2: 1679717, 3: 2143614, 4: 2597895, 5: 3022688, 6: 3422381, 7: 3806060},
    '주거급여': {1: 1230834, 2: 2015660, 3: 2572337, 4: 3117474, 5: 3627225, 6: 4106857, 7: 4567272},
    '교육급여': {1: 1282119, 2: 2099646, 3: 2679518, 4: 3247369, 5: 3778360, 6: 4277976, 7: 4757575},
}
SERVICE_THRESHOLD_KEY = {
    '생계급여': '생계급여',
    '의료급여': '의료급여',
    '주거급여': '주거급여',
    '교육급여': '교육급여',
    '청년주거급여 분리지급': '주거급여',
}

CONFIG_DIR = Path.home() / '.integrated_investigation_team1'
CONFIG_DIR.mkdir(exist_ok=True)
PASSWORD_FILE = CONFIG_DIR / 'excel_password.json'
ASSET_DATE_FILE = CONFIG_DIR / 'asset_date.json'
AUTOSAVE_FILE = CONFIG_DIR / 'autosave_v110.json'
RECENT_FILE = CONFIG_DIR / 'recent_cases_v17.json'


def resource_path(filename: str) -> Path:
    try:
        base = Path(sys._MEIPASS)  # type: ignore[attr-defined]
    except Exception:
        base = Path(__file__).resolve().parent
    return base / filename


def only_digits(value: str) -> str:
    return re.sub(r'[^0-9]', '', str(value or ''))


def money_to_int(value: str) -> int:
    raw = str(value or '').replace(',', '').strip()
    neg = raw.startswith('-')
    digits = re.sub(r'[^0-9]', '', raw)
    if not digits:
        return 0
    n = int(digits)
    return -n if neg else n


def fmt_money(value: int | str) -> str:
    n = int(value) if str(value).strip() else 0
    return f'{n:,}'


def clean(value) -> str:
    if value is None:
        return ''
    s = str(value).strip()
    if s.lower() == 'nan':
        return ''
    return re.sub(r'\s+', ' ', s)


def normalize_birth(value) -> str:
    if value is None:
        return ''
    s = str(value).strip()
    if not s:
        return ''
    s = re.sub(r'\.0$', '', s)
    digits = only_digits(s)
    if len(digits) >= 7:
        return f'{digits[:6]}-{digits[6]}'
    if len(digits) == 6:
        return digits
    return s


def open_workbook_any(path: str, password: str = ''):
    p = Path(path)
    try:
        return load_workbook(p, data_only=True)
    except Exception as first_error:
        if not password:
            raise first_error
        if msoffcrypto is None:
            raise RuntimeError('암호화된 엑셀을 읽으려면 msoffcrypto-tool 설치가 필요합니다.') from first_error
        decrypted = io.BytesIO()
        with open(p, 'rb') as f:
            office_file = msoffcrypto.OfficeFile(f)
            office_file.load_key(password=password)
            office_file.decrypt(decrypted)
        decrypted.seek(0)
        return load_workbook(decrypted, data_only=True)


def row_values(ws, row_num: int, max_col: int | None = None) -> list[str]:
    max_col = max_col or (ws.max_column or 1)
    return [clean(ws.cell(row_num, c).value) for c in range(1, max_col + 1)]


def find_header_map(ws):
    max_row = min(ws.max_row or 1, 80)
    max_col = min(ws.max_column or 1, 80)
    for r in range(1, max_row + 1):
        vals = row_values(ws, r, max_col)
        compact = [v.replace(' ', '') for v in vals]
        if '성명' in compact and '주민등록번호' in compact and '소득재산상세분류코드명' in compact:
            header = {}
            for idx, val in enumerate(compact, start=1):
                if val:
                    header[val] = idx
            return r, header
    return None, {}


def parse_excel_records(path: str, password: str = '') -> list[dict]:
    wb = open_workbook_any(path, password=password)
    all_records = []
    for ws in wb.worksheets:
        header_row, h = find_header_map(ws)
        if not header_row:
            continue
        required = ['성명', '주민등록번호', '소득재산상세분류코드명']
        if not all(x in h for x in required):
            continue
        tmp = []
        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            rec = {}
            for key, col in h.items():
                rec[key] = clean(ws.cell(r, col).value)
            if not rec.get('성명') and not rec.get('소득재산상세분류코드명'):
                continue
            tmp.append(rec)
        if tmp:
            reflected = [x for x in tmp if '반영' in str(x.get('반영여부', ''))]
            if reflected:
                all_records.extend(reflected)
            else:
                amount_rows = [x for x in tmp if amount_int(x) not in (None, 0)]
                all_records.extend(amount_rows if amount_rows else tmp)
    return all_records


def amount_int(row: dict):
    for col in ['반영금액', '최신금액']:
        v = row.get(col, '')
        if v is None or str(v).strip() == '':
            continue
        raw = str(v).replace(',', '').strip()
        try:
            return int(float(raw))
        except Exception:
            digits = re.sub(r'[^0-9-]', '', raw)
            if digits not in ['', '-']:
                try:
                    return int(digits)
                except Exception:
                    pass
    return None


def extract_person_from_excel(path: str, password: str = '') -> dict:
    records = parse_excel_records(path, password=password)
    for rec in records:
        name = clean(rec.get('성명', ''))
        birth = normalize_birth(rec.get('주민등록번호', ''))
        if name or birth:
            return {'name': name or '성명 확인 불가', 'birth': birth}
    # 표 형태가 아니면 라벨 형태도 보조 검색
    wb = open_workbook_any(path, password=password)
    for ws in wb.worksheets:
        max_row = min(ws.max_row or 1, 120)
        max_col = min(ws.max_column or 1, 30)
        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for c in row:
                txt = clean(c.value).replace(' ', '')
                if txt in ['성명', '이름']:
                    right = clean(ws.cell(c.row, c.column + 1).value)
                    below = clean(ws.cell(c.row + 1, c.column).value)
                    name = right or below
                    b = ''
                    # 근처 주민등록번호/생년월일 검색
                    for rr in range(max(1, c.row - 3), min(ws.max_row, c.row + 8) + 1):
                        for cc in range(1, min(ws.max_column, 10) + 1):
                            label = clean(ws.cell(rr, cc).value).replace(' ', '')
                            if label in ['주민등록번호', '주민번호', '생년월일', '생년월일성별']:
                                b = clean(ws.cell(rr, cc + 1).value) or clean(ws.cell(rr + 1, cc).value)
                                break
                        if b:
                            break
                    if name:
                        return {'name': name, 'birth': normalize_birth(b)}
    return {'name': '성명 확인 불가', 'birth': ''}


def threshold_for(service: str, count: int) -> int:
    key = SERVICE_THRESHOLD_KEY.get(service, service)
    table = THRESHOLDS_2026[key]
    if count <= 7:
        return table.get(count, table[7])
    inc = table[7] - table[6]
    return table[7] + inc * (count - 7)


def classify_record(row: dict) -> str:
    code = clean(row.get('소득재산상세분류코드', ''))
    name = clean(row.get('소득재산상세분류코드명', ''))
    if code.startswith('1'):
        return '소득사항'
    if code.startswith('2'):
        return '재산사항'
    if code.startswith(('31', '32', '33')) or code.startswith('3'):
        return '공제'
    if code.startswith('4'):
        return '부채'
    if '공제' in name:
        return '공제'
    if any(k in name for k in ['부채', '대출', '차용']):
        return '부채'
    if any(k in name for k in ['재산', '주택', '건물', '토지', '차량', '금융', '예금', '보험', '증권', '임차보증금']):
        return '재산사항'
    if any(k in name for k in ['급여', '연금', '수당', '소득', '무료임차료', '보상금']):
        return '소득사항'
    return '기타'


def is_financial_asset(name: str) -> bool:
    return any(k in name for k in ['금융', '예금', '적금', '보험', '증권', '주식', '펀드', '채권', '퇴직연금', '연금저축', '기타일시금'])


def deduction_kind(row: dict) -> str:
    code = clean(row.get('소득재산상세분류코드', ''))
    if code.startswith('31'):
        return '소득공제'
    if code.startswith('32'):
        return '일반재산 공제'
    if code.startswith('33'):
        return '금융재산 공제'
    return '기타 공제'


class PersonRow:
    def __init__(self, app, idx: int):
        self.app = app
        self.idx = idx
        self.path = ''
        self.name = ''
        self.birth = ''
        self.records: list[dict] = []
        self.frame = ttk.Frame(app.file_area)
        self.frame.pack(fill='x', pady=2)

        self.lbl_no = ttk.Label(self.frame, text=f'가구원 {idx}', width=8)
        self.lbl_no.grid(row=0, column=0, padx=3, pady=3, sticky='w')
        ttk.Label(self.frame, text='관계').grid(row=0, column=1, padx=(3, 1), pady=3)
        self.relation_var = tk.StringVar(value='가구주' if idx == 1 else '')
        self.cb_relation = ttk.Combobox(self.frame, state='readonly', width=10, values=RELATION_OPTIONS, textvariable=self.relation_var)
        self.cb_relation.grid(row=0, column=2, padx=3, pady=3)
        self.cb_relation.bind('<<ComboboxSelected>>', self.on_relation_change)
        self.e_relation = ttk.Entry(self.frame, width=10, state='disabled')
        self.e_relation.grid(row=0, column=3, padx=3, pady=3)
        self.lbl_person = ttk.Label(self.frame, text='선택된 파일 없음', width=24, anchor='w')
        self.lbl_person.grid(row=0, column=4, padx=4, pady=3, sticky='w')
        ttk.Button(self.frame, text='엑셀 선택', command=self.pick_file).grid(row=0, column=5, padx=4, pady=3)
        self.work_var = tk.StringVar(value='')
        self.work_yes = tk.BooleanVar(value=False)
        self.work_no = tk.BooleanVar(value=False)
        work_box = ttk.LabelFrame(self.frame, text='근로능력', padding=(6, 2))
        work_box.grid(row=0, column=6, padx=(12, 4), pady=2, sticky='w')
        ttk.Checkbutton(work_box, text='있음', variable=self.work_yes, command=lambda: self.set_work('있음')).pack(side='left', padx=(0, 4))
        ttk.Checkbutton(work_box, text='없음', variable=self.work_no, command=lambda: self.set_work('없음')).pack(side='left')
        ttk.Button(self.frame, text='삭제', command=lambda: app.remove_person_row(self)).grid(row=0, column=7, padx=4, pady=3)

        # 첨부 엑셀에 없는 공제는 파일 영역 아래의 별도 '소득공제' 카드에서
        # 가구원별로 입력합니다. 저장 데이터 키는 기존과 동일하게 유지합니다.
        self.deduction_frame = ttk.Frame(app.deduction_area)
        self.deduction_frame.pack(fill='x', pady=2)
        self.lbl_deduction_person = ttk.Label(self.deduction_frame, text=f'가구원 {idx}', width=22, anchor='w')
        self.lbl_deduction_person.pack(side='left', padx=(4, 8))
        ttk.Label(self.deduction_frame, text='근로유인 소득공제').pack(side='left', padx=(0, 4))
        self.e_work_deduction = ttk.Entry(self.deduction_frame, width=18)
        self.e_work_deduction.pack(side='left')
        ttk.Label(self.deduction_frame, text='원').pack(side='left', padx=(3, 0))
        self.e_work_deduction.bind('<FocusOut>', self.app.format_money_entry)

    def on_relation_change(self, event=None):
        if self.relation_var.get() == '직접입력':
            self.e_relation.config(state='normal')
            self.e_relation.focus_set()
        else:
            self.e_relation.delete(0, 'end')
            self.e_relation.config(state='disabled')

    def set_work(self, value: str):
        self.work_var.set(value)
        self.work_yes.set(value == '있음')
        self.work_no.set(value == '없음')

    def pick_file(self):
        path = filedialog.askopenfilename(filetypes=[('Excel files', '*.xlsx *.xlsm *.xls'), ('All files', '*.*')])
        if not path:
            return
        self.path = path
        self.lbl_person.config(text='읽는 중...')
        self.app.root.update_idletasks()
        try:
            password = self.app.e_file_password.get().strip()
            self.records = parse_excel_records(path, password=password)
            info = extract_person_from_excel(path, password=password)
            self.name = info.get('name') or '성명 확인 불가'
            self.birth = info.get('birth') or ''
            self.refresh_label()
            self.app.root.after(100, self.app.refresh_live_preview)
        except Exception as e:
            self.records = []
            self.name = ''
            self.birth = ''
            self.lbl_person.config(text=f'{Path(path).name} / 성명 확인 불가')
            messagebox.showwarning('엑셀 읽기 안내', f'엑셀에서 성명/생년월일을 자동 확인하지 못했습니다.\n\n{e}')

    def refresh_label(self):
        if self.name:
            self.lbl_person.config(text=f'{self.name}, {self.birth}' if self.birth else self.name)
        elif self.path:
            self.lbl_person.config(text=Path(self.path).name)
        else:
            self.lbl_person.config(text='선택된 파일 없음')
        display_name = self.name or f'가구원 {self.idx}'
        self.lbl_deduction_person.config(text=display_name)

    def relation(self) -> str:
        return self.e_relation.get().strip() if self.relation_var.get() == '직접입력' else self.relation_var.get().strip()

    def is_active(self) -> bool:
        return bool(self.name or self.path)

    def validate(self) -> str:
        if not self.is_active():
            return ''
        if not self.relation():
            return f'가구원 {self.idx}의 관계를 선택하거나 입력하세요.'
        if not self.work_var.get():
            return f'가구원 {self.idx}의 근로능력 있음/없음 중 하나를 선택하세요.'
        return ''

    def display_line(self) -> str:
        relation = self.relation()
        name = self.name or (Path(self.path).stem if self.path else '성명 확인 불가')
        birth = self.birth or '000000-0'
        return f'- {relation} {name}, {birth}, 근로능력 {self.work_var.get()}'


class SupporterRow:
    def __init__(self, app, idx: int):
        self.app = app
        self.frame = ttk.Frame(app.supporter_area)
        self.frame.pack(fill='x', pady=2)
        self.lbl_no = ttk.Label(self.frame, text=f'{idx}', width=3, anchor='center')
        self.lbl_no.grid(row=0, column=0, padx=(0, 4), pady=1)
        self.e_relation = ttk.Entry(self.frame, width=14)
        self.e_relation.grid(row=0, column=1, padx=2, pady=1)
        self.e_name = ttk.Entry(self.frame, width=12)
        self.e_name.grid(row=0, column=2, padx=2, pady=1)
        self.e_birth = ttk.Entry(self.frame, width=10)
        self.e_birth.grid(row=0, column=3, padx=2, pady=1)
        self.e_birth.bind('<FocusOut>', self.format_birth)
        self.ability = {}  # 예: {'생계급여': '없음', '의료급여': '있음'}
        self.ability_status = tk.StringVar(value='미입력')
        self.lbl_ability = ttk.Label(self.frame, textvariable=self.ability_status, width=10, anchor='center')
        self.lbl_ability.grid(row=0, column=4, padx=2, pady=1)
        ttk.Button(self.frame, text='부양능력 입력', command=self.open_ability_popup).grid(row=0, column=5, padx=2, pady=1)
        ttk.Button(self.frame, text='삭제', command=lambda: app.remove_supporter_row(self)).grid(row=0, column=6, padx=2, pady=1)

    def format_birth(self, event=None):
        digits = only_digits(self.e_birth.get())[:6]
        self.e_birth.delete(0, 'end')
        self.e_birth.insert(0, digits)

    def services_to_check(self) -> list[str]:
        return [s for s in self.app.selected_services() if s in ('생계급여', '의료급여')]

    def refresh_ability_status(self):
        needed = self.services_to_check()
        if not needed:
            self.ability_status.set('해당없음')
            return
        if all(self.ability.get(s) in ('있음', '없음') for s in needed):
            self.ability_status.set('입력완료')
        else:
            self.ability_status.set('미입력')

    def open_ability_popup(self):
        needed = self.services_to_check()
        if not needed:
            messagebox.showinfo('안내', '생계급여 또는 의료급여를 선택한 경우에만 부양능력을 입력합니다.')
            return
        win = tk.Toplevel(self.app.root)
        win.title('부양능력 판정 입력')
        win.transient(self.app.root)
        win.grab_set()
        win.resizable(False, False)
        pad = ttk.Frame(win, padding=14)
        pad.pack(fill='both', expand=True)
        vars_ = {}
        r = 0
        for service in needed:
            ttk.Label(pad, text=f'{service} 부양능력').grid(row=r, column=0, sticky='w', padx=(0, 8), pady=4)
            var = tk.StringVar(value=self.ability.get(service, '없음'))
            cb = ttk.Combobox(pad, state='readonly', width=12, values=['있음', '없음'], textvariable=var)
            cb.grid(row=r, column=1, sticky='w', pady=4)
            vars_[service] = var
            r += 1
        def save():
            for service, var in vars_.items():
                self.ability[service] = var.get()
            self.refresh_ability_status()
            win.destroy()
        btns = ttk.Frame(pad)
        btns.grid(row=r, column=0, columnspan=2, sticky='e', pady=(12, 0))
        ttk.Button(btns, text='저장', command=save).pack(side='left', padx=4)
        ttk.Button(btns, text='취소', command=win.destroy).pack(side='left', padx=4)

    def set_enabled(self, enabled: bool):
        state = 'normal' if enabled else 'disabled'
        for w in [self.e_relation, self.e_name, self.e_birth]:
            w.config(state=state)
        for child in self.frame.winfo_children():
            if isinstance(child, ttk.Button):
                child.config(state=state)

    def values(self) -> dict:
        return {
            'relation': self.e_relation.get().strip(),
            'name': self.e_name.get().strip(),
            'birth': only_digits(self.e_birth.get())[:6],
        }

    def has_value(self):
        v = self.values()
        return any(v.values())

    def display_line(self) -> str:
        v = self.values()
        parts = []
        if v['relation']:
            parts.append(f"관계 {v['relation']}")
        if v['name']:
            parts.append(f"성명 {v['name']}")
        if v['birth']:
            parts.append(f"생년월일 {v['birth']}")
        ability_parts = []
        for service in self.services_to_check():
            val = self.ability.get(service, '')
            if val:
                ability_parts.append(f'{service} 부양능력 {val}')
        if ability_parts:
            parts.extend(ability_parts)
        return '- ' + ', '.join(parts)


class App:
    def __init__(self, root):
        self.root = root
        self.return_to_launcher = False
        self.root.title(APP_TITLE)
        self.root.geometry('1720x940')
        self.root.minsize(1360, 720)
        self.app_font_size = 10
        self.person_rows: list[PersonRow] = []
        self.supporter_rows: list[SupporterRow] = []
        self.income_entries: dict[str, ttk.Entry] = {}
        self.existing_benefit_vars: dict[str, tk.BooleanVar] = {}
        self.last_snapshot = None
        self._autosave_job = None
        self._restoring = False
        self._live_job = None
        self._last_preview_text = ''
        self._invalid_widgets = []
        apply_pastel_theme(self.root, self.app_font_size)
        self.setup_fonts(self.app_font_size)

        # 좌우 분할 화면: 왼쪽은 입력영역, 오른쪽은 상담내역 출력창입니다.
        main_pane = ttk.PanedWindow(root, orient='horizontal')
        main_pane.pack(fill='both', expand=True)
        left_container = ttk.Frame(main_pane, style='Window.TFrame')
        right_container = ttk.Frame(main_pane, style='Window.TFrame')
        main_pane.add(left_container, weight=3)
        main_pane.add(right_container, weight=2)
        # 최초 실행 시 왼쪽 입력영역의 주요 항목이 한 화면에 보이는 폭을 확보합니다.
        def _stabilize_sash():
            try:
                target = min(1120, max(980, self.root.winfo_width() - 560))
                if main_pane.sashpos(0) < 900:
                    main_pane.sashpos(0, target)
            except tk.TclError:
                pass
        for delay in (80, 300, 800):
            self.root.after(delay, _stabilize_sash)

        # 입력영역은 항목이 많으므로 스크롤 가능하게 구성합니다.
        left_canvas = tk.Canvas(left_container, highlightthickness=0, background=PALETTE['window'])
        left_scroll = ttk.Scrollbar(left_container, orient='vertical', command=left_canvas.yview)
        left_canvas.configure(yscrollcommand=left_scroll.set)
        left_canvas.pack(side='left', fill='both', expand=True)
        left_scroll.pack(side='right', fill='y')
        self.input_root = ttk.Frame(left_canvas, style='Window.TFrame', padding=(4, 4))
        left_window = left_canvas.create_window((0, 0), window=self.input_root, anchor='nw')

        def _sync_scroll(event=None):
            left_canvas.configure(scrollregion=left_canvas.bbox('all'))
            left_canvas.itemconfigure(left_window, width=left_canvas.winfo_width())
        self.input_root.bind('<Configure>', _sync_scroll)
        left_canvas.bind('<Configure>', _sync_scroll)

        def _is_descendant(widget, ancestor):
            current = widget
            while current is not None:
                if current == ancestor:
                    return True
                try:
                    parent_name = current.winfo_parent()
                    current = current._nametowidget(parent_name) if parent_name else None
                except Exception:
                    return False
            return False

        def _on_mousewheel(event):
            # 왼쪽 입력영역에서만 왼쪽 스크롤을 움직입니다.
            # 오른쪽 상담내역 Text는 자체 스크롤 동작을 그대로 사용합니다.
            if _is_descendant(event.widget, left_container):
                delta = -1 if event.delta > 0 else 1
                left_canvas.yview_scroll(delta * 3, 'units')
                return 'break'
            return None
        left_canvas.bind_all('<MouseWheel>', _on_mousewheel, add='+')

        nav_frame = ttk.Frame(self.input_root, padding=(12, 10, 12, 4), style='Window.TFrame')
        nav_frame.pack(fill='x')
        ttk.Button(nav_frame, text='← 뒤로가기', command=self.go_back_to_launcher).pack(side='left')
        tk.Button(nav_frame, text='복지대상자 선정기준표 보기', command=self.open_selection_standard_pdf,
                  fg='#b00000', font=('Malgun Gothic', self.app_font_size, 'bold')).pack(side='left', padx=(8, 0))

        top = ttk.LabelFrame(self.input_root, text='신청 정보', padding=12, style='Card.TLabelframe')
        top.pack(fill='x', padx=12, pady=(4, 8))

        ttk.Label(top, text='보장구분').grid(row=0, column=0, sticky='w', padx=4, pady=4)
        self.service_vars = {}
        service_frame = ttk.Frame(top)
        service_frame.grid(row=0, column=1, columnspan=7, sticky='w', padx=4, pady=4)
        for i, s in enumerate(SERVICE_OPTIONS):
            var = tk.BooleanVar(value=False)
            self.service_vars[s] = var
            ttk.Checkbutton(service_frame, text=s, variable=var, command=self.refresh_income_inputs).grid(row=0, column=i, sticky='w', padx=(0, 10))
        ttk.Label(top, text='기보장구분').grid(row=1, column=0, sticky='w', padx=4, pady=4)
        existing_benefit_frame = ttk.Frame(top)
        existing_benefit_frame.grid(row=1, column=1, columnspan=7, sticky='w', padx=4, pady=4)
        for i, benefit in enumerate(EXISTING_BENEFIT_OPTIONS):
            var = tk.BooleanVar(value=False)
            self.existing_benefit_vars[benefit] = var
            ttk.Checkbutton(existing_benefit_frame, text=benefit, variable=var).grid(row=0, column=i, sticky='w', padx=(0, 10))

        ttk.Label(top, text='가구유형').grid(row=2, column=0, sticky='w', padx=4, pady=4)
        self.house_var = tk.StringVar(value='단독가구')
        self.cb_house = ttk.Combobox(top, state='readonly', width=18, values=HOUSEHOLD_TYPES, textvariable=self.house_var)
        self.cb_house.grid(row=2, column=1, sticky='w', padx=4, pady=4)
        self.cb_house.bind('<<ComboboxSelected>>', self.on_house_change)
        self.e_house_direct = ttk.Entry(top, width=18, state='disabled')
        self.e_house_direct.grid(row=2, column=2, sticky='w', padx=4, pady=4)

        ttk.Label(top, text='신청일자').grid(row=3, column=0, sticky='w', padx=4, pady=4)
        self.e_date = ttk.Entry(top, width=16)
        self.e_date.grid(row=3, column=1, sticky='w', padx=4, pady=4)

        self.supporter_outer = ttk.LabelFrame(self.input_root, text='부양의무자', padding=8, style='Peach.TLabelframe')
        self.supporter_outer.pack(fill='x', padx=12, pady=(0, 8))
        no_sup_frame = ttk.Frame(self.supporter_outer)
        no_sup_frame.pack(fill='x', pady=(0, 4))
        self.no_supporter_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(no_sup_frame, text='부양의무자 없음', variable=self.no_supporter_var, command=self.toggle_no_supporter).pack(side='left', padx=4)

        header = ttk.Frame(self.supporter_outer)
        header.pack(fill='x')
        ttk.Label(header, text='', width=3, anchor='center').grid(row=0, column=0, padx=(0,4), pady=1)
        ttk.Label(header, text='수급자와의 관계', width=14, anchor='center').grid(row=0, column=1, padx=2, pady=1)
        ttk.Label(header, text='성명', width=12, anchor='center').grid(row=0, column=2, padx=2, pady=1)
        ttk.Label(header, text='생년월일', width=10, anchor='center').grid(row=0, column=3, padx=2, pady=1)
        ttk.Label(header, text='부양능력', width=12, anchor='center').grid(row=0, column=4, padx=2, pady=1)
        self.supporter_area = ttk.Frame(self.supporter_outer)
        self.supporter_area.pack(fill='x')
        sf_btn = ttk.Frame(self.supporter_outer)
        sf_btn.pack(fill='x', pady=(4,0))
        ttk.Button(sf_btn, text='부양의무자 추가', command=self.add_supporter_row).pack(side='left', padx=4)
        self.add_supporter_row()

        self.house_frame = ttk.LabelFrame(self.input_root, text='주거 정보', padding=10, style='Mint.TLabelframe')
        self.house_frame.pack(fill='x', padx=12, pady=(0, 8))
        ttk.Label(self.house_frame, text='주거유형').grid(row=0, column=0, padx=4, pady=4, sticky='w')
        self.home_var = tk.StringVar(value='자가')
        self.cb_home = ttk.Combobox(self.house_frame, state='readonly', width=28, values=HOME_OPTIONS, textvariable=self.home_var)
        self.cb_home.grid(row=0, column=1, padx=4, pady=4, sticky='w')
        self.cb_home.bind('<<ComboboxSelected>>', self.on_home_change)
        ttk.Label(self.house_frame, text='주거유형 직접입력').grid(row=0, column=2, padx=4, pady=4, sticky='w')
        self.e_home_direct = ttk.Entry(self.house_frame, width=24, state='disabled')
        self.e_home_direct.grid(row=0, column=3, columnspan=2, padx=4, pady=4, sticky='w')
        ttk.Label(self.house_frame, text='보증금').grid(row=1, column=0, padx=4, pady=4)
        self.e_deposit = ttk.Entry(self.house_frame, width=16)
        self.e_deposit.grid(row=1, column=1, padx=4, pady=4)
        ttk.Label(self.house_frame, text='임대료').grid(row=1, column=2, padx=4, pady=4)
        self.e_rent = ttk.Entry(self.house_frame, width=16)
        self.e_rent.grid(row=1, column=3, padx=4, pady=4)
        for e in [self.e_deposit, self.e_rent]:
            e.bind('<FocusOut>', self.format_money_entry)
        self.on_home_change()

        fw = ttk.LabelFrame(self.input_root, text='가구원별 엑셀 파일', padding=12, style='Lavender.TLabelframe')
        fw.pack(fill='x', padx=12, pady=(0, 8))
        pwf = ttk.Frame(fw)
        pwf.pack(fill='x', pady=(0, 6))
        ttk.Label(pwf, text='엑셀 비밀번호').pack(side='left', padx=4)
        self.e_file_password = ttk.Entry(pwf, width=20, show='*')
        self.e_file_password.pack(side='left', padx=4)
        saved_pwd = self.load_json(PASSWORD_FILE).get('password', '')
        if saved_pwd:
            self.e_file_password.insert(0, saved_pwd)
        ttk.Button(pwf, text='비밀번호 저장', command=self.save_excel_password).pack(side='left', padx=4)
        ttk.Button(pwf, text='비밀번호 초기화', command=self.clear_excel_password).pack(side='left', padx=4)
        self.file_area = ttk.Frame(fw)
        self.file_area.pack(fill='x')
        fb = ttk.Frame(fw)
        fb.pack(fill='x', pady=(8, 0))
        ttk.Button(fb, text='가구원 추가', command=self.add_person_row).pack(side='left', padx=4)
        tk.Button(fb, text='첨부 파일 반영하기', command=self.reload_person_files, fg='black', font=('Malgun Gothic', self.app_font_size, 'bold')).pack(side='left', padx=4)

        self.income_frame = ttk.LabelFrame(self.input_root, text='소득인정액', padding=10, style='Blue.TLabelframe')
        self.income_frame.pack(fill='x', padx=12, pady=(0, 8))
        self.refresh_income_inputs()

        self.deduction_frame = ttk.LabelFrame(self.input_root, text='소득공제', padding=10, style='Mint.TLabelframe')
        self.deduction_frame.pack(fill='x', padx=12, pady=(0, 8))
        ttk.Label(
            self.deduction_frame,
            text='첨부 엑셀에 포함되지 않은 공제 금액을 가구원별로 입력하세요.',
            style='Muted.TLabel',
        ).pack(anchor='w', padx=4, pady=(0, 4))
        self.deduction_area = ttk.Frame(self.deduction_frame)
        self.deduction_area.pack(fill='x')

        asset_df = ttk.Frame(self.input_root, padding=(12, 0, 12, 4))
        asset_df.pack(fill='x')
        ttk.Label(asset_df, text='금융재산 조회기준일(선택사항)').pack(side='left', padx=4)
        self.e_asset_date = ttk.Entry(asset_df, width=18)
        self.e_asset_date.pack(side='left', padx=4)
        saved_asset = self.load_json(ASSET_DATE_FILE).get('asset_date', '')
        if saved_asset:
            self.e_asset_date.insert(0, saved_asset)
        ttk.Button(asset_df, text='조회기준일 저장', command=self.save_asset_date).pack(side='left', padx=4)
        ttk.Button(asset_df, text='조회기준일 초기화', command=self.clear_asset_date).pack(side='left', padx=4)

        bf = ttk.Frame(self.input_root, padding=12)
        bf.pack(fill='x')
        tk.Button(bf, text='상담내역 저장', command=self.copy_clipboard, fg='black', font=('Malgun Gothic', self.app_font_size, 'bold')).pack(side='left', padx=4)
        tk.Button(bf, text='입력내용 초기화', command=self.reset_all, fg='#b00000', font=('Malgun Gothic', self.app_font_size, 'bold')).pack(side='left', padx=4)
        ttk.Button(bf, text='초기화 되돌리기', command=self.undo_reset).pack(side='left', padx=4)
        ttk.Label(bf, text='글자 크기').pack(side='left', padx=(18, 4))
        self.font_size_var = tk.StringVar(value='10')
        cb_font = ttk.Combobox(bf, state='readonly', width=5, values=['9', '10', '11', '12', '13', '14', '16', '18'], textvariable=self.font_size_var)
        cb_font.pack(side='left', padx=4)
        cb_font.bind('<<ComboboxSelected>>', self.change_font_size)

        self.ref_var = tk.BooleanVar(value=False)
        ref_frame = ttk.Frame(self.input_root, padding=(12, 0, 12, 8))
        ref_frame.pack(fill='x')
        ttk.Label(ref_frame, text='참고자료 미조사').pack(side='left', padx=(4, 2))
        ttk.Checkbutton(ref_frame, variable=self.ref_var).pack(side='left', padx=(0, 8))
        ttk.Label(ref_frame, text='* 현재 소득인정액이 기준을 초과하여 참고자료 조사가 필요없을 경우', foreground='gray').pack(side='left', padx=4)

        shortcut_row = ttk.Frame(self.input_root, padding=(12, 0, 12, 10))
        shortcut_row.pack(fill='x')
        shortcut_card = tk.Frame(shortcut_row, bd=1, relief='solid', padx=8, pady=4, background=PALETTE['blue'], highlightbackground=PALETTE['border'])
        shortcut_card.pack(side='left', padx=4)
        tk.Label(shortcut_card, text='⌨ 빠른 단축키', background=PALETTE['blue'], foreground=PALETTE['primary'],
                 font=('Malgun Gothic', 9, 'bold')).grid(row=0, column=0, columnspan=5, sticky='w', pady=(0, 2))
        shortcuts = [('Ctrl+Enter', '생성'), ('Ctrl+S', '저장'), ('Ctrl+R', '초기화'), ('Ctrl+Z', '되돌리기'), ('F1', '안내')]
        for col, (key, action) in enumerate(shortcuts):
            cell = tk.Frame(shortcut_card, background=PALETTE['blue'])
            cell.grid(row=1, column=col, padx=(0 if col == 0 else 7, 0), sticky='w')
            tk.Label(cell, text=key, bd=1, relief='ridge', padx=4, pady=1,
                     background=PALETTE['entry'], foreground=PALETTE['primary'], font=('Malgun Gothic', 8, 'bold')).pack(side='left')
            tk.Label(cell, text=action, background=PALETTE['blue'], foreground=PALETTE['text'],
                     font=('Malgun Gothic', 8)).pack(side='left', padx=(3, 0))

        body = ttk.Frame(right_container, padding=(8, 12, 12, 12), style='Window.TFrame')
        body.pack(fill='both', expand=True)

        right_top = ttk.Frame(body, style='Window.TFrame')
        right_top.pack(fill='x', pady=(0, 6))
        ttk.Label(
            right_top,
            text='상담내역 미리보기',
            font=('Malgun Gothic', self.app_font_size + 1, 'bold'),
            foreground=PALETTE['primary'],
            background=PALETTE['window'],
        ).pack(side='left', padx=(4, 0))
        ttk.Button(right_top, text='최근 입력 10건', command=self.show_recent_cases).pack(side='right', padx=4)
        self.char_count_var = tk.StringVar(value='0자')
        ttk.Label(right_top, textvariable=self.char_count_var).pack(side='right', padx=8)

        preview_frame = ttk.Frame(body, style='Card.TFrame')
        preview_frame.pack(fill='both', expand=True)
        self.txt = tk.Text(preview_frame, wrap='word', font=('Malgun Gothic', 10), undo=True,
                           background=PALETTE['entry'], foreground=PALETTE['text'], insertbackground=PALETTE['text'],
                           selectbackground=PALETTE['blue_active'], relief='solid', bd=1)
        right_scroll = ttk.Scrollbar(preview_frame, orient='vertical', command=self.txt.yview)
        self.txt.configure(yscrollcommand=right_scroll.set)
        self.txt.pack(side='left', fill='both', expand=True)
        right_scroll.pack(side='right', fill='y')
        self.txt.tag_configure('fit', foreground='#0b57d0')
        self.txt.tag_configure('unfit', foreground='#c00000')
        self.txt.tag_configure('changed', background='#fff2a8')
        self.txt.bind('<<Modified>>', self.on_preview_modified)
        self.txt.edit_modified(False)
        self.add_person_row()
        self.bind_shortcuts()
        self.setup_validation_styles()
        self.setup_live_preview()
        self.update_checklist_and_required_styles()
        self.root.protocol('WM_DELETE_WINDOW', self.on_close)
        self.root.after(300, self.restore_autosave_if_available)
        self.schedule_autosave()
        style_plain_widgets(self.root)

    def setup_fonts(self, size=10):
        self.app_font_size = int(size)
        for name in ['TkDefaultFont', 'TkTextFont', 'TkFixedFont', 'TkMenuFont', 'TkHeadingFont', 'TkCaptionFont', 'TkSmallCaptionFont', 'TkIconFont', 'TkTooltipFont']:
            try:
                tkfont.nametofont(name).configure(family='Malgun Gothic', size=self.app_font_size)
            except Exception:
                pass
        try:
            style = ttk.Style()
            style.configure('.', font=('Malgun Gothic', self.app_font_size))
        except Exception:
            pass

    def selected_services(self) -> list[str]:
        return [s for s, v in self.service_vars.items() if v.get()]

    def selected_existing_benefits(self) -> list[str]:
        return [s for s, v in self.existing_benefit_vars.items() if v.get()]

    def refresh_income_inputs(self):
        if not hasattr(self, 'income_frame'):
            return
        existing = {s: e.get() for s, e in self.income_entries.items()}
        for child in self.income_frame.winfo_children():
            child.destroy()
        self.income_entries = {}
        services = self.selected_services()
        self.update_supporter_visibility()
        if not services:
            ttk.Label(self.income_frame, text='보장구분을 선택하면 보장별 소득인정액 입력칸이 표시됩니다.').pack(side='left', padx=4)
            return
        for i, s in enumerate(services):
            ttk.Label(self.income_frame, text=s).grid(row=0, column=i*2, padx=(4 if i == 0 else 12, 3), pady=2, sticky='w')
            e = ttk.Entry(self.income_frame, width=16)
            e.grid(row=0, column=i*2+1, padx=3, pady=2, sticky='w')
            if s in existing:
                e.insert(0, existing[s])
            e.bind('<FocusOut>', self.format_money_entry)
            self.income_entries[s] = e
        if hasattr(self, 'supporter_rows'):
            self.refresh_supporter_ability_statuses()

    def is_housing_only(self) -> bool:
        return self.selected_services() == ['주거급여']

    def update_supporter_visibility(self):
        if not hasattr(self, 'supporter_outer'):
            return
        if self.is_housing_only():
            self.supporter_outer.pack_forget()
        elif not self.supporter_outer.winfo_manager():
            options = {'fill': 'x', 'padx': 12, 'pady': (0, 8)}
            if hasattr(self, 'house_frame'):
                options['before'] = self.house_frame
            self.supporter_outer.pack(**options)

    def on_house_change(self, event=None):
        if self.house_var.get() == '직접입력':
            self.e_house_direct.config(state='normal')
        else:
            self.e_house_direct.delete(0, 'end')
            self.e_house_direct.config(state='disabled')

    def household_display(self) -> str:
        return self.e_house_direct.get().strip() or '직접입력' if self.house_var.get() == '직접입력' else self.house_var.get()

    def set_state(self, entry, state='normal', clear=False):
        entry.config(state='normal')
        if clear:
            entry.delete(0, 'end')
        entry.config(state=state)

    def on_home_change(self, event=None):
        for e in [self.e_deposit, self.e_rent]:
            self.set_state(e, 'disabled', clear=True)
        self.set_state(self.e_home_direct, 'disabled', clear=True)
        k = self.home_var.get().strip()
        if k == '민간(전세)':
            self.set_state(self.e_deposit, 'normal')
        elif k in (
            '공공건설임대(영구임대)', '공공건설임대(국민임대)',
            '공공건설임대(매입임대)', '민간(보증부월세)'
        ):
            self.set_state(self.e_deposit, 'normal')
            self.set_state(self.e_rent, 'normal')
        elif k == '전세임대':
            self.set_state(self.e_rent, 'normal')
        elif k == '민간(월세)':
            self.set_state(self.e_rent, 'normal')
        elif k == '기타(직접입력)':
            self.set_state(self.e_home_direct, 'normal')
            for e in [self.e_deposit, self.e_rent]:
                self.set_state(e, 'normal')

    def format_money_entry(self, event=None):
        widget = event.widget if event else None
        if not widget or str(widget.cget('state')) == 'disabled':
            return
        raw = only_digits(widget.get())
        widget.delete(0, 'end')
        if raw:
            widget.insert(0, f'{int(raw):,}')

    def load_json(self, path: Path) -> dict:
        try:
            if path.exists():
                return json.loads(path.read_text(encoding='utf-8'))
        except Exception:
            pass
        return {}

    def save_excel_password(self):
        PASSWORD_FILE.write_text(json.dumps({'password': self.e_file_password.get()}, ensure_ascii=False), encoding='utf-8')
        messagebox.showinfo('안내', '엑셀 비밀번호를 저장했습니다.')

    def clear_excel_password(self):
        PASSWORD_FILE.write_text(json.dumps({'password': ''}, ensure_ascii=False), encoding='utf-8')
        self.e_file_password.delete(0, 'end')
        messagebox.showinfo('안내', '엑셀 비밀번호를 초기화했습니다.')

    def save_asset_date(self):
        ASSET_DATE_FILE.write_text(json.dumps({'asset_date': self.e_asset_date.get()}, ensure_ascii=False), encoding='utf-8')
        messagebox.showinfo('안내', '금융재산 조회기준일을 저장했습니다.')

    def clear_asset_date(self):
        ASSET_DATE_FILE.write_text(json.dumps({'asset_date': ''}, ensure_ascii=False), encoding='utf-8')
        self.e_asset_date.delete(0, 'end')
        messagebox.showinfo('안내', '금융재산 조회기준일을 초기화했습니다.')

    def open_selection_standard_pdf(self):
        path = resource_path('selection_standard_2026.pdf')
        if not path.exists():
            messagebox.showerror('오류', f'선정기준 PDF를 찾을 수 없습니다.\n{path}')
            return
        try:
            if sys.platform.startswith('win'):
                os.startfile(str(path))  # type: ignore[attr-defined]
            elif sys.platform == 'darwin':
                subprocess.Popen(['open', str(path)])
            else:
                subprocess.Popen(['xdg-open', str(path)])
        except Exception as e:
            messagebox.showerror('오류', f'PDF 열기 중 오류가 발생했습니다.\n{e}')

    def add_person_row(self):
        row = PersonRow(self, len(self.person_rows) + 1)
        self.person_rows.append(row)
        self.renumber_rows()

    def remove_person_row(self, row: PersonRow):
        if len(self.person_rows) <= 1:
            messagebox.showinfo('안내', '가구원 행은 최소 1개가 필요합니다.')
            return
        row.frame.destroy()
        row.deduction_frame.destroy()
        self.person_rows = [r for r in self.person_rows if r is not row]
        self.renumber_rows()

    def renumber_rows(self):
        for i, r in enumerate(self.person_rows, start=1):
            r.idx = i
            r.lbl_no.config(text=f'가구원 {i}')
            r.lbl_deduction_person.config(text=r.name or f'가구원 {i}')

    def toggle_no_supporter(self):
        disabled = self.no_supporter_var.get()
        for r in self.supporter_rows:
            r.set_enabled(not disabled)
        if disabled:
            for r in self.supporter_rows:
                r.ability_status.set('해당없음')

    def refresh_supporter_ability_statuses(self):
        for r in self.supporter_rows:
            r.refresh_ability_status()

    def add_supporter_row(self):
        row = SupporterRow(self, len(self.supporter_rows) + 1)
        self.supporter_rows.append(row)
        self.renumber_supporters()
        if hasattr(self, 'no_supporter_var') and self.no_supporter_var.get():
            row.set_enabled(False)

    def remove_supporter_row(self, row: SupporterRow):
        if len(self.supporter_rows) <= 1:
            messagebox.showinfo('안내', '최소 1개의 부양의무자 입력칸은 유지됩니다.')
            return
        row.frame.destroy()
        self.supporter_rows = [r for r in self.supporter_rows if r is not row]
        self.renumber_supporters()

    def renumber_supporters(self):
        for i, r in enumerate(self.supporter_rows, start=1):
            r.lbl_no.config(text=f'{i}')

    def reload_person_files(self):
        cnt = 0
        for r in self.person_rows:
            if r.path:
                try:
                    password = self.e_file_password.get().strip()
                    r.records = parse_excel_records(r.path, password=password)
                    info = extract_person_from_excel(r.path, password=password)
                    r.name = info.get('name') or '성명 확인 불가'
                    r.birth = info.get('birth') or ''
                    r.refresh_label()
                    cnt += 1
                except Exception:
                    pass
        self.root.after(100, self.refresh_live_preview)
        messagebox.showinfo('안내', f'첨부 파일 {cnt}건을 반영했습니다.')

    def active_rows(self) -> list[PersonRow]:
        return [r for r in self.person_rows if r.is_active()]

    def household_count(self) -> int:
        return max(1, len(self.active_rows()))

    def all_records(self) -> list[dict]:
        records = []
        for r in self.active_rows():
            records.extend(r.records or [])
        return records

    def validate_before_generate(self) -> bool:
        services = self.selected_services()
        if not services:
            messagebox.showwarning('확인', '보장구분을 하나 이상 선택하세요.')
            return False
        if not self.e_date.get().strip():
            messagebox.showwarning('확인', '신청일자를 입력하세요.')
            self.e_date.focus_set()
            return False
        if self.house_var.get() == '직접입력' and not self.e_house_direct.get().strip():
            messagebox.showwarning('확인', '가구유형을 직접 입력하세요.')
            self.e_house_direct.focus_set()
            return False
        if not self.active_rows():
            messagebox.showwarning('확인', '신청 가구원 엑셀 파일을 한 건 이상 첨부하세요.')
            return False
        for service in services:
            entry = self.income_entries.get(service)
            if entry is None or not entry.get().strip():
                messagebox.showwarning('확인', f'{service} 소득인정액을 입력하세요.')
                if entry: entry.focus_set()
                return False
        services = self.selected_services()
        if not services:
            messagebox.showwarning('확인', '보장구분을 1개 이상 선택하세요.')
            return False
        self.refresh_income_inputs()
        for s in services:
            e = self.income_entries.get(s)
            if not e or not e.get().strip():
                messagebox.showwarning('확인', f'{s} 소득인정액을 입력하세요.')
                return False
        rows = self.active_rows()
        if not rows:
            messagebox.showwarning('확인', '가구원별 엑셀 파일을 1개 이상 선택하세요.')
            return False
        for r in rows:
            msg = r.validate()
            if msg:
                messagebox.showwarning('확인', msg)
                return False
        needed = [s for s in services if s in ('생계급여', '의료급여')]
        if needed and not self.no_supporter_var.get():
            supporters = [r for r in self.supporter_rows if r.has_value()]
            if not supporters:
                messagebox.showwarning('확인', '생계급여 또는 의료급여 선택 시 부양의무자를 입력하거나 부양의무자 없음에 체크하세요.')
                return False
            for idx, sup in enumerate(supporters, start=1):
                for s in needed:
                    if sup.ability.get(s) not in ('있음', '없음'):
                        messagebox.showwarning('확인', f'부양의무자 {idx}의 {s} 부양능력 있음/없음을 입력하세요.')
                        return False
        return True

    def asset_date_text(self):
        val = self.e_asset_date.get().strip()
        if not val:
            return ''
        digits = only_digits(val)
        if len(digits) == 8:
            val = f'{digits[:4]}.{digits[4:6]}.{digits[6:8]}.'
        return f' (조회기준일 {val})'

    def category_text(self, category: str) -> str:
        records = self.all_records()
        manual_deductions = [
            (r.name or f'가구원 {r.idx}', money_to_int(r.e_work_deduction.get()))
            for r in self.active_rows()
            if money_to_int(r.e_work_deduction.get())
        ]
        if not records and not (category == '공제' and manual_deductions):
            return '해당사항 없음'
        if category == '소득사항':
            items = {}
            total = 0
            for row in records:
                if classify_record(row) != '소득사항':
                    continue
                amt = amount_int(row)
                if amt in (None, 0):
                    continue
                name = clean(row.get('소득재산상세분류코드명', '')) or '항목명 없음'
                items[name] = items.get(name, 0) + amt
                total += amt
            if not total:
                return '해당사항 없음'
            lines = [f'소득총액 {total:,}원']
            lines += [f'- {k} {v:,}원' for k, v in items.items() if v]
            return '\n'.join(lines)
        if category == '재산사항':
            gen_items, fin_items = {}, {}
            gen_total = fin_total = 0
            for row in records:
                if classify_record(row) != '재산사항':
                    continue
                amt = amount_int(row)
                if amt in (None, 0):
                    continue
                name = clean(row.get('소득재산상세분류코드명', '')) or '항목명 없음'
                if is_financial_asset(name):
                    fin_items[name] = fin_items.get(name, 0) + amt
                    fin_total += amt
                else:
                    gen_items[name] = gen_items.get(name, 0) + amt
                    gen_total += amt
            if not (gen_total or fin_total):
                return '해당사항 없음'
            lines = [f'일반재산 총액 {gen_total:,}원']
            lines += [f'- {k} {v:,}원' for k, v in gen_items.items() if v]
            lines.append('')
            lines.append(f'금융재산 총액 {fin_total:,}원{self.asset_date_text()}')
            lines += [f'- {k} {v:,}원' for k, v in fin_items.items() if v]
            return '\n'.join(lines)
        if category == '공제':
            buckets = [('소득공제', {}), ('일반재산 공제', {}), ('금융재산 공제', {}), ('기타 공제', {})]
            bucket_map = {name: d for name, d in buckets}
            totals = {name: 0 for name, _ in buckets}
            for row in records:
                if classify_record(row) != '공제':
                    continue
                amt = amount_int(row)
                if amt in (None, 0):
                    continue
                kind = deduction_kind(row)
                name = clean(row.get('소득재산상세분류코드명', '')) or '항목명 없음'
                bucket_map.setdefault(kind, {})[name] = bucket_map.setdefault(kind, {}).get(name, 0) + amt
                totals[kind] = totals.get(kind, 0) + amt
            for person_name, amt in manual_deductions:
                item_name = f'{person_name} 근로유인 소득공제'
                bucket_map['소득공제'][item_name] = bucket_map['소득공제'].get(item_name, 0) + amt
                totals['소득공제'] += amt
            lines = []
            for kind, item_map in bucket_map.items():
                total = totals.get(kind, 0)
                if not total:
                    continue
                if lines:
                    lines.append('')
                lines.append(f'{kind} 총액 {total:,}원')
                lines += [f'- {k} {v:,}원' for k, v in item_map.items() if v]
            return '\n'.join(lines) if lines else '해당사항 없음'
        if category == '부채':
            items = {}
            total = 0
            for row in records:
                if classify_record(row) != '부채':
                    continue
                amt = amount_int(row)
                if amt in (None, 0):
                    continue
                name = clean(row.get('소득재산상세분류코드명', '')) or '항목명 없음'
                items[name] = items.get(name, 0) + amt
                total += amt
            if not total:
                return '해당사항 없음'
            lines = [f'부채 총액 {total:,}원']
            lines += [f'- {k} {v:,}원' for k, v in items.items() if v]
            return '\n'.join(lines)
        return '해당사항 없음'

    def format_home(self):
        selected = self.home_var.get().strip()
        k = self.e_home_direct.get().strip() if selected == '기타(직접입력)' else selected
        k = k or '직접입력'
        extra = []
        if selected == '민간(전세)' and self.e_deposit.get().strip():
            extra.append(f'보증금 {self.e_deposit.get().strip()}원')
        elif selected in (
            '공공건설임대(영구임대)', '공공건설임대(국민임대)',
            '공공건설임대(매입임대)', '민간(보증부월세)'
        ):
            if self.e_deposit.get().strip():
                extra.append(f'보증금 {self.e_deposit.get().strip()}원')
            if self.e_rent.get().strip():
                extra.append(f'임대료 {self.e_rent.get().strip()}원')
        elif selected == '전세임대':
            if self.e_rent.get().strip():
                extra.append(f'임대료 {self.e_rent.get().strip()}원')
        elif selected == '민간(월세)':
            if self.e_rent.get().strip():
                extra.append(f'임대료 {self.e_rent.get().strip()}원')
        elif selected == '기타(직접입력)':
            if self.e_deposit.get().strip(): extra.append(f'보증금 {self.e_deposit.get().strip()}원')
            if self.e_rent.get().strip(): extra.append(f'임대료 {self.e_rent.get().strip()}원')
        return f"{k} ({', '.join(extra)})" if k and extra else k

    def section(self, no: int, title: str, content: str) -> str:
        return f'{no}. {title}: {content}' if '\n' not in content else f'{no}. {title}\n{content}'

    def build_preview(self) -> str:
        services = self.selected_services()
        rows = self.active_rows()
        count = self.household_count()
        housing_only = self.is_housing_only()
        supporters = [] if (housing_only or self.no_supporter_var.get()) else [r for r in self.supporter_rows if r.has_value()]

        result_status = {}
        for s in services:
            income = money_to_int(self.income_entries[s].get())
            th = threshold_for(s, count)
            has_able_supporter = (s in ('생계급여', '의료급여') and (not self.no_supporter_var.get()) and any(r.ability.get(s) == '있음' for r in supporters))
            result_status[s] = '부적합' if income > th or has_able_supporter else '책정'

        short = {'생계급여':'생계', '의료급여':'의료', '주거급여':'주거', '교육급여':'교육', '청년주거급여 분리지급':'청년주거급여 분리지급'}
        title_results = ', '.join(f"{short.get(s, s)}({result_status[s]})" for s in services)
        header_lines = [f'<맞춤형복지급여 신청 결과> {title_results}']
        existing_benefits = self.selected_existing_benefits()
        if existing_benefits:
            header_lines.append(f'*기보장내역: {", ".join(existing_benefits)}')
        lines = ['\n'.join(header_lines)]

        applicants = '\n'.join(r.display_line() for r in rows)
        lines.append(f'1. 신청 가구원\n{applicants}')
        lines.append(f'2. 신청일자: {self.e_date.get().strip()}')
        lines.append(f'3. 보장구분: {", ".join(services)}')
        lines.append(f'4. 가구유형: {self.household_display()}')
        lines.append(f'5. 주거유형: {self.format_home()}')
        lines.append(self.section(6, '소득사항', self.category_text('소득사항')))
        lines.append(self.section(7, '재산사항', self.category_text('재산사항')))
        lines.append(self.section(8, '공제', self.category_text('공제')))
        lines.append(self.section(9, '부채', self.category_text('부채')))
        no = 10
        if housing_only:
            pass
        elif self.no_supporter_var.get():
            lines.append(f'{no}. 부양의무자:\n부양의무자 없음'); no += 1
        elif supporters:
            txt = '\n'.join(r.display_line() for r in supporters)
            lines.append(f'{no}. 부양의무자:\n{txt}'); no += 1
        income_text = '\n'.join(f"- {s}: {self.income_entries[s].get().strip()}원" if self.income_entries[s].get().strip() else f'- {s}:' for s in services)
        lines.append(f'{no}. 소득인정액:\n{income_text}'); no += 1
        result_lines = []
        for s in services:
            income = money_to_int(self.income_entries[s].get())
            th = threshold_for(s, count)
            has_able_supporter = (s in ('생계급여', '의료급여') and (not self.no_supporter_var.get()) and any(r.ability.get(s) == '있음' for r in supporters))
            if income > th:
                result_lines.append(f'- {s}: {count}인 가구 소득인정액 {income:,}원으로 {count}인 가구 소득인정액 {th:,}원 이하 기준을 초과하여 부적합.')
            elif has_able_supporter:
                result_lines.append(f'- {s}: {count}인 가구 소득인정액 {income:,}원으로 {count}인 가구 소득인정액 {th:,}원 이하이나, 부양의무자 부양능력 있음으로 부적합.')
            else:
                result_lines.append(f'- {s}: {count}인 가구 소득인정액 {income:,}원으로 {count}인 가구 소득인정액 {th:,}원 이하로 적합.')
        lines.append(f'{no}. 조사결과:\n' + '\n'.join(result_lines)); no += 1
        if self.ref_var.get():
            lines.append(self.section(no, '참고사항', '현재 소득인정액만으로 기준 초과하여 참고자료 조사가 필요하지 않아 참고자료 미조사')); no += 1
        return '\n\n'.join(lines)

    def generate_preview(self):
        self.update_checklist_and_required_styles(focus_first=True)
        if not self.validate_before_generate():
            return
        text = self.build_preview()
        self.set_preview_text(text, highlight_changes=True)
        self.save_recent_case()
        self.save_autosave()

    def set_preview_text(self, text: str, highlight_changes: bool = False):
        old = self.txt.get('1.0', 'end-1c')
        old_lines = old.splitlines()
        new_lines = text.splitlines()
        self.txt.delete('1.0', 'end')
        self.txt.insert('1.0', text)
        self.apply_result_colors()
        if highlight_changes and old:
            import difflib
            matcher = difflib.SequenceMatcher(a=old_lines, b=new_lines)
            changed = set()
            for tag, a1, a2, b1, b2 in matcher.get_opcodes():
                if tag != 'equal':
                    changed.update(range(b1 + 1, b2 + 1))
            for line_no in changed:
                self.txt.tag_add('changed', f'{line_no}.0', f'{line_no}.end')
            if changed:
                self.root.after(2000, lambda: self.txt.tag_remove('changed', '1.0', 'end'))
        self._last_preview_text = text
        self.update_char_count()
        self.txt.edit_modified(False)

    def on_preview_modified(self, event=None):
        if self.txt.edit_modified():
            self.update_char_count()
            self.txt.edit_modified(False)

    def update_char_count(self):
        if hasattr(self, 'char_count_var'):
            n = len(self.txt.get('1.0', 'end-1c'))
            self.char_count_var.set(f'{n:,}자')

    def setup_validation_styles(self):
        style = ttk.Style()
        style.configure('Invalid.TEntry', fieldbackground=PALETTE['rose'])
        style.configure('Normal.TEntry', fieldbackground=PALETTE['entry'])
        style.configure('Invalid.TCombobox', fieldbackground=PALETTE['rose'])

    def required_status(self):
        items = []
        items.append(('보장구분', bool(self.selected_services()), None))
        items.append(('신청일자', bool(self.e_date.get().strip()), self.e_date))
        direct_ok = self.house_var.get() != '직접입력' or bool(self.e_house_direct.get().strip())
        items.append(('가구유형', direct_ok, self.e_house_direct if self.house_var.get() == '직접입력' else None))
        items.append(('신청 가구원 파일', bool(self.active_rows()), None))
        for service in self.selected_services():
            entry = self.income_entries.get(service)
            items.append((f'{service} 소득인정액', bool(entry and entry.get().strip()), entry))
        for r in self.active_rows():
            items.append((f'가구원 {r.idx} 관계', bool(r.relation()), r.e_relation if r.relation_var.get() == '직접입력' else r.cb_relation))
            items.append((f'가구원 {r.idx} 근로능력', bool(r.work_var.get()), None))
        return items

    def update_checklist_and_required_styles(self, focus_first=False):
        items = self.required_status()
        missing = [(name, widget) for name, ok, widget in items if not ok]
        lines = [('✓ ' if ok else '□ ') + name for name, ok, _ in items]
        if not lines:
            lines = ['□ 보장구분', '□ 신청일자', '□ 신청 가구원 파일']
        if hasattr(self, 'checklist_var'):
            self.checklist_var.set('   |   '.join(lines))
            self.checklist_label.config(fg='#b00000' if missing else '#176b2c', bg='#fff1f1' if missing else '#eef8f0')
        # ttk Entry/Combobox에 누락 표시
        candidates = [self.e_date, self.e_house_direct] + list(self.income_entries.values())
        for w in candidates:
            try: w.configure(style='Normal.TEntry')
            except Exception: pass
        for _, widget in missing:
            if widget is not None:
                try:
                    widget.configure(style='Invalid.TCombobox' if isinstance(widget, ttk.Combobox) else 'Invalid.TEntry')
                except Exception: pass
        if focus_first:
            for _, widget in missing:
                if widget is not None:
                    try: widget.focus_set(); break
                    except Exception: pass
        return not missing

    def setup_live_preview(self):
        def changed(event=None):
            if self._restoring:
                return
            self.update_checklist_and_required_styles()
            if self._live_job:
                try: self.root.after_cancel(self._live_job)
                except Exception: pass
            self._live_job = self.root.after(500, self.refresh_live_preview)
        self.root.bind_all('<KeyRelease>', changed, add='+')
        self.root.bind_all('<ButtonRelease-1>', changed, add='+')
        self.root.bind_all('<<ComboboxSelected>>', changed, add='+')
        self._live_change_handler = changed

    def refresh_live_preview(self):
        self._live_job = None
        if self._restoring:
            return
        self.update_checklist_and_required_styles()
        # 최소 필수값이 준비된 경우에만 팝업 없이 미리보기 갱신
        if not self.selected_services() or not self.active_rows():
            return
        if any(not self.income_entries.get(s) or not self.income_entries[s].get().strip() for s in self.selected_services()):
            return
        try:
            text = self.build_preview()
        except Exception:
            return
        if text != self.txt.get('1.0', 'end-1c'):
            self.set_preview_text(text, highlight_changes=True)
        self.save_autosave()

    def save_recent_case(self):
        state = self.collect_state()
        rows = self.active_rows()
        title = ', '.join([r.name for r in rows if r.name]) or state.get('date') or '이름 없음'
        item = {'title': title, 'saved_at': __import__('datetime').datetime.now().strftime('%Y-%m-%d %H:%M'), 'state': state}
        data = self.load_json(RECENT_FILE)
        cases = data.get('cases', []) if isinstance(data, dict) else []
        cases.insert(0, item)
        # 같은 제목·신청일은 최신 1건만 유지
        unique = []
        seen = set()
        for x in cases:
            key = (x.get('title'), x.get('state', {}).get('date'))
            if key in seen:
                continue
            seen.add(key)
            unique.append(x)
        try:
            CONFIG_DIR.mkdir(parents=True, exist_ok=True)
            temp_file = RECENT_FILE.with_suffix('.tmp')
            temp_file.write_text(json.dumps({'cases': unique[:10]}, ensure_ascii=False, indent=2), encoding='utf-8')
            temp_file.replace(RECENT_FILE)
            return True
        except Exception:
            return False

    def show_recent_cases(self):
        data = self.load_json(RECENT_FILE)
        cases = data.get('cases', []) if isinstance(data, dict) else []
        if not cases:
            messagebox.showinfo('최근 입력', '저장된 최근 입력내역이 없습니다.')
            return
        win = tk.Toplevel(self.root); win.title('최근 입력 10건'); win.geometry('520x330'); win.transient(self.root)
        frame = ttk.Frame(win, padding=12); frame.pack(fill='both', expand=True)
        lb = tk.Listbox(frame, font=('Malgun Gothic', 10)); lb.pack(fill='both', expand=True)
        for x in cases: lb.insert('end', f"{x.get('saved_at','')}  |  {x.get('title','이름 없음')}")
        def load_selected():
            sel = lb.curselection()
            if not sel: return
            self.restore_state(cases[sel[0]].get('state', {})); self.update_checklist_and_required_styles(); win.destroy()
        ttk.Button(frame, text='선택한 내역 불러오기', command=load_selected).pack(pady=(8,0))
        lb.bind('<Double-Button-1>', lambda e: load_selected())

    def copy_clipboard(self):
        text = self.txt.get('1.0', 'end').strip()
        if not text:
            self.generate_preview()
            text = self.txt.get('1.0', 'end').strip()
        if text:
            self.root.clipboard_clear()
            self.root.clipboard_append(text)
            recent_saved = self.save_recent_case()
            self.save_autosave()
            if recent_saved:
                messagebox.showinfo('안내', '상담내역을 클립보드에 복사하고 최근 입력내역에 저장했습니다.')
            else:
                messagebox.showwarning('안내', '상담내역은 복사했지만 최근 입력내역 저장에 실패했습니다.')

    def snapshot(self):
        return self.collect_state()

    def reset_all(self):
        if not messagebox.askyesno('입력내용 초기화', '현재 입력한 내용을 모두 초기화할까요?'):
            return
        self.last_snapshot = self.collect_state()
        self.e_date.delete(0, 'end')
        for e in list(self.income_entries.values()): e.delete(0, 'end')
        self.e_deposit.delete(0, 'end'); self.e_rent.delete(0, 'end'); self.e_home_direct.delete(0, 'end')
        self.e_asset_date.delete(0, 'end')
        self.home_var.set('자가'); self.on_home_change()
        self.ref_var.set(False)
        for v in self.existing_benefit_vars.values(): v.set(False)
        if hasattr(self, 'no_supporter_var'):
            self.no_supporter_var.set(False)
        for v in self.service_vars.values(): v.set(False)
        self.refresh_income_inputs()
        for r in list(self.person_rows):
            r.frame.destroy()
            r.deduction_frame.destroy()
        self.person_rows = []; self.add_person_row()
        for r in list(self.supporter_rows): r.frame.destroy()
        self.supporter_rows = []; self.add_supporter_row()
        self.txt.delete('1.0', 'end')
        self.update_char_count()
        self.update_checklist_and_required_styles()
        self.save_autosave()

    def undo_reset(self):
        if not self.last_snapshot:
            messagebox.showinfo('안내', '되돌릴 초기화 내용이 없습니다.')
            return
        self.restore_state(self.last_snapshot)
        self.last_snapshot = None
        messagebox.showinfo('안내', '초기화 전 입력내용을 복원했습니다.')

    def bind_shortcuts(self):
        self.root.bind('<Control-Return>', lambda e: self.generate_preview())
        self.root.bind('<Control-s>', lambda e: self.copy_clipboard())
        self.root.bind('<Control-S>', lambda e: self.copy_clipboard())
        self.root.bind('<Control-r>', lambda e: self.reset_all())
        self.root.bind('<Control-R>', lambda e: self.reset_all())
        self.root.bind('<Control-z>', lambda e: self.undo_reset())
        self.root.bind('<Control-Z>', lambda e: self.undo_reset())
        self.root.bind('<F1>', lambda e: self.show_shortcuts())

    def show_shortcuts(self):
        messagebox.showinfo('단축키 안내', 'Ctrl+Enter : 상담내역 생성\nCtrl+S : 상담내역 저장(클립보드 복사)\nCtrl+R : 입력내용 초기화\nCtrl+Z : 초기화 되돌리기\nF1 : 단축키 안내')

    def apply_result_colors(self):
        self.txt.tag_remove('fit', '1.0', 'end')
        self.txt.tag_remove('unfit', '1.0', 'end')
        for line_no, line in enumerate(self.txt.get('1.0', 'end-1c').splitlines(), start=1):
            if '부적합' in line:
                self.txt.tag_add('unfit', f'{line_no}.0', f'{line_no}.end')
            elif line.strip().endswith('적합.'):
                self.txt.tag_add('fit', f'{line_no}.0', f'{line_no}.end')

    def collect_state(self):
        return {'services': self.selected_services(), 'existing_benefits': self.selected_existing_benefits(), 'house': self.house_var.get(), 'house_direct': self.e_house_direct.get().strip(),
                'date': self.e_date.get().strip(), 'income': {k:v.get().strip() for k,v in self.income_entries.items()},
                'home': self.home_var.get(), 'home_direct': self.e_home_direct.get().strip(), 'deposit': self.e_deposit.get().strip(), 'rent': self.e_rent.get().strip(),
                'asset_date': self.e_asset_date.get().strip(),
                'ref': self.ref_var.get(), 'no_supporter': self.no_supporter_var.get(),
                'persons': [{'relation':r.relation_var.get(),'relation_direct':r.e_relation.get().strip(),'work':r.work_var.get(),'work_deduction':r.e_work_deduction.get().strip(),'path':r.path} for r in self.person_rows],
                'supporters': [dict(r.values(), ability=r.ability.copy()) for r in self.supporter_rows],
                'preview': self.txt.get('1.0','end-1c')}

    def restore_state(self, data):
        if not isinstance(data, dict): return
        self._restoring=True
        try:
            for name,var in self.service_vars.items(): var.set(name in data.get('services',[]))
            for name,var in self.existing_benefit_vars.items(): var.set(name in data.get('existing_benefits',[]))
            self.refresh_income_inputs()
            self.house_var.set(data.get('house','단독가구')); self.on_house_change()
            if self.house_var.get()=='직접입력': self.e_house_direct.insert(0,data.get('house_direct',''))
            self.e_date.delete(0,'end'); self.e_date.insert(0,data.get('date',''))
            for k,v in data.get('income',{}).items():
                if k in self.income_entries: self.income_entries[k].insert(0,v)
            old_home = data.get('home','자가')
            restored_home = OLD_HOME_OPTION_MAP.get(old_home, old_home)
            if restored_home not in HOME_OPTIONS:
                restored_home = '기타(직접입력)'
            self.home_var.set(restored_home); self.on_home_change()
            if self.home_var.get() == '기타(직접입력)':
                direct_home = data.get('home_direct','') or (old_home if old_home not in ('직접입력', '기타(직접입력)') else '')
                self.e_home_direct.insert(0, direct_home)
            for e,key in [(self.e_deposit,'deposit'),(self.e_rent,'rent')]:
                if str(e.cget('state'))!='disabled': e.insert(0,data.get(key,''))
            self.e_asset_date.delete(0,'end'); self.e_asset_date.insert(0,data.get('asset_date',''))
            self.ref_var.set(bool(data.get('ref',False)))
            for r in list(self.person_rows):
                r.frame.destroy()
                r.deduction_frame.destroy()
            self.person_rows=[]
            for item in data.get('persons',[]) or [{}]:
                self.add_person_row(); r=self.person_rows[-1]
                r.relation_var.set(item.get('relation','가구주')); r.on_relation_change()
                if r.relation_var.get()=='직접입력': r.e_relation.insert(0,item.get('relation_direct',''))
                r.set_work(item.get('work',''))
                r.e_work_deduction.insert(0, item.get('work_deduction',''))
                path=item.get('path','')
                if path and Path(path).exists():
                    r.path=path
                    try:
                        pwd=self.e_file_password.get().strip(); r.records=parse_excel_records(path,pwd); info=extract_person_from_excel(path,pwd)
                        r.name=info.get('name',''); r.birth=info.get('birth','')
                    except Exception: pass
                    r.refresh_label()
            for r in list(self.supporter_rows): r.frame.destroy()
            self.supporter_rows=[]
            for item in data.get('supporters',[]) or [{}]:
                self.add_supporter_row(); r=self.supporter_rows[-1]
                # 예전 저장 자료의 income 키는 무시하여 하위 호환성을 유지합니다.
                for e,key in [(r.e_relation,'relation'),(r.e_name,'name'),(r.e_birth,'birth')]: e.insert(0,item.get(key,''))
                r.ability=dict(item.get('ability',{})); r.refresh_ability_status()
            self.no_supporter_var.set(bool(data.get('no_supporter',False))); self.toggle_no_supporter()
            self.update_supporter_visibility()
            self.set_preview_text(data.get('preview',''), highlight_changes=False)
            self.update_checklist_and_required_styles()
        finally: self._restoring=False

    def save_autosave(self):
        if self._restoring: return
        try: AUTOSAVE_FILE.write_text(json.dumps(self.collect_state(),ensure_ascii=False,indent=2),encoding='utf-8')
        except Exception: pass

    def schedule_autosave(self):
        self.save_autosave(); self._autosave_job=self.root.after(30000,self.schedule_autosave)

    def restore_autosave_if_available(self):
        data=self.load_json(AUTOSAVE_FILE)
        if not data: return
        meaningful=bool(data.get('date') or data.get('services') or data.get('preview') or any(x.get('path') for x in data.get('persons',[])))
        if meaningful and messagebox.askyesno('자동저장 복원','이전에 자동저장된 입력내용을 복원할까요?'): self.restore_state(data)

    def on_close(self):
        self.save_autosave(); self.root.destroy()

    def has_input_content(self):
        try:
            if self.e_date.get().strip() or self.selected_services() or self.txt.get('1.0', 'end-1c').strip():
                return True
            return any(row.name or row.birth or row.path for row in self.person_rows)
        except Exception:
            return True

    def go_back_to_launcher(self):
        if self.has_input_content() and not messagebox.askyesno(
            '초기화면으로 돌아가기',
            '입력 중인 내용이 있습니다. 자동저장 후 초기화면으로 돌아갈까요?'
        ):
            return
        self.save_autosave()
        self.return_to_launcher = True
        self.root.destroy()

    def change_font_size(self, event=None):
        try:
            size = int(self.font_size_var.get())
        except Exception:
            size = 10
        self.setup_fonts(size)
        self.txt.configure(font=('Malgun Gothic', size))


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == '__main__':
    main()
